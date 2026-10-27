# -*- coding: utf-8 -*-
"""
generar_informe_tecnico_gestor_f1.py
====================================
Genera el Informe Técnico y Funcional del Gestor F1 (Mentora IdanIA) en Word.
Mismo método, estructura y estilo visual que el informe del Dashboard de IA.

Salida: doc/Informe_Tecnico_Funcional_Gestor_F1.docx
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "doc" / "Informe_Tecnico_Funcional_Gestor_F1.docx"

MORADO = RGBColor(0x55, 0x31, 0x71)
MORADO_OSCURO = RGBColor(0x3D, 0x23, 0x52)
MORADO_MEDIO = RGBColor(0x6D, 0x4A, 0x96)
GRIS = RGBColor(0x5F, 0x55, 0x6E)
GRIS_CLARO = RGBColor(0x8B, 0x7D, 0x9E)
VERDE = RGBColor(0x1B, 0x9D, 0x6B)
AMBAR = RGBColor(0xB0, 0x6A, 0x00)
ROJO = RGBColor(0xA3, 0x2A, 0x2A)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

DOC = None
TABLE_N = 0


# ── Helpers (mismos del generador de referencia) ────────────────────────────
def add_field(paragraph, instruction, placeholder="1"):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = placeholder
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, text, end):
        run._r.append(el)


def add_toc():
    paragraph = DOC.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    add_field(paragraph, 'TOC \\o "1-2" \\h \\z \\u', "Actualice el índice con F9.")


def set_update_fields():
    settings = DOC.settings.element
    for element in settings.findall(qn("w:updateFields")):
        settings.remove(element)
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true")
    settings.append(update)


def set_cell_bg(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tc_pr.insert_element_before(
        shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
        "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange",
    )


def set_cell_text(cell, text, bold=False, size=9.0, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run("" if text is None else str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color is not None:
        run.font.color.rgb = color
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table, color="D9CCEA", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)
        borders.append(element)
    tbl_pr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
        "w:tblCaption", "w:tblDescription", "w:tblPrChange",
    )


def set_table_cell_margins(table, top=70, bottom=70, left=100, right=100):
    tbl_pr = table._tbl.tblPr
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value)); element.set(qn("w:type"), "dxa")
        margins.append(element)
    tbl_pr.insert_element_before(margins, "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader"); header.set(qn("w:val"), "true")
    tr_pr.insert_element_before(header, "w:tblCellSpacing", "w:jc", "w:hidden", "w:ins", "w:del", "w:trPrChange")


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.insert_element_before(
        cant_split, "w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc",
        "w:hidden", "w:ins", "w:del", "w:trPrChange",
    )


def add_caption(text):
    global TABLE_N
    TABLE_N += 1
    paragraph = DOC.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(f"Tabla {TABLE_N}. {text}")
    run.italic = True
    run.font.size = Pt(8.8)
    run.font.color.rgb = GRIS_CLARO


def add_table(headers, rows, widths, caption=None, font_size=9.0, header_size=8.8,
              bold_first_col=False, alternate=True):
    if caption:
        add_caption(caption)
    table = DOC.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, width in enumerate(widths):
        table.columns[index].width = Cm(width)
    set_table_borders(table)
    set_table_cell_margins(table)

    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        header_cells[index].width = Cm(widths[index])
        set_cell_text(header_cells[index], header, bold=True, size=header_size, color=BLANCO)
        set_cell_bg(header_cells[index], "553171")
    repeat_header(table.rows[0]); prevent_split(table.rows[0])

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for column_index, value in enumerate(row):
            cells[column_index].width = Cm(widths[column_index])
            bold = bold_first_col and column_index == 0
            set_cell_text(cells[column_index], value, bold=bold, size=font_size)
            if alternate and row_index % 2 == 1:
                set_cell_bg(cells[column_index], "F9F5FF")
        prevent_split(table.rows[-1])

    after = DOC.add_paragraph()
    after.paragraph_format.space_after = Pt(4)
    after.paragraph_format.space_before = Pt(0)
    return table


STATUS_COLORS = {"Cumple": VERDE, "Parcial": AMBAR, "Por implementarse": ROJO, "No aplica": GRIS}


def add_matrix(title, rows):
    """Matriz de 21 requisitos: N.º · Bloque · Requisito · Obligatoriedad · Estado."""
    add_caption(title)
    table = DOC.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [0.9, 3.2, 8.2, 2.0, 2.4]
    for index, width in enumerate(widths):
        table.columns[index].width = Cm(width)
    set_table_borders(table)
    set_table_cell_margins(table, top=55, bottom=55, left=80, right=80)

    headers = ["N.º", "Bloque", "Requisito (resumen)", "Obligatoriedad", "Estado"]
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Cm(widths[index])
        set_cell_text(cell, header, bold=True, size=8.6, color=BLANCO)
        set_cell_bg(cell, "553171")
    repeat_header(table.rows[0]); prevent_split(table.rows[0])

    for row_index, (number, block, requirement, obligation, status) in enumerate(rows):
        cells = table.add_row().cells
        values = (number, block, requirement, obligation, status)
        for index, value in enumerate(values):
            cells[index].width = Cm(widths[index])
            if index == 4:
                key = "Parcial" if str(status).startswith("Parcial") else (
                    "No aplica" if str(status).startswith("No aplica") else (
                        "Por implementarse" if str(status).startswith("Por implementarse") else "Cumple"))
                set_cell_text(cells[index], value, size=8.3, bold=True, color=STATUS_COLORS[key])
            else:
                set_cell_text(cells[index], value, size=8.3, bold=(index == 0))
            if row_index % 2 == 1:
                set_cell_bg(cells[index], "F9F5FF")
        prevent_split(table.rows[-1])
    after = DOC.add_paragraph(); after.paragraph_format.space_after = Pt(4)


def H1(text):
    paragraph = DOC.add_heading(text, level=1)
    paragraph.paragraph_format.page_break_before = True
    return paragraph


def H2(text):
    return DOC.add_heading(text, level=2)


def H3(text):
    return DOC.add_heading(text, level=3)


def P(text, bold=False, italic=False, size=None, space_after=None, align=None):
    paragraph = DOC.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold; run.italic = italic
    if size is not None:
        run.font.size = Pt(size)
    if space_after is not None:
        paragraph.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        paragraph.alignment = align
    return paragraph


def BULLETS(items):
    for item in items:
        paragraph = DOC.add_paragraph(item, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.1


def NOTE(text):
    paragraph = DOC.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.4)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.italic = True; run.font.size = Pt(9.2); run.font.color.rgb = GRIS


def HORIZONTAL_RULE(paragraph, color="553171", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6"); bottom.set(qn("w:color"), color)
    borders.append(bottom); p_pr.append(borders)


# ═══════════════════════════════════════════════════════════════════════════
# Contenido
# ═══════════════════════════════════════════════════════════════════════════
def cover():
    section = DOC.sections[0]
    section.page_width = Cm(21.0); section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2); section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3); section.right_margin = Cm(2.3)
    section.header_distance = Cm(1.0); section.footer_distance = Cm(1.0)
    section.different_first_page_header_footer = True

    pg_num = OxmlElement("w:pgNumType"); pg_num.set(qn("w:start"), "0")
    section._sectPr.insert_element_before(
        pg_num, "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
        "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings", "w:sectPrChange",
    )

    normal = DOC.styles["Normal"]
    normal.font.name = "Calibri"; normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x2A, 0x1F, 0x3D)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal_rpr = normal.element.get_or_add_rPr()
    if normal_rpr.find(qn("w:lang")) is None:
        lang = OxmlElement("w:lang"); lang.set(qn("w:val"), "es-PE"); lang.set(qn("w:eastAsia"), "es-PE")
        normal_rpr.append(lang)

    for name, size, color, before, after in (
        ("Heading 1", 17, MORADO, 8, 10),
        ("Heading 2", 13, MORADO_OSCURO, 14, 6),
        ("Heading 3", 11.5, MORADO_MEDIO, 10, 4),
    ):
        style = DOC.styles[name]
        style.font.name = "Calibri"; style.font.size = Pt(size); style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True; style.paragraph_format.line_spacing = 1.05

    for name in ("List Bullet", "List Number"):
        style = DOC.styles[name]
        style.font.name = "Calibri"; style.font.size = Pt(10.5)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    hr = header.add_run("ULADECH Católica · Informe técnico y funcional del Gestor F1")
    hr.font.size = Pt(8.2); hr.font.color.rgb = GRIS_CLARO

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Documento de uso interno · Página ")
    fr.font.size = Pt(8.2); fr.font.color.rgb = GRIS_CLARO
    add_field(footer, "PAGE", "1")
    fr2 = footer.add_run(" de "); fr2.font.size = Pt(8.2); fr2.font.color.rgb = GRIS_CLARO
    add_field(footer, "NUMPAGES", "1")
    for run in footer.runs:
        run.font.size = Pt(8.2); run.font.color.rgb = GRIS_CLARO

    DOC.core_properties.title = "Informe Técnico y Funcional — Gestor F1 Mentora IdanIA"
    DOC.core_properties.subject = "Gestor F1 y chat de alumnos de ULADECH Católica"
    DOC.core_properties.author = "Equipo técnico del ecosistema EducaIA"
    DOC.core_properties.last_modified_by = "Equipo técnico del ecosistema EducaIA"
    DOC.core_properties.keywords = "ULADECH, EducaIA, gestor, IdanIA, RAG, Moodle, seguridad"
    DOC.core_properties.comments = "Documento técnico de uso interno. No contiene credenciales ni datos personales."

    p = DOC.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(2)
    run = p.add_run("ULADECH CATÓLICA"); run.bold = True; run.font.size = Pt(11); run.font.color.rgb = MORADO
    p = DOC.add_paragraph(); p.paragraph_format.space_after = Pt(30)
    run = p.add_run("Ecosistema EducaIA"); run.font.size = Pt(10.5); run.font.color.rgb = GRIS
    HORIZONTAL_RULE(p, color="E3D6F5", size="6")

    p = DOC.add_paragraph(); p.paragraph_format.space_before = Pt(24); p.paragraph_format.space_after = Pt(4)
    run = p.add_run("Informe técnico y funcional"); run.bold = True; run.font.size = Pt(27); run.font.color.rgb = MORADO_OSCURO
    p = DOC.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    run = p.add_run("Gestor F1 — Mentora IdanIA"); run.bold = True; run.font.size = Pt(21); run.font.color.rgb = MORADO
    p = DOC.add_paragraph(); p.paragraph_format.space_after = Pt(22)
    run = p.add_run("Roles, módulos, funcionalidades del asistente, seguridad e integración con Moodle")
    run.font.size = Pt(13); run.font.color.rgb = MORADO_MEDIO

    p = DOC.add_paragraph(); p.paragraph_format.space_after = Pt(30)
    run = p.add_run(
        "Arquitectura, modelo de datos, funcionalidades por rol, procesamiento de contenidos, "
        "seguridad, despliegue, calidad y trabajo pendiente del gestor institucional y del chat de alumnos."
    )
    run.font.size = Pt(11); run.font.color.rgb = GRIS; run.italic = True

    for label, value in [
        ("Documento", "Informe técnico y funcional"),
        ("Versión", "1.0"),
        ("Fecha", "Octubre de 2026"),
        ("Clasificación", "Uso interno"),
        ("Ámbito", "Gestor F1 (administrador y docente) y chat público de alumnos"),
        ("Elaboración", "Equipo técnico del ecosistema EducaIA"),
    ]:
        paragraph = DOC.add_paragraph(); paragraph.paragraph_format.space_after = Pt(1)
        lr = paragraph.add_run(f"{label}: "); lr.bold = True; lr.font.size = Pt(9.5); lr.font.color.rgb = MORADO_OSCURO
        vr = paragraph.add_run(value); vr.font.size = Pt(9.5); vr.font.color.rgb = GRIS

    DOC.add_paragraph()
    add_table(
        ["Versión", "Fecha", "Descripción"],
        [["1.0", "Octubre de 2026", "Versión inicial del informe."]],
        [2.0, 3.2, 11.0], caption="Control de versiones del documento.", font_size=9.0,
    )

    DOC.add_page_break()
    p = DOC.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    run = p.add_run("Índice"); run.bold = True; run.font.size = Pt(15); run.font.color.rgb = MORADO
    add_toc()
    set_update_fields()
    DOC.add_page_break()


def sec_resumen():
    H1("1. Resumen ejecutivo")
    P(
        "El Gestor F1 es la aplicación de ULADECH Católica que administra la base de conocimiento "
        "de los asistentes de IA (mentora IdanIA) y su chat público. Se ejecuta como una aplicación "
        "FastAPI montada bajo el prefijo /f1, detrás de un proxy inverso con HTTPS, y se despliega con "
        "contenedores sobre una base PostgreSQL y una base vectorial Qdrant."
    )
    P(
        "La aplicación cubre tres roles. El administrador configura todo el sistema: colecciones, "
        "aulas RD, asistentes, estrategias del pipeline RAG, claves de proveedores, usuarios y un visor "
        "de lo que el sistema sabe de cada curso. El docente entra desde Moodle y ve únicamente sus "
        "cursos; puede revisar el material vectorizado, marcar documentos que no desea usar, activar o "
        "desactivar funcionalidades, ajustar el número de mensajes, crear su asistente y probarlo. El "
        "alumno usa el chat público del asistente, que responde con el material del curso, datos de "
        "navegación del aula y sus propias calificaciones."
    )
    P(
        "El asistente apoya su respuesta en tres funcionalidades: Documentos (RAG híbrido sobre el "
        "material vectorizado), Navegabilidad del aula (semanas, actividades, fechas y ubicación de "
        "recursos) y Notas (calificaciones y avance del propio alumno). Cada una puede desactivarse por "
        "asistente y por aula; si se apaga, el asistente devuelve un mensaje fijo y no recurre al RAG."
    )
    P(
        "En materia de integración segura con Moodle, el sistema ya cuenta con controles básicos "
        "(lista blanca de rutas por rol, verificación de curso contra el ERP, contraseñas con PBKDF2, "
        "cookie firmada, secretos en el entorno y descargas con token de servicio). La arquitectura "
        "objetivo con token cifrado AES-256-GCM y sesión propia se documenta en el apartado 9: en la "
        "matriz del acceso del docente hay 1 requisito que cumple, 6 parciales y 14 por implementarse; "
        "en la del chat de alumnos hay 2 parciales, 1 no aplica y 18 por implementarse."
    )
    NOTE(
        "Convención del documento: «Por implementarse» indica una capacidad descrita en el modelo de "
        "seguridad objetivo que todavía no está en el código. No se presenta como hecha."
    )


def sec_alcance():
    H1("2. Alcance y objetivos")
    P(
        "El informe describe todas las funcionalidades del gestor y del chat de los alumnos, "
        "organizadas por rol (administrador, docente y alumno), y las características técnicas de la "
        "integración con Moodle. Es un documento de redacción: no modifica el código de la aplicación "
        "ni el servidor, y no incluye credenciales, tokens, direcciones internas ni datos personales."
    )
    H2("2.1 Objetivos")
    BULLETS([
        "Explicar la arquitectura técnica, las fuentes de datos y los flujos principales (pregunta del alumno y vectorización).",
        "Describir los módulos del gestor, para quién son y qué reglas aplican.",
        "Detallar las funcionalidades del asistente y el comportamiento del chat.",
        "Documentar el procesamiento de contenidos (ETL) y los estados del material.",
        "Presentar la seguridad actual, la arquitectura objetivo de integración con Moodle y dos matrices de cumplimiento de 21 requisitos.",
        "Resumir el despliegue, la operación, las pruebas y el trabajo pendiente.",
    ])
    H2("2.2 Fuera de alcance")
    BULLETS([
        "Cambios de código o de configuración en producción.",
        "Credenciales, claves, tokens reales, direcciones internas y datos personales.",
        "El detalle de los proyectos F2 y F3, que solo se mencionan como aplicaciones sujetas a la misma matriz de seguridad.",
    ])


def sec_arquitectura():
    H1("3. Arquitectura técnica")
    P(
        "La aplicación es un servicio FastAPI con plantillas Jinja. Toda la interfaz se sirve desde el "
        "mismo proceso, que también expone la API interna del gestor y las rutas públicas del chat. La "
        "persistencia se reparte entre PostgreSQL (configuración y datos de negocio), Qdrant (vectores) "
        "y servicios externos (Moodle y el ERP)."
    )
    H2("3.1 Componentes")
    add_table(
        ["Componente", "Función"],
        [
            ["Aplicación FastAPI + Jinja", "Sirve el gestor (/f1/gestor), el login, la API interna y el chat público (/f1/asistente/…). Se ejecuta bajo el prefijo /f1."],
            ["Proxy inverso con HTTPS", "Termina TLS y publica la aplicación en el subdominio institucional; la aplicación no configura TLS por sí misma."],
            ["PostgreSQL", "Usuarios y roles, RDs, colecciones, asistentes, sesiones y mensajes, cola de archivos (ledger), caché del Course Map, campos estructurados, fuente de autoridad, configuración por curso, exclusiones y configuración general."],
            ["Qdrant", "Base vectorial híbrida: un vector denso y uno disperso (BM25) por fragmento, con fusión y filtros por metadatos."],
            ["API de servicios web de Moodle", "Cursos, contenidos, usuarios, calificaciones, calendario y descarga de archivos, siempre con token de servicio. No se inicia sesión en Moodle con usuario y contraseña."],
            ["ERP MySQL (solo lectura)", "Vista de matrícula del periodo vigente: relación entre el documento del docente y sus cursos."],
            ["Proveedores de modelos", "Embeddings, generación, reranking y transcripción, configurables por etapa; pueden ser locales o de API."],
            ["Docker Compose", "Empaqueta la aplicación; la caché de modelos persiste en un volumen."],
        ],
        [3.6, 12.6], caption="Componentes de la arquitectura del Gestor F1.",
    )

    H2("3.2 Flujo de una pregunta del alumno")
    add_table(
        ["Paso", "Qué ocurre"],
        [
            ["1. Entrada", "El alumno abre /f1/asistente/<token> con el curso y su identidad (o documento y nombre). La página crea o reutiliza una sesión."],
            ["2. Mensaje", "El navegador envía la pregunta al endpoint público del asistente."],
            ["3. Configuración", "Se resuelve la colección del curso y la configuración por curso del asistente (funcionalidades, máximo de mensajes, activo y documentos excluidos)."],
            ["4. Query Router", "Reglas deterministas clasifican la pregunta. Navegación, fechas, semana actual, notas y avance se responden sin RAG, con datos reales de Moodle."],
            ["5. Recuperación", "Para Documentos se embebe la consulta (denso y disperso) y Qdrant recupera candidatos, excluyendo los documentos marcados como «No considerar»."],
            ["6. Evidencia", "Se aplican reranking, refuerzos por tipo de fuente y semana, y el índice de campos estructurados extraídos de tablas."],
            ["7. Generación", "El modelo redacta la respuesta con el prompt del asistente, la memoria de la conversación y una compuerta de evidencia que permite abstenerse."],
            ["8. Registro", "La respuesta, sus fuentes y la configuración usada se guardan en la sesión; el alumno puede calificarla."],
        ],
        [2.6, 13.6], caption="Flujo de una pregunta del alumno hasta la respuesta.",
    )

    H2("3.3 Flujo de vectorización")
    add_table(
        ["Paso", "Qué ocurre"],
        [
            ["1. Origen", "El administrador elige un curso de Moodle (Automático), sube archivos (Semiautomático) o reemplaza un documento (Actualización)."],
            ["2. Listado", "Se consultan los recursos del curso vía servicios web y se agrupan por semana o sección."],
            ["3. Validación", "Se verifica que el archivo no esté vacío, que su tamaño y tipo estén permitidos y que la firma binaria corresponda."],
            ["4. Duplicados", "Se detecta si el documento ya existe en la colección y se decide reemplazarlo."],
            ["5. ETL", "Se extrae el texto (PDF, Word, PowerPoint, Excel, imágenes con OCR, audio/video con transcripción, enlaces) y se fragmenta."],
            ["6. Embeddings", "Se generan los vectores densos y dispersos con las estrategias configuradas para la colección."],
            ["7. Qdrant", "Se escriben los fragmentos con metadatos (archivo, semana, sección, tipo de fuente, páginas)."],
            ["8. Ledger y campos", "Se actualiza la cola de archivos del curso y se re-sincroniza el índice de campos estructurados."],
        ],
        [2.6, 13.6], caption="Flujo de vectorización de contenidos.",
    )
    P(
        "El módulo Automático incorpora un selector de modelo de embeddings. Si se deja en blanco, usa el "
        "valor por defecto del sistema; si se elige un modelo, el backend vectoriza con ese modelo. Esto "
        "evita el error de dimensiones cuando la colección destino no usa el embedding por defecto."
    )


def sec_datos():
    H1("4. Modelo de datos y fuentes")
    H2("4.1 Tablas principales")
    add_table(
        ["Tabla", "Para qué sirve"],
        [
            ["usuarios", "Cuentas del gestor: administradores y docentes, con rol, documento, identificador de Moodle y nombre. Los docentes no tienen contraseña local."],
            ["rds", "Aulas RD (Redes de Datos) y su embedding maestro."],
            ["colecciones", "Curso de Moodle asociado a una RD o colección normal, colección de Qdrant, esquema, modo e intervalo de migración y si el asistente está activo para ese curso."],
            ["asistentes", "Asistente público: nombre, curso o RD, prompt, mensajes, token, estado, estrategias del pipeline y funcionalidades."],
            ["asistente_sesiones", "Una conversación por asistente, curso y usuario; incluye marca de prueba y quién la probó."],
            ["asistente_mensajes", "Historial de preguntas y respuestas, fuentes, configuración usada, calificación y comentario."],
            ["asistente_carpetas", "Agrupación de conversaciones de un alumno."],
            ["asistente_curso_config", "Ajustes del docente por aula: funcionalidades, máximo de mensajes y activo, sin tocar la configuración global de la RD."],
            ["asistente_documentos_excluidos", "Documentos que el docente decidió no usar, por asistente y por curso."],
            ["curso_archivos", "Cola de indexación: estado de cada archivo del curso (nuevo, posible actualización, vectorizado, eliminado, no soportado)."],
            ["course_map_cache", "Estructura del curso guardada desde Moodle, con fecha de actualización e invalidación."],
            ["coleccion_campos_estructurados", "Campos extraídos de las tablas de los documentos, normalizados y con su documento de origen."],
            ["fuente_autoridad", "Prioridad entre tipos de fuente cuando dos documentos se contradicen."],
            ["configuracion", "Configuración mutable del gestor y credenciales de proveedores, marcadas como secretas."],
        ],
        [4.2, 12.0], caption="Tablas principales de PostgreSQL.",
    )

    H2("4.2 Origen de los datos")
    add_table(
        ["Fuente", "Qué aporta"],
        [
            ["Moodle (servicios web)", "Cursos, secciones, actividades, archivos, usuarios, calificaciones, progreso y fechas de vencimiento."],
            ["ERP MySQL (solo lectura)", "Matrícula del periodo vigente: documento del docente y shortname de sus cursos."],
            ["Qdrant", "Fragmentos vectoriales con su texto y metadatos; se usa para recuperar y para listar documentos."],
            ["PostgreSQL", "Toda la configuración, los roles, los asistentes, las sesiones, el ledger y los índices estructurados."],
            ["Modelos y proveedores", "Embeddings, reranking, generación y transcripción, según la estrategia activa de cada etapa."],
        ],
        [5.0, 11.2], caption="Fuentes de datos y su aporte.",
    )


def sec_roles():
    H1("5. Roles y accesos")
    H2("5.1 Administrador")
    P(
        "Accede con usuario y contraseña. La contraseña se guarda con PBKDF2-SHA256 y sal propia. Ve todo "
        "el menú y puede configurar colecciones, RDs, asistentes, estrategias del pipeline, usuarios y "
        "claves de proveedores. También dispone del visor de conocimiento por curso."
    )
    H2("5.2 Docente")
    P(
        "Entra desde una actividad URL de Moodle con su documento, un token compartido y su identificador "
        "de Moodle. El identificador se valida en vivo contra Moodle y debe corresponder al documento. "
        "Los cursos se obtienen del ERP del periodo vigente. Una lista blanca de rutas limita su acceso a "
        "su panel y a sus propios cursos; cada operación por curso vuelve a verificar la pertenencia."
    )
    H2("5.3 Alumno")
    P(
        "Usa el chat público del asistente. El token de la ruta identifica al asistente, no al alumno. La "
        "identidad del alumno viaja como identificador en la URL o como documento y nombre escritos por "
        "el visitante; el nombre se consulta en Moodle."
    )
    H2("5.4 Modo prueba")
    P(
        "El docente o el administrador puede abrir el chat en modo prueba desde el gestor. La sesión se "
        "marca como prueba y se excluye de las estadísticas; usa la misma configuración del curso, de modo "
        "que reproduce exactamente lo que verá el alumno."
    )
    add_table(
        ["Módulo / capacidad", "Administrador", "Docente", "Alumno"],
        [
            ["Mis cursos y detalle del curso", "Sí (visor admin)", "Sí (solo sus cursos)", "No"],
            ["Conocimiento idanIA y exclusiones", "Sí", "Sí (solo sus cursos)", "No"],
            ["Funcionalidades y mensajes por aula", "Sí", "Sí (solo sus cursos)", "No"],
            ["Probar chat", "Sí", "Sí", "No"],
            ["Vectorización (Semiautomático, Automático, Actualización)", "Sí", "No", "No"],
            ["Gestión RD y Course Map", "Sí", "No", "No"],
            ["Asistentes (lista, edición, creación)", "Sí", "Solo el de su curso", "No"],
            ["Gestión Qdrant", "Sí", "No", "No"],
            ["Usuarios, conocimiento del curso y claves", "Sí", "No", "No"],
            ["Chat público del asistente", "Sí (modo prueba)", "Sí (modo prueba)", "Sí"],
        ],
        [6.6, 3.2, 3.4, 3.0], caption="Acceso por rol y módulo.",
    )


def sec_modulos():
    H1("6. Módulos funcionales del gestor")
    P(
        "El menú del gestor se organiza en cuatro secciones para el administrador y una vista propia para "
        "el docente. Los módulos que se describen a continuación corresponden al código actual."
    )

    H2("6.1 Mis cursos (docente)")
    P("El panel del docente sigue el patrón lista y detalle. La lista muestra tarjetas compactas por curso, con el estado del asistente y el avance del material. Al abrir un curso se accede a cuatro pestañas:")
    BULLETS([
        "Resumen: estado del asistente, colección y vectores, funcionalidades activas y las últimas conversaciones de prueba.",
        "Conocimiento idanIA: todos los materiales del curso tomados de Moodle, agrupados por semana con el mismo criterio del módulo Automático. Cada archivo muestra su estado (vectorizado, pendiente, no soportado o excluido), la fecha de subida, los fragmentos y la opción «No considerar». Permite buscar, filtrar por estado y por tipo, y vectorizar los pendientes con una barra de progreso. Los cambios de exclusión se confirman con una barra de «cambios sin guardar».",
        "Funcionalidades: tarjetas de Documentos, Navegabilidad del aula y Notas, con su interruptor, una descripción, ejemplos de preguntas y una vista previa del contenido real. Incluye el máximo de mensajes por conversación y el estado activo. En asistentes de una RD, los cambios se guardan por aula y no modifican la configuración global de la RD.",
        "Probar chat: incrusta el chat en modo prueba con la configuración de esa aula.",
    ])
    NOTE(
        "La vista por funcionalidad, el wizard Curso → Identidad → Funcionalidades → Avanzado → Revisar y "
        "la Navegabilidad en línea de tiempo están implementados en el código actual del gestor."
    )

    H2("6.2 Vectorización")
    H3("Semiautomático")
    P(
        "Carga manual de documentos multiformato. Se elige la colección destino, se suben los archivos, se "
        "valida cada uno y se confirma la vectorización con el modelo elegido. Es el flujo indicado para "
        "material que no está en Moodle."
    )
    H3("Automático")
    P(
        "Vectoriza un curso de Moodle a partir de su identificador. Consulta los recursos, los muestra en "
        "un árbol por semana con filtros por tipo, marca duplicados y permite seleccionar archivos. Incluye "
        "un selector de modelo de embeddings que debe coincidir con el de la colección destino. Cada archivo "
        "se puede clasificar como material del curso o bibliografía."
    )
    H3("Actualización")
    P(
        "Reemplaza un documento ya vectorizado por una versión nueva, desde Moodle o desde un archivo local. "
        "Elimina la versión anterior en Qdrant y vectoriza la nueva con el mismo modelo de la colección."
    )

    H2("6.3 Asistentes")
    BULLETS([
        "Búsqueda semántica: consulta por similitud sobre una colección y muestra los fragmentos recuperados.",
        "Gestión RD: crea y edita aulas RD, su curso de referencia y sus cursos y colecciones hijas; permite activar o desactivar el asistente por curso.",
        "Course Map: muestra la estructura real del aula (secciones, actividades, visibilidad y datos generales) tomada de Moodle.",
        "Asistentes: lista los asistentes con sus funcionalidades, estado, URL pública y calificación; permite editar y eliminar, y ofrece la vista «Ver por funcionalidad» con interruptores rápidos.",
        "Crear asistente: wizard guiado en el orden Curso → Identidad → Funcionalidades → Avanzado (opcional) → Revisar. Precarga la plantilla Idania, usa los embeddings por defecto para asistentes y, al terminar, ofrece la prueba de chat y el acceso a la edición.",
        "Prueba de asistente: ejecuta una consulta contra el pipeline completo con la configuración activa.",
    ])

    H2("6.4 Gestión Qdrant")
    BULLETS([
        "Dashboard: resumen de colecciones y vectores.",
        "Colecciones: listado con esquema, puntos, estado y migración; incluye la vista por semana, el detalle de fragmentos y la recategorización de un documento ya vectorizado.",
        "Crear colección: colección normal, o hija o padre de una RD; esquema híbrido o legado; embedding denso y disperso; métrica de distancia.",
        "Verificar existencia: comprueba si una colección existe y su esquema.",
        "API Reference: referencia de los endpoints de administración de colecciones.",
    ])

    H2("6.5 Administración")
    BULLETS([
        "Usuarios: creación y eliminación de cuentas con rol administrador o docente; un docente se crea con su documento, sin contraseña, y su acceso es por el enlace de Moodle. La pantalla incluye la ayuda para pegar el enlace como actividad URL.",
        "Conocimiento del curso: visor de solo lectura de lo que el sistema guardó y usa para responder: estructura guardada, datos generales del aula, campos de tablas, fechas de entrega, archivos y cobertura, y fuente de autoridad. La pestaña de estructura usa la misma representación de navegación del panel docente.",
        "Claves de API: alta, prueba y estado de las claves de los proveedores.",
    ])


def sec_funcionalidades():
    H1("7. Funcionalidades del asistente")
    P(
        "El asistente decide en cada pregunta si usa una respuesta determinista (navegación, fechas, notas) "
        "o el RAG sobre los documentos. Las tres funcionalidades pueden activarse o desactivarse por "
        "asistente y por aula."
    )
    H2("7.1 Documentos")
    P(
        "Responde con el material del curso mediante recuperación híbrida: embedding denso y BM25 en "
        "Qdrant, fusión, reranking, refuerzo por tipo de fuente y por semana, y evidencia de campos "
        "estructurados extraídos de tablas. Aplica una compuerta de evidencia (se abstiene si no hay "
        "fuente) y una validación de fundamentación. Respeta los documentos excluidos con un filtro que "
        "los deja fuera de la búsqueda."
    )
    P("Ejemplos: «¿Qué es…?», «Explícame el tema de la semana 3».")
    H2("7.2 Navegabilidad del aula")
    P(
        "Responde sin RAG con el Course Map y el calendario de Moodle: semana actual, actividades de una "
        "semana, ubicación y enlace de un recurso, fechas de entrega y horario. Si el calendario no está "
        "disponible, lo informa y no inventa fechas."
    )
    P("Ejemplos: «¿Qué tengo esta semana?», «¿Dónde está la tarea 2?», «¿Cuándo vence…?».")
    H2("7.3 Notas")
    P(
        "Consulta en vivo Moodle para el alumno de la sesión: calificaciones y avance. Cada alumno ve solo "
        "sus propias notas; el asistente no consulta a otros estudiantes."
    )
    P("Ejemplos: «¿Cuál es mi nota?», «¿Qué me falta entregar?».")
    H2("7.4 Query Router: intents")
    add_table(
        ["Intent", "Qué responde", "Funcionalidad"],
        [
            ["STRUCTURED", "Conteos y rangos de unidades o semanas desde el índice estructurado.", "Documentos"],
            ["EVALUATION_CRITERIA", "Criterios de evaluación tomados del material.", "Documentos"],
            ["TEACHER_INFO", "Docente y datos de contacto del aula.", "Documentos"],
            ["NAVIGATION", "Ubicación y enlace de una actividad o recurso.", "Navegabilidad"],
            ["CURRENT_WEEK", "Semana actual del curso y su tema.", "Navegabilidad"],
            ["DUE_DATES", "Fechas de entrega reales de Moodle.", "Navegabilidad"],
            ["SCHEDULE", "Horario del curso.", "Navegabilidad"],
            ["COURSE_DATE", "Fechas de inicio y fin del curso.", "Navegabilidad"],
            ["STUDY_HELP", "Orientación de estudio con la estructura del aula.", "Navegabilidad"],
            ["GRADES", "Calificaciones del alumno.", "Notas"],
            ["PROGRESS", "Avance del alumno en el curso.", "Notas"],
        ],
        [4.0, 8.6, 3.6], caption="Intents del Query Router y funcionalidad asociada.",
    )
    H2("7.5 Funcionalidad desactivada")
    P(
        "Si el intent pertenece a una funcionalidad apagada, el asistente devuelve un mensaje fijo («Esta "
        "funcionalidad no está habilitada en este asistente») y no cae al RAG, para no inventar notas o "
        "fechas a partir de documentos. Si Documentos está apagado y el router no responde, devuelve un "
        "mensaje indicando que el asistente no está configurado para responder sobre los documentos."
    )
    H2("7.6 Elementos del chat")
    BULLETS([
        "Saludo y reencuentro personalizados con la semana actual del curso.",
        "Historial de conversaciones, carpetas y renombrado.",
        "Calificación «Me ayudó / No me ayudó» con comentario opcional.",
        "Tope de preguntas por conversación, configurable por asistente y por aula.",
        "Modo prueba con aviso visible y exclusión de las estadísticas.",
    ])


def sec_etl():
    H1("8. Procesamiento de contenidos (ETL)")
    H2("8.1 Formatos admitidos")
    add_table(
        ["Tipo", "Formatos", "Procesamiento"],
        [
            ["Documentos", "PDF", "Texto y tablas; OCR para páginas sin texto."],
            ["Documentos", "Word (DOCX), PowerPoint (PPTX y PPT heredado)", "Texto, tablas e imágenes."],
            ["Documentos", "Excel (XLSX)", "Hojas y tablas."],
            ["Imágenes", "JPG, JPEG, PNG, WEBP", "OCR con detección de idioma."],
            ["Audio", "MP3, WAV, M4A, OGG", "Transcripción automática."],
            ["Video", "MP4, MOV, WEBM, MKV", "Extracción de audio y transcripción; subtítulos de YouTube cuando existen."],
            ["Web", "Enlaces y páginas", "Extracción del texto principal."],
        ],
        [3.0, 7.0, 6.2], caption="Formatos admitidos y procesamiento.",
    )
    H2("8.2 Validaciones y estados")
    BULLETS([
        "Validación de archivo vacío, tamaño máximo, extensión permitida y firma binaria según el tipo.",
        "Detección de duplicados por nombre de archivo en la colección; si existe, se reemplaza o se bloquea según el flujo.",
        "Clasificación por tipo de fuente: material del curso o bibliografía.",
        "Metadatos de semana y sección tomados de la estructura de Moodle.",
        "Cola de archivos con estados: nuevo, posible actualización, sin cambios, vectorizado, eliminado y no soportado.",
        "Re-sincronización del índice de campos estructurados después de vectorizar.",
        "Limpieza de archivos temporales y barrido periódico de huérfanos.",
    ])


def sec_seguridad():
    H1("9. Seguridad e integración con Moodle")
    H2("9.1 Controles existentes")
    BULLETS([
        "Lista blanca de rutas por rol en el middleware de acceso.",
        "Verificación de que el curso pertenece al docente contra el ERP, en cada operación por curso.",
        "Contraseñas de administrador con PBKDF2-SHA256 y sal propia.",
        "Cookie de sesión firmada, HttpOnly y SameSite=Lax.",
        "Secretos en variables de entorno, fuera del repositorio.",
        "Descargas de Moodle con token de servicio, sin credenciales de usuario.",
        "Validación en vivo del identificador del docente contra Moodle.",
        "Exclusión de las sesiones de prueba en las estadísticas de calificaciones.",
        "Mensajes de error genéricos para el usuario.",
    ])

    H2("9.2 Arquitectura objetivo: token seguro de Moodle")
    P(
        "El modelo de referencia es la integración con la actividad «URL segura» de Moodle. Moodle emite "
        "un token cifrado con AES-256-GCM que la aplicación descifra con la librería oficial, valida y "
        "canjea por una sesión propia. Lo que sigue se marca como «Por implementarse» salvo lo indicado "
        "en las matrices."
    )
    add_table(
        ["Paso", "Responsable", "Acción"],
        [
            ["1", "Moodle", "Emite el token cifrado con identificador de aplicación, versión de clave, identificador de usuario, vencimiento corto, identificador único y parámetros permitidos."],
            ["2", "Aplicación", "Recibe el token por el parámetro configurado y lo descifra con AES-256-GCM."],
            ["3", "Aplicación", "Valida formato, aplicación, versión de clave, firma y vencimiento con margen de tolerancia."],
            ["4", "Aplicación", "Registra el identificador único y rechaza su reutilización durante la vigencia."],
            ["5", "Aplicación", "Identifica al usuario solo por el identificador firmado y lo busca en sus datos propios."],
            ["6", "Aplicación", "Crea una sesión nueva, redirige sin el token y responde con caché deshabilitada."],
            ["7", "Aplicación", "Aplica cabeceras de seguridad, protección de origen y límites de peticiones."],
            ["8", "Ambas", "Audita los intentos fallidos y permite la rotación de claves."],
        ],
        [1.6, 3.0, 11.6], caption="Flujo objetivo de la integración con token seguro.",
    )
    P(
        "El proyecto de referencia implementa este flujo con una librería de verificación, una sesión "
        "nueva por login, anti-reuso del identificador único en un almacén propio, auditoría, cabeceras "
        "CSP, HSTS y X-Frame-Options, protección CSRF por cabecera y origen, límite de peticiones en el "
        "proxy y contenedor sin privilegios. Todo ello se considera «Por implementarse» en el Gestor F1."
    )
    P(
        "Ámbito de aplicación: (a) el acceso del rol docente al gestor, que reemplaza el token compartido y "
        "el par documento e identificador; (b) las aplicaciones integradas a Moodle, en particular el chat "
        "de alumnos, que reemplaza el identificador en la URL y el acceso por documento y nombre. Las "
        "aplicaciones F2 y F3 deben cumplir la misma matriz."
    )

    H2("9.3 Matrices de cumplimiento")
    P(
        "Se presentan dos matrices con los 21 requisitos del documento de características técnicas: una "
        "para el acceso del docente al gestor y otra para las aplicaciones, en particular el chat de "
        "alumnos. Los estados posibles son Cumple, Parcial, Por implementarse y No aplica."
    )

    matrix_a = [
        ("1", "Autenticación", "Endpoint HTTPS que acepta el token", "Obligatorio", "Parcial (HTTPS sí; el token actual es un secreto compartido, no un token cifrado)"),
        ("2", "Autenticación", "Descifrado AES-256-GCM con librería oficial", "Obligatorio", "Por implementarse"),
        ("3", "Autenticación", "Identificador de aplicación y claves en variables de entorno", "Obligatorio", "Por implementarse"),
        ("4", "Autenticación", "Rotación de claves (aceptar la anterior)", "Obligatorio", "Por implementarse"),
        ("5", "Autenticación", "Validar formato, aplicación, versión, firma y vencimiento", "Obligatorio", "Por implementarse"),
        ("6", "Autenticación", "Mensaje genérico ante cualquier fallo", "Obligatorio", "Por implementarse"),
        ("7", "Identidad", "Identificar solo por el identificador firmado", "Obligatorio", "Por implementarse (hoy: documento e identificador validados en Moodle)"),
        ("8", "Identidad", "No confiar en los parámetros del token", "Obligatorio", "Por implementarse"),
        ("9", "Identidad", "Buscar al usuario por identificador en datos propios", "Obligatorio", "Parcial (cursos desde el ERP por documento; por implementarse con el identificador firmado)"),
        ("10", "Identidad", "Listas de autorización separadas por nivel", "Obligatorio (si aplica)", "Parcial (roles admin y docente; lista por identificador de Moodle por implementarse)"),
        ("11", "Sesión", "Cookie firmada HttpOnly, Secure y SameSite", "Obligatorio", "Parcial (firmada, HttpOnly y SameSite; Secure por implementarse)"),
        ("12", "Sesión", "Redirección a la URL sin el token", "Obligatorio", "Cumple (redirige al panel sin parámetros)"),
        ("13", "Sesión", "Nuevo identificador de sesión en cada login", "Obligatorio", "Parcial (sí en el login por enlace; no en el login con contraseña)"),
        ("14", "Sesión", "Anti-reuso del identificador único", "Obligatorio", "Por implementarse"),
        ("15", "Sesión", "Caché deshabilitada en la respuesta con token", "Obligatorio", "Por implementarse"),
        ("16", "Sin puerta trasera", "No habilitar accesos más débiles", "Obligatorio", "Por implementarse (retirar el token compartido)"),
        ("17", "Sin puerta trasera", "Alternativas igual de fuertes o desactivadas", "Obligatorio", "Parcial (administradores con contraseña PBKDF2)"),
        ("18", "Moodle", "URL base fija al registrar la aplicación", "Obligatorio", "Por implementarse"),
        ("19", "Moodle", "Parámetros adicionales bloqueados", "Obligatorio", "Por implementarse"),
        ("20", "Pruebas", "Tokens de prueba locales", "Recomendado", "Por implementarse"),
        ("21", "Pruebas", "Casos válido, alterado, caducado, reusado y otra aplicación", "Recomendado", "Por implementarse"),
    ]
    add_matrix("Matriz A. Acceso del docente al gestor.", matrix_a)
    P(
        "Resumen de la Matriz A: 1 requisito cumple, 6 son parciales y 14 están por implementarse. "
        "El único requisito que cumple es la redirección sin el token; los parciales corresponden a "
        "controles básicos ya presentes, como HTTPS, la cookie firmada y la validación del docente."
    )

    matrix_b = [
        ("1", "Autenticación", "Endpoint HTTPS que acepta el token", "Obligatorio", "Parcial (HTTPS sí; el token de la ruta identifica al asistente, no al alumno)"),
        ("2", "Autenticación", "Descifrado AES-256-GCM con librería oficial", "Obligatorio", "Por implementarse"),
        ("3", "Autenticación", "Identificador de aplicación y claves en variables de entorno", "Obligatorio", "Por implementarse"),
        ("4", "Autenticación", "Rotación de claves (aceptar la anterior)", "Obligatorio", "Por implementarse"),
        ("5", "Autenticación", "Validar formato, aplicación, versión, firma y vencimiento", "Obligatorio", "Por implementarse"),
        ("6", "Autenticación", "Mensaje genérico ante cualquier fallo", "Obligatorio", "Por implementarse"),
        ("7", "Identidad", "Identificar solo por el identificador firmado", "Obligatorio", "Por implementarse (hoy: identificador en la URL o documento y nombre)"),
        ("8", "Identidad", "No confiar en los parámetros del token", "Obligatorio", "Por implementarse"),
        ("9", "Identidad", "Buscar al usuario por identificador en datos propios", "Obligatorio", "Parcial (nombre desde Moodle por identificador no autenticado)"),
        ("10", "Identidad", "Listas de autorización separadas por nivel", "Obligatorio (si aplica)", "No aplica (un solo perfil: alumno)"),
        ("11", "Sesión", "Cookie firmada HttpOnly, Secure y SameSite", "Obligatorio", "Por implementarse (sin sesión propia del alumno)"),
        ("12", "Sesión", "Redirección a la URL sin el token", "Obligatorio", "Por implementarse"),
        ("13", "Sesión", "Nuevo identificador de sesión en cada login", "Obligatorio", "Por implementarse"),
        ("14", "Sesión", "Anti-reuso del identificador único", "Obligatorio", "Por implementarse"),
        ("15", "Sesión", "Caché deshabilitada en la respuesta con token", "Obligatorio", "Por implementarse"),
        ("16", "Sin puerta trasera", "No habilitar accesos más débiles", "Obligatorio", "Por implementarse (retirar identificador en la URL y documento y nombre)"),
        ("17", "Sin puerta trasera", "Alternativas igual de fuertes o desactivadas", "Obligatorio", "Por implementarse"),
        ("18", "Moodle", "URL base fija al registrar la aplicación", "Obligatorio", "Por implementarse"),
        ("19", "Moodle", "Parámetros adicionales bloqueados", "Obligatorio", "Por implementarse"),
        ("20", "Pruebas", "Tokens de prueba locales", "Recomendado", "Por implementarse"),
        ("21", "Pruebas", "Casos válido, alterado, caducado, reusado y otra aplicación", "Recomendado", "Por implementarse"),
    ]
    add_matrix("Matriz B. Aplicaciones: chat de alumnos.", matrix_b)
    P(
        "Resumen de la Matriz B: ningún requisito cumple, 2 son parciales, 1 no aplica y 18 están por "
        "implementarse. El chat actual no tiene sesión propia del alumno y acepta el identificador de la "
        "URL, por lo que la mayor parte de la matriz queda pendiente."
    )


def sec_despliegue():
    H1("10. Despliegue, operación y calidad")
    H2("10.1 Despliegue")
    BULLETS([
        "La aplicación se ejecuta en un contenedor definido con Docker Compose y se publica a través del proxy inverso con HTTPS.",
        "Los scripts de despliegue suben el código y reconstruyen el contenedor. El archivo .dockerignore excluye material temporal, credenciales, resultados de pruebas y artefactos pesados del contexto de construcción.",
        "La caché de modelos persiste en un volumen para no volver a descargarlos en cada despliegue.",
        "Las migraciones SQL se aplican a mano; no hay herramienta automática de migraciones.",
        "El registro de la aplicación informa el estado de la base, el modelo por defecto y las tareas de mantenimiento.",
    ])
    H2("10.2 Calidad y pruebas")
    BULLETS([
        "Pruebas unitarias con pytest para roles, configuración por curso, funcionalidades y sesiones de prueba.",
        "Pruebas de extremo a extremo con Playwright para login, colecciones, asistentes, chat y panel del docente.",
        "Evaluaciones de calidad del asistente sobre conjuntos de 200 preguntas por curso, con métricas de acierto, alucinación y fundamentación.",
    ])
    add_table(
        ["Conjunto de evaluación", "Preguntas", "Acierto", "Alucinación", "Fundamentación"],
        [
            ["Batch 1 (alucinaciones)", "200", "80,0 %", "9,5 %", "8,48 / 10"],
            ["Batch 2 (ejemplos nuevos)", "200", "72,5 %", "14,5 %", "8,32 / 10"],
        ],
        [5.2, 2.4, 2.6, 2.8, 3.2],
        caption="Resumen de las evaluaciones de calidad (sin datos personales).",
    )
    NOTE(
        "Los porcentajes son un resumen agregado de los archivos de evaluación del repositorio; no incluyen "
        "preguntas, respuestas ni identificadores de personas."
    )


def sec_limitaciones():
    H1("11. Limitaciones y trabajo pendiente")
    P("Se listan las capacidades que todavía no están implementadas o que requieren una decisión de producto. Todas se marcan como «Por implementarse».")
    BULLETS([
        "Integración con token cifrado de Moodle en el acceso del docente: descifrado AES-256-GCM, rotación de claves, anti-reuso y redirección sin token.",
        "Sesión propia del alumno en el chat público: hoy el identificador viaja en la URL y el acceso por documento y nombre no se verifica.",
        "Cabeceras de seguridad desde la aplicación (CSP, HSTS y X-Frame-Options), caché deshabilitada en la respuesta con token y protección CSRF.",
        "Cookie de sesión con atributo Secure desde la aplicación; hoy el HTTPS lo aporta el proxy.",
        "Regeneración de la sesión en el login de administrador con contraseña.",
        "Listas de autorización por identificador de Moodle separadas por nivel, además de los roles actuales.",
        "Tokens de prueba locales y pruebas automatizadas de token válido, alterado, caducado, reusado y de otra aplicación.",
        "Bloqueo de parámetros adicionales y URL base fija en el registro de la aplicación dentro de Moodle.",
        "Eliminación gradual de los métodos de acceso más débiles una vez adoptado el token cifrado.",
    ])
    P(
        "Además, el informe no ejecuta una auditoría de seguridad ofensiva ni sustituye una revisión de "
        "código; las matrices reflejan lo observado en el código y en la configuración descrita."
    )


def sec_anexos():
    H1("12. Anexos")
    H2("12.1 Glosario")
    add_table(
        ["Término", "Significado"],
        [
            ["RAG", "Generación aumentada por recuperación: se buscan fragmentos relevantes y el modelo responde con ellos."],
            ["RD", "Red de Datos o aula de referencia que agrupa cursos y fija un embedding maestro."],
            ["Chunk", "Fragmento de un documento que se vectoriza y se recupera por separado."],
            ["Embedding", "Representación numérica de un texto que permite medir similitud."],
            ["BM25", "Modelo de búsqueda léxica por frecuencia de términos."],
            ["Reranking", "Reordenamiento de los candidatos por un modelo de relevancia."],
            ["Intent", "Intención detectada en la pregunta por reglas deterministas."],
            ["JTI", "Identificador único de un token, usado para evitar su reutilización."],
            ["Ledger", "Registro del estado de cada archivo del curso para la indexación."],
            ["Course Map", "Estructura del curso tomada de Moodle (secciones, actividades y enlaces)."],
            ["Compuerta de evidencia", "Regla que permite al asistente abstenerse cuando no hay fuente suficiente."],
            ["Sesión de prueba", "Conversación del docente o administrador que no cuenta en las estadísticas."],
        ],
        [3.6, 12.6], caption="Glosario de términos.",
    )

    H2("12.2 Endpoints principales")
    add_table(
        ["Método", "Ruta", "Rol", "Propósito"],
        [
            ["GET", "/f1/login", "Público", "Formulario de login y enlace de docente."],
            ["POST", "/f1/login", "Público", "Login de administrador con usuario y contraseña."],
            ["GET", "/f1/gestor", "Admin y docente", "Interfaz del gestor."],
            ["GET", "/f1/api/me", "Admin y docente", "Identidad, rol y cursos de la sesión."],
            ["GET", "/f1/asistente/<token>", "Público", "Página del chat del asistente."],
            ["POST", "/f1/asistente/<token>/mensaje", "Público", "Enviar una pregunta al asistente."],
            ["GET", "/f1/asistente/<token>/conversaciones", "Público", "Historial de conversaciones del alumno."],
            ["GET", "/f1/api/docente/cursos", "Docente", "Cursos del docente con colección y asistente."],
            ["GET", "/f1/api/docente/cursos/{id}/conocimiento", "Docente", "Materiales del curso y su estado."],
            ["GET", "/f1/api/docente/cursos/{id}/funcionalidad/{nombre}", "Docente", "Contenido real de una funcionalidad."],
            ["POST", "/f1/api/docente/cursos/{id}/vectorizar", "Docente", "Vectorizar pendientes del curso."],
            ["GET", "/f1/api/docente/cursos/{id}/vectorizar/estado", "Docente", "Estado de la vectorización."],
            ["PATCH", "/f1/api/docente/cursos/{id}/asistente", "Docente", "Configuración del asistente por aula."],
            ["PUT", "/f1/api/docente/cursos/{id}/asistente/documentos-excluidos", "Docente", "Documentos que no se usarán."],
            ["GET", "/f1/api/docente/cursos/{id}/pruebas", "Docente", "Conversaciones de prueba del curso."],
            ["POST", "/f1/api/auto/preview", "Admin", "Vista previa del curso de Moodle."],
            ["POST", "/f1/api/auto/download-validate", "Admin", "Descarga y validación de archivos."],
            ["POST", "/f1/api/auto/confirm", "Admin", "Confirmar y vectorizar."],
            ["POST", "/f1/api/semi/*", "Admin", "Carga manual y vectorización."],
            ["POST", "/f1/api/update/*", "Admin", "Reemplazo o actualización de documentos."],
            ["GET", "/f1/api/collections", "Admin", "Listar colecciones y su esquema."],
            ["POST", "/f1/api/collections", "Admin", "Crear una colección."],
            ["DELETE", "/f1/api/collections/{nombre}", "Admin", "Eliminar una colección."],
            ["GET", "/f1/api/collections/{nombre}/documents", "Admin", "Documentos de una colección."],
            ["GET", "/f1/api/rds", "Admin", "Listar aulas RD."],
            ["POST", "/f1/api/rds", "Admin", "Crear un aula RD."],
            ["POST", "/f1/api/rds/{id}/courseids", "Admin", "Agregar un curso a una RD."],
            ["GET", "/f1/api/asistentes", "Admin", "Listar asistentes."],
            ["POST", "/f1/api/asistentes", "Admin", "Crear un asistente."],
            ["PATCH", "/f1/api/asistentes/{id}", "Admin", "Editar un asistente."],
            ["GET", "/f1/api/asistentes/plantilla", "Admin", "Plantilla Idania."],
            ["GET", "/f1/api/config/{etapa}", "Admin", "Estrategia activa y catálogo de una etapa."],
            ["GET", "/f1/api/config/defaults-asistente", "Admin", "Embeddings por defecto para asistentes."],
            ["GET", "/f1/api/usuarios", "Admin", "Listar usuarios y roles."],
            ["GET", "/f1/api/admin/cursos/{id}/conocimiento", "Admin", "Visor de conocimiento del curso."],
            ["GET", "/f1/api/jobs/{id}", "Admin", "Estado de un trabajo de vectorización."],
            ["GET", "/f1/api/health", "Público", "Estado del servicio."],
        ],
        [1.6, 7.4, 2.4, 4.8],
        caption="Endpoints principales agrupados por prefijo y rol.",
        font_size=8.3, header_size=8.4,
    )


def build():
    global DOC
    DOC = Document()
    cover()

    sec_resumen()
    sec_alcance()
    sec_arquitectura()
    sec_datos()
    sec_roles()
    sec_modulos()
    sec_funcionalidades()
    sec_etl()
    sec_seguridad()
    sec_despliegue()
    sec_limitaciones()
    sec_anexos()

    OUT.parent.mkdir(parents=True, exist_ok=True)
    DOC.save(str(OUT))
    print(f"Generado: {OUT}")


if __name__ == "__main__":
    build()
