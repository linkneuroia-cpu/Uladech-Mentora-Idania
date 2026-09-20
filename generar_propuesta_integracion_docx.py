"""
generar_propuesta_integracion_docx.py
=====================================
Genera el documento Word de propuesta de integración entre las tres
funcionalidades del ecosistema educaia.uladech.edu.pe:

    F1 = Idania (asistente RAG de curso)  -> este repositorio
    F2 = Asistente de retroalimentación Moodle
    F3 = Gestor de evaluación oral con IA

Ejecutar:  python generar_propuesta_integracion_docx.py
Salida:    Propuesta_Integracion_F1_F2_F3.docx

Corresponde a la Parte D.3 del PLAN.md.
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


AZUL = RGBColor(0x1F, 0x3B, 0x73)
AZUL_CLARO = RGBColor(0x2E, 0x5C, 0xA8)
GRIS = RGBColor(0x44, 0x44, 0x44)


def set_heading_style(doc):
    for i in range(1, 4):
        try:
            style = doc.styles[f"Heading {i}"]
            style.font.color.rgb = AZUL
            style.font.name = "Calibri"
        except KeyError:
            pass


def add_bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9.5)
    return table


def build():
    doc = Document()
    set_heading_style(doc)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # ── Portada ────────────────────────────────────────────────────
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Propuesta de Integración F1 ↔ F2 ↔ F3")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = AZUL

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Idania (asistente RAG) · Retroalimentación Moodle · Evaluación oral con IA")
    r.font.size = Pt(13)
    r.font.color.rgb = AZUL_CLARO

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "Ecosistema educaia.uladech.edu.pe\n"
        "ULADECH Católica — Funcionalidad 1\n"
        "Documento de propuesta técnica"
    )
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = GRIS

    doc.add_page_break()

    # ── 1. Resumen ejecutivo ───────────────────────────────────────
    doc.add_heading("1. Resumen Ejecutivo", level=1)
    doc.add_paragraph(
        "Hoy las tres funcionalidades que conviven bajo el dominio "
        "educaia.uladech.edu.pe operan de forma totalmente aislada entre sí: cada "
        "una tiene su propio backend y su propio espacio de URL, sin ningún canal "
        "de comunicación. Esto genera una desconexión concreta y visible para el "
        "estudiante: Idania (F1) puede ver que en una semana existe una "
        "\"Evaluación Oral/Escrita - IA N° 3\", pero no sabe nada de lo que ocurre "
        "dentro de esa herramienta (F3), ni si F2 ya le dejó retroalimentación a "
        "una entrega."
    )
    doc.add_paragraph(
        "Esta propuesta plantea una integración incremental en tres niveles, de "
        "menor a mayor esfuerzo y coordinación entre equipos. El Nivel 1 es "
        "implementable de forma unilateral por F1 con lo que ya existe hoy (el "
        "mapa del aula de Moodle). Los Niveles 2 y 3 requieren acuerdos con los "
        "equipos de F2 y F3: el Nivel 2 se apoya en el propio Moodle como "
        "contrato de datos y el Nivel 3 propone una API directa entre servicios."
    )

    # ── 2. Contexto ────────────────────────────────────────────────
    doc.add_heading("2. Contexto: arquitectura actual compartida", level=1)
    doc.add_paragraph(
        "Las tres funcionalidades comparten un mismo dominio y un mismo punto de "
        "entrada (un proxy inverso nginx en producción), que enruta por ruta. Cada "
        "funcionalidad vive en su propio proceso y su propio repositorio, sin "
        "comunicación entre ellas:"
    )
    add_bullets(doc, [
        "nginx (educaia.uladech.edu.pe) → location /f1/ → Idania (F1), backend "
        "propio. Es el asistente RAG que indexa el material de cada curso.",
        "location /f2/ → Asistente de retroalimentación Moodle (F2), backend propio "
        "en producción en 127.0.0.1:7001.",
        "location /f3/ → Gestor de evaluación oral con IA (F3), backend propio en "
        "producción en 127.0.0.1:7301.",
    ])
    doc.add_paragraph(
        "En producción, F1 solo puede AGREGAR su propio bloque de configuración en "
        "nginx; nunca debe tocar los bloques de F2/F3, que se mantienen tal cual "
        "para no pisar el trabajo de esos equipos. Lo mismo aplica en el sentido "
        "inverso: esta propuesta no asume ningún cambio unilateral sobre F2/F3."
    )
    doc.add_paragraph(
        "En el servidor de preproducción la arquitectura es equivalente pero "
        "basada en contenedores Docker detrás de Traefik (dominio "
        "iauladech.linkneuroia.com, ruteo por prefijo de ruta). Allí se "
        "confirmó, inspeccionando los contenedores en vivo, la presencia de "
        "servicios equivalentes a F2 y F3:"
    )
    add_bullets(doc, [
        "f1_retro_app — proyecto f1-retroalimentador, puerto 5001, ruta "
        "/retroalimentador_tarea.py → F2 (retroalimentación Moodle).",
        "f1_oral_app — proyecto f1-evaluacion-oral, puerto 5000, ruta "
        "/evaluacion_oral.py → F3 (evaluación oral con IA).",
        "f1_asistente_app (f1-asistente-general, puerto 5003) y f1_validador_app "
        "(f1-validador-recursos, puerto 5002): otros servicios del mismo "
        "ecosistema, no son este repositorio de F1 (Idania).",
    ])
    doc.add_paragraph(
        "Nota de alcance: este documento se redacta sin acceso al código de F2 ni "
        "de F3, por lo que describe puntos de integración, no implementaciones "
        "internas de esos sistemas. La identidad de los contenedores de "
        "preproducción como versiones de F2/F3 ya fue confirmada en vivo; lo que "
        "queda por acordar es el contenido de cada integración."
    )

    # ── 3. Motivación ──────────────────────────────────────────────
    doc.add_heading("3. Motivación: un caso real", level=1)
    doc.add_paragraph(
        "En el curso \"Introducción a la Odontología\" (courseid 1472 en "
        "producción), la Semana 3 incluye un recurso real de Moodle llamado "
        "\"Evaluación Oral/Escrita - IA N° 3\" (enlace real "
        "mod/url/view.php?id=54728). Ese enlace apunta, con altísima "
        "probabilidad, a la herramienta de evaluación oral (F3)."
    )
    doc.add_paragraph(
        "Cuando un estudiante le pregunta a Idania cómo prepararse para esa "
        "semana, Idania puede listar el recurso por su nombre y darle el enlace "
        "real, pero no puede decirle nada más: no sabe qué formato exacto tiene, "
        "qué se evalúa por dentro, ni si ya tiene un resultado. Es exactamente el "
        "tipo de pregunta donde una integración aportaría valor inmediato."
    )
    doc.add_paragraph(
        "Un caso equivalente aplica a F2: si el docente dejó retroalimentación "
        "sobre una entrega en la herramienta de retroalimentación, Idania no tiene "
        "forma de traerla al chat, aunque el estudiante la pida explícitamente."
    )

    # ── 4. Propuesta en tres niveles ───────────────────────────────
    doc.add_heading("4. Propuesta: tres niveles de integración", level=1)
    doc.add_paragraph(
        "Los niveles son acumulativos (el Nivel 2 no invalida el Nivel 1) y están "
        "ordenados de menor a mayor esfuerzo y coordinación entre equipos."
    )

    doc.add_heading("4.1 Nivel 1 — Detección en el mapa del curso (solo F1)", level=2)
    doc.add_paragraph(
        "Ya implementable hoy por F1, sin coordinación con nadie. Idania ya lee el "
        "mapa real del aula de Moodle (Course Map) y conoce el nombre y el enlace "
        "de cada recurso de cada semana. La propuesta es que, cuando el nombre de "
        "un recurso coincida con patrones genéricos de evaluación oral/escrita o "
        "de retroalimentación (por palabras clave del propio nombre, no "
        "hardcodeado a un curso puntual), Idania lo mencione de forma explícita: "
        "\"esta semana tenés una evaluación oral; se realiza con la herramienta "
        "del enlace X\"."
    )
    doc.add_paragraph(
        "Este nivel no requiere saber nada de lo que pasa DENTRO de F2/F3. Solo "
        "usa el enlace real que Moodle ya expone y que Idania ya puede leer."
    )

    doc.add_heading("4.2 Nivel 2 — Integración vía Moodle (sin API nueva)", level=2)
    doc.add_paragraph(
        "Si F2 y/o F3 escriben sus resultados en estructuras estándar de Moodle "
        "(una nota en el libro de calificaciones, un evento de calendario con "
        "fecha límite, un comentario de retroalimentación), Idania ya puede leer "
        "esos datos con los mecanismos que tiene: el intent de calificaciones "
        "reales (gradebook) y el de fechas de vencimiento. El \"contrato\" de "
        "integración pasa a ser Moodle mismo, sin ningún desarrollo nuevo entre "
        "equipos."
    )
    doc.add_paragraph(
        "Este nivel requiere una decisión de los equipos de F2/F3 sobre DÓNDE "
        "escriben sus resultados, pero no una API nueva."
    )

    doc.add_heading("4.3 Nivel 3 — API directa entre servicios", level=2)
    doc.add_paragraph(
        "El nivel más potente, pero el que exige mayor acuerdo entre equipos. Si "
        "F2 y F3 exponen un endpoint propio (por ejemplo, "
        "GET /f2/api/retroalimentacion/{userid}/{courseid} o "
        "GET /f3/api/resultado/{userid}/{courseid}), autenticado, Idania podría "
        "responder preguntas como \"¿qué retroalimentación me dio F2 en mi última "
        "entrega?\" o \"¿cómo me fue en mi evaluación oral?\" con datos reales de "
        "esos sistemas, consultados bajo demanda desde el chat."
    )
    doc.add_paragraph(
        "Esto no puede construirse de forma unilateral: requiere que los equipos "
        "de F2 y F3 definan y mantengan un contrato de API. Es, por lo tanto, la "
        "conversación formal que esta propuesta busca iniciar."
    )

    # ── 5. Qué se necesita de cada equipo ──────────────────────────
    doc.add_heading("5. Qué se necesita de cada equipo", level=1)
    add_table(
        doc,
        ["Nivel", "Qué aporta F1", "Qué se necesita de F2", "Qué se necesita de F3"],
        [
            [
                "Nivel 1",
                "Detección por nombre en el Course Map y mención explícita al "
                "estudiante (trabajo de F1, sin dependencias).",
                "Nada.",
                "Nada.",
            ],
            [
                "Nivel 2",
                "Lectura de notas/eventos/comentarios de Moodle con mecanismos ya "
                "existentes.",
                "Definir en qué estructura de Moodle escribe su retroalimentación "
                "(nota, comentario, etc.).",
                "Definir en qué estructura de Moodle escribe su resultado (nota, "
                "evento de calendario, etc.).",
            ],
            [
                "Nivel 3",
                "Cliente que consume el endpoint de F2/F3 bajo demanda, con "
                "aislamiento por estudiante.",
                "Exponer un endpoint autenticado de retroalimentación por "
                "estudiante/curso.",
                "Exponer un endpoint autenticado de resultados de evaluación por "
                "estudiante/curso.",
            ],
        ],
    )

    # ── 6. Riesgos y consideraciones ───────────────────────────────
    doc.add_heading("6. Riesgos y consideraciones", level=1)
    add_bullets(doc, [
        "Aislamiento de datos por estudiante: cualquier dato que F1 traiga de F2 o "
        "F3 debe respetar el mismo criterio ya aplicado en calificaciones y "
        "progreso — nunca exponer datos de un estudiante a otro. El userid de la "
        "sesión de F1 es el único identificador válido para consultar.",
        "Propiedad del contrato de API: si el Nivel 3 avanza, hay que definir "
        "explícitamente qué equipo es dueño del contrato y qué pasa cuando cambia "
        "(versionado, compatibilidad hacia atrás).",
        "Autenticación entre servicios: el Nivel 3 requiere un mecanismo de "
        "autenticación entre backends (token de servicio o red interna), que debe "
        "acordarse con los equipos.",
        "Dependencia de disponibilidad: si F2/F3 no responden, F1 debe degradar de "
        "forma honesta (\"no pude consultar esa herramienta ahora\") y nunca "
        "inventar un resultado.",
        "Preproducción: ya se confirmó que f1_retro_app (F2) y f1_oral_app (F3) "
        "corren en el servidor de preproducción; puede usarse ese entorno para "
        "probar la integración sin tocar producción.",
    ])

    # ── 7. Recomendación ───────────────────────────────────────────
    doc.add_heading("7. Recomendación", level=1)
    doc.add_paragraph(
        "Se recomienda arrancar de inmediato por el Nivel 1: es viable sin "
        "coordinación, no toca a los otros equipos y ya mejora de forma concreta "
        "la experiencia del estudiante (mencionar explícitamente las evaluaciones "
        "orales y las instancias de retroalimentación que aparecen en la semana)."
    )
    doc.add_paragraph(
        "En paralelo, se propone abrir una conversación formal con los equipos de "
        "F2 y F3 para evaluar el Nivel 3 (API directa), usando el Nivel 2 como "
        "paso intermedio de bajo costo. El caso real de la \"Evaluación "
        "Oral/Escrita - IA N° 3\" de la Semana 3 del curso 1472 es el ejemplo "
        "concreto con el que se puede ilustrar el valor de la integración."
    )

    out = "Propuesta_Integracion_F1_F2_F3.docx"
    doc.save(out)
    print(f"OK -> {out}")


if __name__ == "__main__":
    build()
