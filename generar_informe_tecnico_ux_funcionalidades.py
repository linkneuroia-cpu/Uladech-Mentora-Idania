"""
generar_informe_tecnico_ux_funcionalidades.py
=============================================
Genera el informe técnico (Word) de la reorganización UX por funcionalidad
del gestor F1: wizard admin reordenado, Navegabilidad estructurada y
Conocimiento idanIA con el mismo árbol que Automático.

Ejecutar:  python generar_informe_tecnico_ux_funcionalidades.py
Salida:    doc/Informe_Tecnico_UX_Funcionalidades.docx
"""
import os

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

AZUL = RGBColor(0x1F, 0x3B, 0x73)
AZUL_CLARO = RGBColor(0x2E, 0x5C, 0xA8)
GRIS = RGBColor(0x44, 0x44, 0x44)


def _kv(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9.5)
    return table


def _bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def build():
    doc = Document()
    for i in range(1, 4):
        try:
            style = doc.styles[f"Heading {i}"]
            style.font.color.rgb = AZUL
            style.font.name = "Calibri"
        except KeyError:
            pass
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # Portada
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Informe técnico — UX por funcionalidad")
    run.font.size = Pt(24)
    run.font.color.rgb = AZUL
    run.bold = True

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Gestor F1 ULADECH · wizard admin, Navegabilidad y Conocimiento idanIA")
    run.font.size = Pt(13)
    run.font.color.rgb = AZUL_CLARO

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Producción: https://educaia.uladech.edu.pe/f1/gestor\nFecha: 20/10/2026")
    run.font.size = Pt(10.5)
    run.font.color.rgb = GRIS
    doc.add_page_break()

    # 1. Resumen
    doc.add_heading("1. Resumen ejecutivo", level=1)
    doc.add_paragraph(
        "Se reorganizó la experiencia de uso del gestor en torno a las funcionalidades "
        "del asistente (Documentos, Navegabilidad del aula y Notas), tanto para el "
        "administrador como para el docente. El cambio no modifica la lógica del RAG: "
        "reutiliza los mismos datos y componentes para que la creación de un asistente "
        "completo sea guiada, y para que lo que se muestra en el panel del docente sea "
        "idéntico a lo que usa la vectorización automática."
    )
    _kv(doc, ["Ámbito", "Resultado"], [
        ["Wizard admin", "Orden Curso → Identidad → Funcionalidades → Avanzado → Revisar; plantilla Idania y embeddings por defecto precargados."],
        ["Navegabilidad (docente/admin)", "Línea de tiempo por semana con resumen en texto plano, semana actual resaltada, actividades por tipo y vencimientos cruzados."],
        ["Conocimiento idanIA", "Usa el MISMO árbol, orden y criterios que el módulo Automático (verificado en el curso 1472)."],
        ["Lista/edición admin", "Chips de funcionalidades, filtro y vista \"Ver por funcionalidad\"; edición por pestañas, incluida la de prueba de chat."],
        ["Pruebas", "29 pruebas automatizadas en verde."],
    ])
    doc.add_paragraph()

    # 2. Problemas reportados
    doc.add_heading("2. Problemas reportados y diagnóstico", level=1)
    doc.add_paragraph(
        "El pedido de Miguel fue: (1) una análisis UX del admin para que la configuración "
        "se vea por funcionalidad y sea fácil crear un asistente completo; (2) que la "
        "Navegabilidad del docente se vea estructurada; y (3) que Conocimiento idanIA se "
        "vea ordenado igual que la vectorización automática."
    )
    doc.add_paragraph(
        "El diagnóstico encontró tres causas concretas:"
    )
    _bullets(doc, [
        "El wizard seguía un orden técnico (Configuración RAG en el paso 2) y no precargaba la plantilla Idania; en asistentes normales tomaba el embedding por defecto global (MiniLM) en vez del default para asistentes (openai/text-embedding-3-small), repitiendo el riesgo del bug de dimensiones del curso 1440.",
        "La Navegabilidad mostraba el HTML crudo de Moodle escapado y cortado (se veían etiquetas <p>), todas las actividades en una línea y las fechas sueltas, sin resaltar la semana actual.",
        "Conocimiento idanIA agrupaba en el backend con su propia lógica, mientras Automático agrupaba en el frontend con groupDocsBySemana y computeRepeats; el mismo curso se veía distinto en las dos pantallas.",
    ])
    doc.add_paragraph()

    # 3. Cambios backend
    doc.add_heading("3. Cambios en el backend", level=1)
    doc.add_paragraph(
        "Se mantuvieron los endpoints existentes y se ampliaron sin cambiar la lógica de recuperación:"
    )
    _kv(doc, ["Cambio", "Detalle"], [
        ["GET /api/docente/cursos/{id}/conocimiento",
         "Devuelve 'modules' con la misma forma que GET /api/auto/preview (module_name, resource_id, files[...]), respetando el orden de Moodle. Cada archivo incluye estado (excluido > vectorizado > no_soportado > pendiente), is_duplicate, existing_chunks/date, moodle_upload_date y fileurl SIN token. Se mantienen resumen, semana_actual, job_vectorizacion y collection_name. Se eliminó 'semanas'."],
        ["_conocimiento_modules / _estado_archivo",
         "Funciones puras que calculan los estados y el resumen, sin consultar la base; reemplazan el agrupamiento propio."],
        ["GET /api/docente/cursos/{id}/funcionalidad/navegacion",
         "Navegación estructurada: secciones con numero, nombre, visible, es_actual, resumen_texto (HTML de Moodle convertido a texto con _html_a_texto), actividades con tipo, tipo_label (MODNAME_LABELS), url, visible y 'vence' cruzado por nombre normalizado; proximas_entregas (3) y fechas_status/fechas_message. Si el calendario está 'unavailable', no se inventan fechas."],
        ["GET /api/asistentes/plantilla",
         "Solo admin. Devuelve la plantilla Idania (prompt maestro, bienvenida y reencuentro) para precargar el wizard y la edición."],
        ["GET /api/config/defaults-asistente",
         "Devuelve {dense: openai/text-embedding-3-small, sparse: bm25}, los valores por defecto con los que se crean asistentes."],
        ["AutoCore.preview_course",
         "Acepta collection_name=None para listar todos los materiales de un curso todavía sin colección, reutilizando el mismo recorrido de Moodle."],
        ["query_router.MODNAME_LABELS",
         "Se expuso el mapa de etiquetas de actividades como constante pública y se reutiliza en la navegación."],
    ])
    doc.add_paragraph()

    # 4. Cambios frontend
    doc.add_heading("4. Cambios en la interfaz", level=1)
    doc.add_paragraph(
        "Se aplicó el principio de reutilizar, no duplicar. Los tres componentes compartidos son:"
    )
    _kv(doc, ["Componente", "Uso"], [
        ["materialesTreeHTML(mods, opts)",
         "Extraído del render de Automático sin cambiar su comportamiento. Incluye computeRepeats, filtro por tipo, groupDocsBySemana, encabezados con contadores y filas con ícono, estado, fechas, 'Ver' y tamaño. Se parametrizó estadoDe, rowLeft, rowExtra, statusBadge, headerCounter, openGroup e idPrefix. Lo usan Automático, Conocimiento idanIA del docente y la edición del admin."],
        ["navegacionHTML(data)",
         "Encabezado con semana actual y próximas entregas; línea de tiempo por sección con resumen limpio (2 líneas + 'ver más'), actividades agrupadas por tipo, enlaces a Moodle, vencimientos en su actividad, ocultas atenuadas y acordeón (abre la semana actual y la siguiente)."],
        ["funcionalidadCardHTML(...)",
         "Tarjeta por funcionalidad con interruptor, descripción, ejemplos y vista previa perezosa del contenido real. Si está apagada, muestra el mensaje fijo que dará el asistente. La usan el wizard, la edición admin y el panel docente."],
    ])
    doc.add_paragraph()
    doc.add_paragraph("Cambios por pantalla:")
    _bullets(doc, [
        "Wizard admin: nuevo orden Curso → Identidad → Funcionalidades → Avanzado (opcional, colapsado) → Revisar. Identidad precargada con la plantilla Idania y botón 'Restaurar plantilla'. Los embeddings se toman de defaults-asistente (no del default global). Al crear muestra la URL, un iframe de prueba (?prueba=1) y el botón 'Ir al asistente'.",
        "Lista de asistentes: columnas de funcionalidades con chips (atenuados/tachados si están apagados), buscador, filtro por funcionalidad y selector 'Ver por: Asistente | Funcionalidad'. En la vista por funcionalidad hay tres bloques con interruptores rápidos y recuentos.",
        "Edición de asistente: pestañas General · Funcionalidades · Conocimiento idanIA · Probar chat · Avanzado (RAG). En asistentes de RD, las pestañas por curso piden elegir el aula de la RD.",
        "Panel del docente: Conocimiento idanIA con chips de estado y tipo, buscador, botón 'Vectorizar pendientes (N)' con progreso, barra fija de cambios sin guardar, y la misma nota de alcance por aula para asistentes de RD. Navegabilidad con navegacionHTML. Funcionalidades con funcionalidadCardHTML.",
        "Visor admin 'Qué sabe el asistente': la pestaña Estructura ahora usa navegacionHTML (misma representación del curso).",
    ])
    doc.add_paragraph()

    # 5. Criterio de aceptación
    doc.add_heading("5. Criterio de aceptación: un solo árbol del curso", level=1)
    doc.add_paragraph(
        "Se comparó, en producción y para el curso 1472, el orden y los nombres de los "
        "módulos de GET /api/auto/preview contra los de GET /api/docente/.../conocimiento. "
        "El resultado fue idéntico: mismos módulos, mismo orden y mismos archivos por grupo. "
        "La única diferencia entre ambas pantallas es la columna de acciones."
    )
    _kv(doc, ["Comprobación (curso 1472)", "Resultado"], [
        ["Módulos Automático", "119"],
        ["Módulos Conocimiento idanIA", "119"],
        ["Mismo orden y nombres", "Sí"],
        ["Archivos totales", "124"],
        ["Vectorizados / Pendientes / No soportados / Excluidos", "111 / 6 / 7 / 0"],
    ])
    doc.add_paragraph()

    # 6. Verificación en producción
    doc.add_heading("6. Verificación en producción", level=1)
    _bullets(doc, [
        "Contenedor gestor_f1 arriba y 'Application startup complete' sin Traceback.",
        "Admin: GET /api/asistentes/plantilla → 200; GET /api/config/defaults-asistente → {dense: openai/text-embedding-3-small, sparse: bm25}.",
        "Docente 41607454 (curso 1472): Conocimiento 200 con 119 módulos y resumen correcto; la respuesta ya no incluye 'semanas'.",
        "Navegabilidad 1472: semana_actual 8, 18 secciones, resumen_texto limpio ('Datos del curso…'), es_actual marcado. El calendario respondió 'unavailable' y por eso no se listaron fechas inventadas.",
        "Automático sigue respondiendo (preview 119 módulos, 124 archivos) y su árbol conserva el comportamiento anterior.",
    ])
    doc.add_paragraph()

    # 7. Pruebas
    doc.add_heading("7. Pruebas", level=1)
    doc.add_paragraph("Se ejecutó la suite automatizada: 29 pruebas en verde.")
    _kv(doc, ["Archivo", "Cobertura"], [
        ["tests/test_navegacion.py", "_html_a_texto, cruce de vencimientos por nombre normalizado, calendario unavailable sin fechas inventadas, es_actual y actividades ocultas."],
        ["tests/test_docente_conocimiento.py", "modules conserva el orden de Moodle, fileurl sin token, prioridad de estados y resumen; 409 con job en curso; filenames=null encola solo pendientes."],
        ["tests/test_roles.py", "Permisos por rol y resolución de cursos en vivo (incluye sesión por contraseña y ERP caído → 503)."],
        ["tests/test_funcionalidades.py", "Mensajes fijos por funcionalidad apagada, documentos=False sin retrieval y query_filter must_not de exclusiones."],
        ["tests/test_docente_config.py", "Resolución de asistente/colección normal o de RD y exclusiones por curso."],
        ["tests/test_prueba_chat.py", "Sesiones de prueba aisladas y excluidas de estadísticas."],
    ])
    doc.add_paragraph(
        "Los specs E2E (tests/e2e/) no se ejecutaron completos en esta sesión porque "
        "apuntan a producción y algunos crean/modifican datos; la verificación se hizo "
        "por API y con renders del gestor. Quedan como regresión pendiente de una corrida "
        "controlada."
    )
    doc.add_paragraph()

    # 8. Despliegue
    doc.add_heading("8. Despliegue", level=1)
    _bullets(doc, [
        "Se ejecutó bash deploy/upload.sh y bash deploy/deploy.sh; el build reutilizó la caché y el contenedor se recreó correctamente.",
        "Este incremento no requirió migraciones de base de datos.",
        "No se hizo commit ni push (según la regla del plan).",
    ])
    doc.add_paragraph()

    # 9. Limitaciones y notas
    doc.add_heading("9. Limitaciones y notas", level=1)
    _bullets(doc, [
        "Las exclusiones marcadas en la vista previa de Documentos del wizard se aplican al crear solo si esa vista previa se abrió (los checks se renderizan al cargar el árbol). Para asistentes de RD, la creación aplica exclusiones/vectorización por curso solo en asistentes normales con curso; la gestión por aula queda en el panel del docente y el visor admin.",
        "El botón 'Vectorizar pendientes al crear' es opcional y se ejecuta en segundo plano; la pantalla final muestra la prueba de chat, no una barra de progreso persistente del job.",
        "La administración de configuraciones por aula de una RD sigue siendo por curso (no se toca la configuración global de la RD).",
    ])
    doc.add_paragraph()

    # 10. Archivos tocados
    doc.add_heading("10. Archivos principales modificados", level=1)
    _kv(doc, ["Archivo", "Cambio"], [
        ["app.py", "modules en /conocimiento; _html_a_texto y _navegacion_payload; endpoints de plantilla y defaults; preview sin colección."],
        ["query_router.py", "MODNAME_LABELS público."],
        ["frontend/interfaz.html", "materialesTreeHTML, navegacionHTML, funcionalidadCardHTML; refactor de aRenderTree; wizard, lista, edición, panel docente y visor admin."],
        ["tests/test_navegacion.py", "Nuevo."],
        ["tests/test_docente_conocimiento.py", "Actualizado a modules."],
    ])
    doc.add_paragraph()

    # Cierre
    p = doc.add_paragraph()
    run = p.add_run(
        "Conclusión: la configuración ahora se entiende por funcionalidad, el docente ve "
        "el curso con la misma representación que la vectorización automática y el admin "
        "puede crear un asistente completo con solo elegir el curso y avanzar, con la "
        "plantilla Idania y los embeddings correctos por defecto."
    )
    run.font.color.rgb = GRIS

    os.makedirs("doc", exist_ok=True)
    out = os.path.join("doc", "Informe_Tecnico_UX_Funcionalidades.docx")
    doc.save(out)
    print(f"Generado: {out}")


if __name__ == "__main__":
    build()
