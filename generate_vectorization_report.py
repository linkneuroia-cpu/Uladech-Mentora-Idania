"""
generate_vectorization_report.py
=================================
Arma el reporte Word "Idania - Evaluacion de Vectorizacion (1472 y 1541).docx"
a partir de evaluate_vectorization_results.json (ver evaluate_vectorization_quality.py).
Mismo estilo/tono que los demás informes en Data prueba/ (Título, Heading 1,
párrafos Normal, tablas "Light Grid Accent 1").
"""
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).parent
DATA = json.loads((ROOT / "evaluate_vectorization_results.json").read_text(encoding="utf-8"))
OUT = ROOT / "Data prueba" / "Idania - Evaluacion de Vectorizacion (1472 y 1541).docx"

C1472 = DATA["Curso_1472"]
C1541 = DATA["Curso_1541"]

FLAG_LABELS = {
    "too_short": "Demasiado corto (<20 car.)",
    "too_long": "Demasiado largo (>3000 car.)",
    "replacement_char": "Carácter de reemplazo Unicode (encoding roto)",
    "repeated_char_run": "Secuencia de 8+ caracteres repetidos",
    "low_alpha_ratio": "Ratio de letras <35% en chunk de prosa",
}


def add_bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    return table


def pct(n, d):
    return f"{n / d * 100:.1f}%" if d else "N/A"


doc = Document()

# ── Portada / título ────────────────────────────────────────────────
doc.add_heading("Evaluación de vectorización — Curso_1472 y Curso_1541", level=0)
intro = doc.add_paragraph()
intro.add_run(
    "Auditoría de la calidad de EXTRACCIÓN, CHUNKING E INDEXACIÓN real de los archivos de "
    "ambos cursos, ejecutada contra el mirror local de desarrollo (Qdrant localhost:6333 + "
    "Postgres U_F1_Profundizacion_desarrollo), copia byte-a-byte de las colecciones reales de "
    "producción Curso_1472 (Introducción a la Odontología) y Curso_1541 (Filosofía de la "
    "Educación). Es distinta de la evaluación de calidad de RESPUESTAS del tutor (200 pruebas "
    "de alucinaciones, rendimiento) ya reportada por separado: acá se audita qué tan bien se "
    "extrajo, chunkeó e indexó el contenido de cada archivo, no la calidad de lo que el LLM "
    "responde. 100% solo lectura — ningún vector, payload, chunk ni fila del ledger fue "
    "modificado."
)

# ── Resumen ejecutivo ────────────────────────────────────────────────
doc.add_heading("Resumen ejecutivo", level=1)
resumen = [
    f"Cobertura de contenido: Curso_1472 tiene 106 documentos / 4,337 chunks vectorizados; "
    f"Curso_1541 tiene 73 documentos / 3,593 chunks. En ambos, el 100% de los archivos "
    f"presentes en Qdrant tienen chunks reales y legibles — no hay documentos \"fantasma\" "
    f"(marcados como vectorizados sin contenido).",
    f"Hallazgo operativo real: la tabla curso_archivos (el \"ledger\" de ledger.py) tiene "
    f"CERO filas para Curso_1472 y CERO para Curso_1541 — nunca se ejecutó sync_curso_archivos() "
    f"sobre estos dos cursos. El contenido en Qdrant es correcto, pero no hay cola de trabajo "
    f"incremental activa: un cambio futuro en Moodle no se detectará automáticamente.",
    f"Hallazgo de calidad más importante: el fallback de OCR de fórmulas/gráficos vectoriales "
    f"(_find_formula_regions + pix2tex, en definitions.py) dispara sobre contenido que NO es "
    f"una fórmula (tablas sin líneas completas, diagramas, fotos) y produce texto "
    f"pseudo-LaTeX sin sentido en decenas de chunks de AMBAS colecciones (80 chunks en 1472, "
    f"35 en 1541, con el patrón 'begin{{array}}'; hasta el literal '[object Object]' aparece "
    f"filtrado en 6 chunks). En al menos un caso confirmado (Curso_1541) esto reemplazó una "
    f"tabla real de datos (estadística de suicidios por grupo de edad) por texto sin sentido.",
    f"Hallazgo de chunking: los chunks chunk_type='table_row' (fila individual de tabla) se "
    f"generan sin ningún límite de tamaño en chunk_structured_blocks() — a diferencia de "
    f"'table' y 'prose', que sí pasan por un splitter. Se confirmaron 2 chunks de "
    f"\"Sistema de notación dentaria.pdf\" (Curso_1472) de 5,049 y 5,349 caracteres, "
    f"5.6-5.9× el CHUNK_SIZE configurado (900).",
    f"El Índice de Campos Estructurados funciona como está diseñado en ambos cursos: el "
    f"sílabo real de cada curso aparece 100% marcado is_repeated_template=False (evidencia de "
    f"alta autoridad) — 116/116 filas en 1472, 143/143 filas (docx+pdf) en 1541.",
    f"Cobertura semanal: Curso_1541 cubre las 16 semanas completas. Curso_1472 cubre 13/16 — "
    f"faltan semanas 4, 8 y 16. La semana 4 SÍ tiene contenido vectorizado, pero con un typo "
    f"real en Moodle (\"Senana 04\" en vez de \"Semana 04\"), lo que rompe cualquier "
    f"coincidencia por texto \"Semana N\"; las semanas 8 y 16 no tienen ningún chunk asociado.",
]
add_bullets(doc, resumen)

# ── Metodología ──────────────────────────────────────────────────────
doc.add_heading("Metodología", level=1)
doc.add_paragraph(
    "Script: evaluate_vectorization_quality.py (raíz del proyecto). Resultados intermedios: "
    "evaluate_vectorization_results.json. 100% de lectura vía scroll() de Qdrant (con_payload "
    "completo, sin límite — se recorrieron los 4,337 + 3,593 puntos reales de ambas "
    "colecciones, no una muestra) y consultas SELECT de solo lectura contra curso_archivos y "
    "coleccion_campos_estructurados en Postgres (mirror dev, ENV=DESARROLLO). No se hizo ningún "
    "upsert, set_payload ni escritura de ningún tipo."
)
metodo_items = [
    "Cobertura ledger vs Qdrant: se comparó el listado de curso_archivos por colección contra "
    "los document_filename distintos realmente presentes en Qdrant.",
    "Calidad de extracción/chunking: por cada chunk real se midió longitud de texto y se "
    "aplicaron 5 heurísticas sobre el texto real del payload — chunk <20 caracteres, "
    ">3,000 caracteres, presencia del carácter de reemplazo Unicode U+FFFD (encoding roto), "
    "una secuencia de 8+ caracteres idénticos consecutivos, y ratio de letras <35% en chunks "
    "de prosa (no aplica a chunks de tabla, donde pipes/números son legítimos).",
    "Contenido 'garbled' pseudo-LaTeX: búsqueda de 4 patrones de texto (mathrm{...} con "
    "palabras sin sentido, scriptstyle, begin{array}, el literal '[object Object]') que son "
    "la firma del fallback de OCR de fórmulas (pix2tex) alucinando sobre contenido que no es "
    "una fórmula real — ver definitions.py::_find_formula_regions/_ocr_page_region.",
    "Separación tabla/fila/prosa: para cada documento pdf/docx (los únicos que pasan por "
    "chunk_structured_blocks(), ver definitions.py) se verificó si sus chunks incluyen "
    "chunk_type='table'/'table_row' además de 'prose', o si quedaron solo como prosa.",
    "Índice de Campos Estructurados: consulta directa a coleccion_campos_estructurados "
    "(campos distintos, filas por is_repeated_template, y las filas del/los archivo(s) de "
    "sílabo de cada curso).",
    "Semanas: se parseó el número de semana desde section_name (metadata real de Moodle "
    "persistida en el payload) con la expresión 'Semana N', y se comparó contra las 16 "
    "semanas típicas del curso.",
]
add_bullets(doc, metodo_items)

# ── Tabla resumen comparativa ────────────────────────────────────────
doc.add_heading("Tabla resumen comparativa", level=1)

sil_1472 = C1472["campos_estructurados"]["silabo_docs"]
sil_1541 = C1541["campos_estructurados"]["silabo_docs"]
sil_1472_rows = sum(d["rows_unique"] + d["rows_repeated_template"] for d in sil_1472)
sil_1472_false = sum(d["rows_unique"] for d in sil_1472)
sil_1541_rows = sum(d["rows_unique"] + d["rows_repeated_template"] for d in sil_1541)
sil_1541_false = sum(d["rows_unique"] for d in sil_1541)

comp_rows = [
    ("Documentos únicos vectorizados", C1472["total_documents"], C1541["total_documents"]),
    ("Chunks totales", C1472["total_chunks"], C1541["total_chunks"]),
    ("Chunks — prosa", C1472["global_chunk_type_counts"].get("prose", 0), C1541["global_chunk_type_counts"].get("prose", 0)),
    ("Chunks — tabla completa", C1472["global_chunk_type_counts"].get("table", 0), C1541["global_chunk_type_counts"].get("table", 0)),
    ("Chunks — fila de tabla", C1472["global_chunk_type_counts"].get("table_row", 0), C1541["global_chunk_type_counts"].get("table_row", 0)),
    ("Chunks con alguna bandera de calidad", "151 (3.5%)", "208 (5.8%)"),
    ("Chunks con patrón 'begin{array}' (garbled OCR fórmula)", C1472["garbled_pattern_counts"].get("begin_array", 0), C1541["garbled_pattern_counts"].get("begin_array", 0)),
    ("Chunks con literal '[object Object]'", C1472["garbled_pattern_counts"].get("object_object", 0), C1541["garbled_pattern_counts"].get("object_object", 0)),
    ("Docs pdf/docx con separación tabla+prosa funcionando",
     f"{len(C1472['structured_docs_ok'])}/{len(C1472['structured_docs_ok']) + len(C1472['structured_docs_prose_only']) + len(C1472['structured_docs_table_only'])} ({pct(len(C1472['structured_docs_ok']), len(C1472['structured_docs_ok']) + len(C1472['structured_docs_prose_only']) + len(C1472['structured_docs_table_only']))})",
     f"{len(C1541['structured_docs_ok'])}/{len(C1541['structured_docs_ok']) + len(C1541['structured_docs_prose_only']) + len(C1541['structured_docs_table_only'])} ({pct(len(C1541['structured_docs_ok']), len(C1541['structured_docs_ok']) + len(C1541['structured_docs_prose_only']) + len(C1541['structured_docs_table_only']))})"),
    ("Docs pdf/docx solo prosa (sin tabla detectada)", len(C1472["structured_docs_prose_only"]), len(C1541["structured_docs_prose_only"])),
    ("Filas en curso_archivos (ledger)", C1472["ledger_comparison"]["ledger_total_filas"], C1541["ledger_comparison"]["ledger_total_filas"]),
    ("Documentos en Qdrant sin fila en el ledger", len(C1472["ledger_comparison"]["en_qdrant_pero_no_en_ledger"]), len(C1541["ledger_comparison"]["en_qdrant_pero_no_en_ledger"])),
    ("Campos estructurados distintos", C1472["campos_estructurados"]["distinct_fields"], C1541["campos_estructurados"]["distinct_fields"]),
    ("Filas en el Índice de Campos Estructurados", C1472["campos_estructurados"]["total_field_records"], C1541["campos_estructurados"]["total_field_records"]),
    ("Filas del sílabo con is_repeated_template=False (autoridad alta)",
     f"{sil_1472_false}/{sil_1472_rows} ({pct(sil_1472_false, sil_1472_rows)})",
     f"{sil_1541_false}/{sil_1541_rows} ({pct(sil_1541_false, sil_1541_rows)})"),
    ("Semanas (1-16) con contenido vectorizado", f"{len(C1472['weeks_present'])}/16", f"{len(C1541['weeks_present'])}/16"),
    ("Chunks sin campo chunk_type en el payload (legado)",
     f"{C1472['missing_chunk_type_field_total']} ({pct(C1472['missing_chunk_type_field_total'], C1472['total_chunks'])})",
     f"{C1541['missing_chunk_type_field_total']} ({pct(C1541['missing_chunk_type_field_total'], C1541['total_chunks'])})"),
]
add_table(doc, ["Métrica", "Curso_1472", "Curso_1541"], comp_rows)

# ── Hallazgo 1: ledger vacío ─────────────────────────────────────────
doc.add_heading("Hallazgo 1 — El ledger (curso_archivos) está vacío para ambos cursos", level=1)
doc.add_paragraph(
    "curso_archivos tiene 874 filas en total en el mirror, repartidas entre 14 colecciones "
    "distintas (Curso_173: 110, Curso_533: 101, Curso_1398: 71, etc.) — pero CERO filas para "
    "Curso_1472 y CERO para Curso_1541, exactamente los dos cursos evaluados. El contenido en "
    "Qdrant es correcto y completo (verificado abajo): esto no es una falla de vectorización, "
    "es que sync_curso_archivos() (ledger.py) nunca se ejecutó sobre estos dos cursos — "
    "probablemente porque su carga inicial se hizo antes de que existiera el ledger, o por una "
    "vía distinta (migración manual/semiautomática) que no llama a sync_curso_archivos()/"
    "marcar_vectorizado()."
)
doc.add_paragraph(
    "Consecuencia práctica: hoy no hay forma automática de detectar si un profesor sube un "
    "archivo nuevo, reemplaza uno existente, o borra uno en Moodle para Curso_1472/Curso_1541 — "
    "el mecanismo que existe para eso (ledger.py) no tiene ningún punto de partida para estos "
    "dos cursos."
)

# ── Hallazgo 2: garbled OCR de fórmulas ──────────────────────────────
doc.add_heading("Hallazgo 2 — Contenido \"garbled\" del fallback de OCR de fórmulas", level=1)
doc.add_paragraph(
    "definitions.py::_find_formula_regions() identifica regiones de una página con trazos "
    "vectoriales (fuera de tablas ya detectadas) como candidatas a fórmula, y las manda a "
    "_ocr_page_region() → pix2tex (LaTeX-OCR). Cuando la región NO es realmente una fórmula "
    "(una tabla sin líneas completas que pdfplumber no reconoció como tabla, un diagrama, una "
    "foto), pix2tex alucina texto pseudo-LaTeX sin sentido — ese texto queda igual insertado "
    "en el chunk como si fuera contenido real del curso."
)
doc.add_paragraph("Volumen medido (búsqueda de patrones sobre el texto real de cada chunk):")
garb_rows = [
    ("Patrón \"begin{array}\"", C1472["garbled_pattern_counts"].get("begin_array", 0), C1541["garbled_pattern_counts"].get("begin_array", 0)),
    ("Patrón \"mathrm{palabra_sin_sentido}\"", C1472["garbled_pattern_counts"].get("mathrm_gibberish", 0), C1541["garbled_pattern_counts"].get("mathrm_gibberish", 0)),
    ("Patrón \"scriptstyle\"", C1472["garbled_pattern_counts"].get("scriptstyle", 0), C1541["garbled_pattern_counts"].get("scriptstyle", 0)),
    ("Literal \"[object Object]\"", C1472["garbled_pattern_counts"].get("object_object", 0), C1541["garbled_pattern_counts"].get("object_object", 0)),
]
add_table(doc, ["Patrón detectado", "Curso_1472 (chunks)", "Curso_1541 (chunks)"], garb_rows)
doc.add_paragraph(
    "En Curso_1541 el patrón aparece salpicado en casi todos los \"SESION_XX_FILOSOFIA...pdf\" "
    "(al menos 1-3 chunks por sesión semanal) además de 19 chunks en la bibliografía "
    "\"Filosofía, ética y educación para la ciudadanía en la era digital.pdf\". En Curso_1472 "
    "se concentra sobre todo en \"diagnostico-situacion-amalgama-dental.pdf\" (36 chunks), pero "
    "también toca \"ABC del consultorio odontológico.pdf\", \"Artículo científico.pdf\", "
    "\"TESIS.pdf\" y \"LIBRO.pdf\", entre otros."
)
doc.add_paragraph("Ejemplo confirmado de PÉRDIDA de contenido real (no solo ruido añadido):")
example_p = doc.add_paragraph()
example_p.add_run(
    "Curso_1541 — \"Filosofía, ética y educación para la ciudadanía en la era digital.pdf\", "
    "chunk 1659 (página 416): el chunk debía contener la tabla real \"Suicidios por grupo de "
    "edad entre enero de 2021 y diciembre de 2023\", pero el texto extraído es:"
).italic = False
quote = doc.add_paragraph()
quote.paragraph_format.left_indent = Pt(24)
run = quote.add_run(
    "\"Tabla 3 Suicidios por grupo de edad entre enero de 2021 y diciembre de 2023\\n\\n"
    "( left[begin{array}{l}{{mathrm{Aio}}} {{2o21}} {{2021}} {{2021}} {{mathrm...\""
)
run.italic = True
run.font.size = Pt(10)
doc.add_paragraph(
    "Es decir: en vez de los datos reales de la tabla (edades, cifras de suicidios por año), "
    "el chunk vectorizado contiene texto pseudo-LaTeX sin ningún valor semántico — cualquier "
    "pregunta sobre esa tabla específica no tiene forma de responderse bien porque el dato real "
    "nunca llegó a indexarse."
)
doc.add_paragraph("Otro ejemplo — Curso_1541, \"SESION_16_FILOSOFIA_DE_LA_EDUCACION_2026II.pdf\", chunk 1 (página 1):")
quote2 = doc.add_paragraph()
quote2.paragraph_format.left_indent = Pt(24)
run2 = quote2.add_run(
    "\"( begin{array}{l}{{left[mathrm{popamadesualo}}} {{scriptstyle[mathrm{pemadsealsuponell}}} "
    "{{scriptstyle[mathrm{lemata}mathrm{tinyopearolosesionellonel...\""
)
run2.italic = True
run2.font.size = Pt(10)
doc.add_paragraph(
    "\"popamadesualo\", \"pemadsealsuponell\" no son palabras reales en español ni notación "
    "matemática válida — es alucinación del modelo pix2tex sobre una región que probablemente "
    "es un logo, encabezado gráfico o diagrama de la plantilla de sesión, no una fórmula."
)
doc.add_paragraph(
    "El caso \"[object Object]\" es aparte y más llamativo: es la representación por defecto de "
    "un objeto JavaScript — su sola presencia en el texto extraído (5 chunks en Curso_1472, 1 "
    "en Curso_1541, repartidos en \"Artículo científico.pdf\", \"ABC del consultorio "
    "odontológico.pdf\", \"diagnostico-situacion-amalgama-dental.pdf\", \"TESIS.pdf\" y "
    "\"Filosofía, ética y educación...pdf\") confirma que se trata de salida cruda de un modelo "
    "sin ningún filtro de sanidad antes de insertarse como texto del chunk."
)

# ── Hallazgo 3: table_row sin límite de tamaño ───────────────────────
doc.add_heading("Hallazgo 3 — chunk_type='table_row' sin límite de tamaño", level=1)
doc.add_paragraph(
    "En definitions.py::chunk_structured_blocks(), los chunks 'prose' pasan por prose_splitter "
    "y los chunks 'table' por table_splitter — ambos respetan CHUNK_SIZE=900. Pero cada fila "
    "individual generada por _parse_markdown_table_rows() (chunk_type='table_row') se agrega "
    "directamente como chunk sin pasar por ningún splitter: si una celda de esa fila contiene "
    "un párrafo largo, el chunk resultante no tiene techo de tamaño."
)
doc.add_paragraph(
    "Confirmado con datos reales — Curso_1472, \"Sistema de notación dentaria.pdf\": chunk 27 "
    "(página 3) con 5,049 caracteres y chunk 13 (página 2) con 5,349 caracteres, ambos "
    "chunk_type='table_row' — 5.6× y 5.9× el CHUNK_SIZE configurado (900 caracteres, ver .env). "
    "Un chunk tan por encima del tamaño objetivo del modelo de embeddings puede perder "
    "precisión semántica y, más importante, reintroduce exactamente el problema que la "
    "separación por fila fue diseñada para evitar (ver docstring de chunk_structured_blocks): "
    "diluir la relevancia de un dato puntual dentro de un bloque de texto grande."
)

# ── Hallazgo 4: separación tabla/prosa ───────────────────────────────
doc.add_heading("Hallazgo 4 — Separación tabla/fila/prosa por documento", level=1)
doc.add_paragraph(
    "chunk_structured_blocks() solo aplica a pdf/docx. De los documentos pdf/docx de cada "
    "curso, la mayoría sí produce chunks de tabla y de prosa por separado (evidencia de que la "
    "detección de tablas de pdfplumber funcionó): "
    f"{len(C1472['structured_docs_ok'])}/{len(C1472['structured_docs_ok']) + len(C1472['structured_docs_prose_only']) + len(C1472['structured_docs_table_only'])} "
    f"en Curso_1472, {len(C1541['structured_docs_ok'])}/{len(C1541['structured_docs_ok']) + len(C1541['structured_docs_prose_only']) + len(C1541['structured_docs_table_only'])} "
    "en Curso_1541. El resto quedó como un solo bloque de prosa sin ninguna tabla detectada — "
    "esto NO es automáticamente una falla: varios de estos documentos son artículos/glosarios "
    "sin tablas reales. Se lista para que el usuario confirme cuáles sí deberían tener tabla:"
)
doc.add_paragraph("Curso_1472 — documentos pdf/docx solo-prosa (13):")
add_bullets(doc, C1472["structured_docs_prose_only"])
doc.add_paragraph("Curso_1541 — documentos pdf/docx solo-prosa (4):")
add_bullets(doc, C1541["structured_docs_prose_only"])

# ── Hallazgo 5: campos estructurados ──────────────────────────────────
doc.add_heading("Hallazgo 5 — Índice de Campos Estructurados: cobertura y autoridad", level=1)
doc.add_paragraph(
    f"Curso_1472: {C1472['campos_estructurados']['distinct_fields']} campos distintos, "
    f"{C1472['campos_estructurados']['total_field_records']} filas totales "
    f"({C1472['campos_estructurados']['repeated_template_false']} de tabla única/alta "
    f"autoridad, {C1472['campos_estructurados']['repeated_template_true']} de plantilla "
    f"repetida/baja autoridad)."
)
doc.add_paragraph(
    f"Curso_1541: {C1541['campos_estructurados']['distinct_fields']} campos distintos, "
    f"{C1541['campos_estructurados']['total_field_records']} filas totales "
    f"({C1541['campos_estructurados']['repeated_template_false']} de tabla única/alta "
    f"autoridad, {C1541['campos_estructurados']['repeated_template_true']} de plantilla "
    f"repetida/baja autoridad)."
)
doc.add_paragraph(
    "En ambos cursos el sílabo real aparece con TODAS sus filas marcadas "
    "is_repeated_template=False (evidencia de alta autoridad) — validado directamente contra "
    "coleccion_campos_estructurados:"
)
sil_rows = []
for d in sil_1472:
    sil_rows.append(("Curso_1472", d["filename"], d["rows_unique"], d["rows_repeated_template"]))
for d in sil_1541:
    sil_rows.append(("Curso_1541", d["filename"], d["rows_unique"], d["rows_repeated_template"]))
add_table(doc, ["Curso", "Archivo de sílabo", "Filas alta autoridad", "Filas plantilla repetida"], sil_rows)
doc.add_paragraph(
    "Esto confirma que la jerarquía de evidencia de resolve_field() (campos_estructurados.py) "
    "tiene con qué funcionar en ambos cursos: cuando se pregunta por un campo del sílabo "
    "(docente, modalidad, etc.), existe al menos una fuente de alta autoridad real e inequívoca "
    "para resolverlo."
)

# ── Hallazgo 6: cobertura semanal ─────────────────────────────────────
doc.add_heading("Hallazgo 6 — Cobertura por semana (1-16)", level=1)
doc.add_paragraph(
    f"Curso_1541: 16/16 semanas tienen al menos un chunk vectorizado — cobertura completa."
)
doc.add_paragraph(
    f"Curso_1472: 13/16 semanas — faltan semanas 4, 8 y 16 en el campo section_name de los "
    f"chunks. Semanas presentes: {', '.join(str(w) for w in C1472['weeks_present'])}."
)
doc.add_paragraph(
    "La semana 4 es un caso particular: SÍ tiene contenido vectorizado, pero el nombre real de "
    "la sección en Moodle trae un typo — \"Senana 04 : Odontología peruana en los siglos XX y "
    "XXI...\" en vez de \"Semana 04\" — verificado directamente en el payload section_name de "
    "los chunks de esa semana. Cualquier código (presente o futuro) que busque section_name por "
    "el patrón \"Semana N\" no la va a encontrar; el typo está en el dato de origen (Moodle), no "
    "en el pipeline de vectorización. Las semanas 8 y 16 no tienen ningún section_name asociado "
    "en absoluto — no hay chunk cuyo payload mencione esas semanas."
)

# ── Hallazgo 7: payload legado sin chunk_type ─────────────────────────
doc.add_heading("Hallazgo 7 (menor) — Chunks con payload de una versión anterior del pipeline", level=1)
doc.add_paragraph(
    f"El campo chunk_type está ausente por completo (no solo vacío) en "
    f"{C1472['missing_chunk_type_field_total']} de {C1472['total_chunks']} chunks de "
    f"Curso_1472 ({pct(C1472['missing_chunk_type_field_total'], C1472['total_chunks'])}) y en "
    f"{C1541['missing_chunk_type_field_total']} de {C1541['total_chunks']} de Curso_1541 "
    f"({pct(C1541['missing_chunk_type_field_total'], C1541['total_chunks'])}). El código actual "
    "de HybridVectorizationService.process_and_store() siempre escribe chunk_type (con "
    "default 'prose' si no aplica chunking estructurado) — su ausencia total en estos puntos "
    "indica que fueron vectorizados con una versión más antigua del pipeline, antes de que ese "
    "campo existiera, y nunca se re-vectorizaron con el código actual."
)
doc.add_paragraph(
    "No rompe nada hoy porque todo el código que lee chunk_type lo hace con manejo explícito de "
    "ausencia — pero es la diferencia real entre 1472 (38.3%, afecta sobre todo a los 37 videos, "
    "18 páginas web y 9 pptx del curso) y 1541 (1.5%, prácticamente marginal) — 1472 tiene una "
    "porción más grande de su contenido corriendo con metadata desactualizada."
)

# ── Recomendaciones ────────────────────────────────────────────────
doc.add_heading("Recomendaciones priorizadas", level=1)
recs = [
    "PRIORIDAD ALTA — Acotar cuándo dispara el fallback de OCR de fórmulas "
    "(_find_formula_regions/_ocr_page_region en definitions.py): hoy alucina texto "
    "pseudo-LaTeX sobre contenido que no es fórmula (visto en tablas, diagramas y elementos de "
    "plantilla en AMBOS cursos) y en al menos un caso confirmado destruyó una tabla de datos "
    "real. Vale la pena: (a) un filtro de sanidad sobre la salida de pix2tex antes de insertarla "
    "(rechazar si contiene patrones como '[object Object]', o una proporción alta de tokens "
    "'mathrm{palabra_larga_sin_sentido}'), y (b) revisar si el área mínima/criterio de "
    "_find_formula_regions necesita subir más, ya que el ajuste hecho para \"Precálculo\" "
    "(según el comentario ya existente en el código) no generaliza a Odontología ni Filosofía.",
    "PRIORIDAD MEDIA — Aplicar un splitter con techo de tamaño a los chunks chunk_type="
    "'table_row' en chunk_structured_blocks() (definitions.py), igual que ya se hace para "
    "'table' y 'prose' — evita chunks de 5,000+ caracteres cuando una celda de tabla trae un "
    "párrafo largo.",
    "PRIORIDAD MEDIA — Ejecutar sync_curso_archivos('Curso_1472', ...) y "
    "sync_curso_archivos('Curso_1541', ...) al menos una vez para poblar el ledger — hoy no hay "
    "forma automática de detectar altas/cambios/bajas de archivos para estos dos cursos.",
    "PRIORIDAD BAJA — Corregir en Moodle el typo \"Senana 04\" → \"Semana 04\" en la sección de "
    "Curso_1472 (fuera del alcance del código de este repositorio, es dato de origen), y "
    "confirmar si las semanas 8 y 16 de Curso_1472 deberían tener material propio (hoy no "
    "tienen ningún archivo vectorizado con esa semana).",
    "PRIORIDAD BAJA — Revisar manualmente los 13 documentos de Curso_1472 y 4 de Curso_1541 que "
    "quedaron solo-prosa (sin chunk de tabla) para confirmar si genuinamente no tienen tablas o "
    "si pdfplumber falló en detectarlas (ver lista completa en el Hallazgo 4).",
    "OPCIONAL — Re-vectorizar los documentos de Curso_1472 cuyo payload no tiene chunk_type "
    "(sobre todo los 37 videos y 18 páginas web) para que queden con la metadata completa del "
    "pipeline actual, si en algún momento se necesita filtrar/reportar por chunk_type sobre el "
    "100% real del curso.",
]
add_bullets(doc, recs)

# ── Anexo ───────────────────────────────────────────────────────────
doc.add_heading("Anexo — Trazabilidad", level=1)
doc.add_paragraph(
    "Script de evaluación: evaluate_vectorization_quality.py (raíz del proyecto). Resultados "
    "completos: evaluate_vectorization_results.json (incluye, por documento: chunks, tipos de "
    "chunk, longitudes, banderas de calidad con ejemplos de texto real, semanas y secciones). "
    "Ambos archivos quedan en el repositorio junto a este informe para que cualquier hallazgo "
    "pueda verificarse de nuevo corriendo el script contra el mismo mirror local."
)

doc.save(str(OUT))
print(f"Guardado: {OUT}")
