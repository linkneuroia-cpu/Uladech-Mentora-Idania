"""
generar_analisis_docx.py
=======================
Genera un documento Word (.docx) con el análisis funcional y técnico del
proyecto "Sistema de Vectorización Documental v2 / Gestor RAG ULADECH".

Ejecutar:  python generar_analisis_docx.py
Salida:    Analisis_Proyecto_Gestor_RAG.docx
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


AZUL = RGBColor(0x1F, 0x3B, 0x73)
AZUL_CLARO = RGBColor(0x2E, 0x5C, 0xA8)
GRIS = RGBColor(0x44, 0x44, 0x44)


def set_heading_style(doc):
    for i in range(1, 4):
        try:
            style = doc.styles[f"Heading {i}"]
            style.font.color.rgb = AZUL
            style.font.name = "Calibri"
        except KeyError:
            pass


def add_bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def add_numbered(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Number")


def add_kv_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9.5)
    return table


def build():
    doc = Document()
    set_heading_style(doc)

    # Estilo base del cuerpo
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # ── Portada ────────────────────────────────────────────────────
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Sistema de Vectorización Documental v2")
    r.bold = True
    r.font.size = Pt(26)
    r.font.color.rgb = AZUL

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Gestor RAG y Asistentes Académicos Inteligentes (ULADECH Católica)")
    r.font.size = Pt(14)
    r.font.color.rgb = AZUL_CLARO

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("Documento de Análisis del Proyecto\nFuncionalidad 1 — Profundización")
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = GRIS

    doc.add_paragraph()
    doc.add_paragraph(
        "Este documento describe, en lenguaje funcional y técnico, de qué trata el "
        "proyecto, qué problema resuelve, cómo está construido y cuáles son sus "
        "componentes principales. La información se obtuvo directamente del código "
        "fuente, la configuración, el esquema de base de datos y los scripts de "
        "evaluación incluidos en el repositorio."
    )

    doc.add_page_break()

    # ── 1. Resumen ejecutivo ───────────────────────────────────────
    doc.add_heading("1. Resumen Ejecutivo", level=1)
    doc.add_paragraph(
        "El proyecto es una plataforma de inteligencia artificial aplicada a la "
        "educación. Su propósito principal es transformar todo el material "
        "documental de un curso (PDF, DOCX, PPTX, XLSX, audio, video y páginas web) "
        "en un asistente académico conversacional capaz de responder preguntas de "
        "los estudiantes citando las fuentes reales del curso."
    )
    doc.add_paragraph(
        "Técnicamente es un sistema de Generación Aumentada por Recuperación (RAG — "
        "Retrieval-Augmented Generation) construido sobre FastAPI, Qdrant (base de "
        "datos vectorial) y PostgreSQL, con integración directa al Campus Virtual "
        "(Moodle) de la universidad. Todo el pipeline RAG es configurable mediante "
        "el patrón de diseño Strategy, lo que permite intercambiar motores de "
        "extracción, embeddings, reranking y generación sin tocar el código."
    )
    doc.add_paragraph(
        "El sistema expone tres mecanismos de vectorización (Semiautomático, "
        "Automático y de Actualización), un gestor web con login, una API REST "
        "documentada, y una ruta pública de chat (/asistente/{token}) que los "
        "estudiantes abren directamente desde Moodle."
    )

    # ── 2. Problema que resuelve ───────────────────────────────────
    doc.add_heading("2. Problema que Resuelve", level=1)
    doc.add_paragraph("El proyecto ataca varios problemas concretos del proceso de enseñanza-aprendizaje:")
    add_bullets(doc, [
        "Dispersión del material: los contenidos de un curso viven en decenas de "
        "archivos heterogéneos (PDFs, diapositivas, hojas de cálculo, videos, "
        "enlaces web, tablas del aula virtual), difíciles de consultar de forma ágil.",
        "Falta de respuestas confiables: un chatbot genérico 'alucina' (inventa) "
        "respuestas. Este sistema obliga al modelo a responder solo con base en el "
        "material real del curso y a citar la fuente (archivo, diapositiva, minuto, fragmento).",
        "Preguntas de datos exactos: preguntas como '¿cuántas semanas tiene la unidad 2?' "
        "o '¿quién es el docente?' se resuelven con motores deterministas (índice de "
        "campos estructurados y Structured Query Engine), no con el LLM, para evitar conteos erróneos.",
        "Navegación por el curso: el asistente indica enlaces reales a tareas, "
        "cuestionarios y recursos usando el mapa del aula de Moodle (Course Map).",
        "Trabajo manual de indexación: en lugar de cargar archivos a mano, el sistema "
        "detecta altas, modificaciones y bajas de archivos en Moodle y sincroniza automáticamente.",
    ])

    # ── 3. Visión funcional ────────────────────────────────────────
    doc.add_heading("3. Visión Funcional: Tres Asistentes", level=1)
    doc.add_paragraph(
        "La aplicación FastAPI concentra en un solo puerto (8100) tres flujos de "
        "trabajo, además de endpoints compartidos:"
    )

    doc.add_heading("3.1 Asistente Semiautomático (/api/semi/*)", level=2)
    doc.add_paragraph(
        "Flujo guiado en tres pasos: el usuario sube archivos, el sistema los "
        "guarda y valida (estructura, tamaño, extensión, compatibilidad de "
        "dimensiones con la colección), y solo tras la confirmación del usuario se "
        "vectorizan en segundo plano. El estado del trabajo se sigue por un job_id."
    )

    doc.add_heading("3.2 Asistente Automático (/api/auto/*)", level=2)
    doc.add_paragraph(
        "Se conecta directamente a Moodle con un curid (ID del curso). Descarga los "
        "PDFs/recursos de los módulos seleccionados, marca los que ya estaban "
        "vectorizados (duplicados) y, tras la confirmación, los indexa. Es el flujo "
        "pensado para cargar un curso completo desde el aula virtual."
    )

    doc.add_heading("3.3 Asistente de Actualización (/api/update/*)", level=2)
    doc.add_paragraph(
        "Reemplaza un documento ya vectorizado por una versión nueva, provista de "
        "forma manual (subida de archivo) o descargada desde Moodle. Elimina el "
        "documento anterior de la colección y vectoriza el nuevo."
    )

    doc.add_heading("3.4 Chat Público del Estudiante (/asistente/{token})", level=2)
    doc.add_paragraph(
        "Cada asistente configurado tiene un token y una URL pública con la que el "
        "estudiante conversa desde el navegador, sin necesidad de cuenta en el gestor. "
        "El estudiante se identifica con su userid de Moodle, o bien con documento de "
        "identidad y nombre (para visitantes sin cuenta). La interfaz incluye:"
    )
    add_bullets(doc, [
        "Historial de conversaciones con carpetas, opción de renombrar y reabrir charlas previas.",
        "Mensaje de bienvenida y mensaje de reencuentro configurables.",
        "Límite de preguntas por conversación (max_actividades_sesion).",
        "Botones de calificación de utilidad ('me ayudó' / 'no me ayudó') con comentario opcional.",
        "Citas de las fuentes usadas en cada respuesta.",
    ])

    # ── 4. Arquitectura técnica ────────────────────────────────────
    doc.add_heading("4. Arquitectura Técnica", level=1)
    doc.add_paragraph(
        "El sistema sigue una arquitectura por capas, con separación clara entre "
        "web, lógica de negocio, persistencia e integración con servicios externos."
    )

    doc.add_heading("4.1 Stack Tecnológico", level=2)
    add_kv_table(
        doc,
        ["Capa", "Tecnología"],
        [
            ["Backend / API", "FastAPI 0.115, Uvicorn, Pydantic"],
            ["Interfaz web", "Jinja2 (plantillas HTML) + JavaScript en el navegador"],
            ["Base de datos vectorial", "Qdrant (vector) 1.16, esquema dense + sparse (híbrido)"],
            ["Base de datos relacional", "PostgreSQL (psycopg2)"],
            ["Embeddings", "Sentence-Transformers, Azure/OpenAI, Jina v3, E5-large"],
            ["Extracción documental", "PyMuPDF, pdfplumber, python-docx, python-pptx, openpyxl"],
            ["OCR / Fórmulas", "EasyOCR, pix2tex (LaTeX-OCR)"],
            ["Audio / Video", "faster-whisper, yt-dlp, youtube-transcript-api"],
            ["Reranking / Generación", "BGE local, Cohere, GPT-4o, Gemini, DeepSeek"],
            ["Integración LMS", "Moodle Web Services (vía token)"],
            ["Despliegue", "Docker / docker-compose (CPU-only), nginx (proxy inverso)"],
        ],
    )
    doc.add_paragraph()

    doc.add_heading("4.2 Componentes de Código Principales", level=2)
    add_kv_table(
        doc,
        ["Archivo", "Responsabilidad"],
        [
            ["app.py", "Aplicación FastAPI: todas las rutas de API, login, chat público y arranque."],
            ["settings.py", "Configuración central, catálogo de modelos y estrategias, validación de entorno."],
            ["core.py", "Núcleo de los 3 mecanismos (SemiAutoCore, AutoCore, UpdateCore) y el store de jobs."],
            ["definitions.py", "Extracción/conversión de documentos, OCR, tablas a Markdown, servicio de vectorización."],
            ["rag_pipeline.py", "Orquestación de la lectura: recuperación híbrida, reranking, boosts y generación."],
            ["strategies/", "Implementaciones intercambiables por etapa (etl_document, etl_audio, contextual, dense, sparse, rerank, generation)."],
            ["campos_estructurados.py", "Índice de campos (campo→valor) y Structured Query Engine (COUNT/MIN/MAX deterministas)."],
            ["query_router.py", "Router de reglas deterministas (regex) para navegación, conteos y progreso del estudiante."],
            ["course_map.py", "Mapa real del curso desde Moodle (secciones, actividades, fechas, semana actual)."],
            ["moodle_tools.py", "Búsqueda de actividades, progreso real del alumno e integración con Moodle."],
            ["assistants.py", "CRUD de asistentes, sesiones, mensajes, carpetas y calificaciones."],
            ["rds.py", "Gestión de RD (aulas), colecciones y embeddings maestros."],
            ["qdrant_admin.py", "Administración de colecciones en Qdrant (crear, borrar, inspeccionar, documentos)."],
            ["ledger.py", "Detección de altas/modificaciones/bajas de archivos del curso (cola de indexación)."],
            ["migracion.py", "Scheduler de re-sincronización automática con Moodle (modo periódico)."],
            ["auth.py / credentials.py", "Login del gestor y gestión segura de API keys."],
            ["db.py / db_schema.sql", "Acceso a PostgreSQL y modelo de datos."],
        ],
    )
    doc.add_paragraph()

    doc.add_heading("4.3 Estrategias Configurables (Strategy Pattern)", level=2)
    doc.add_paragraph(
        "Todo el pipeline RAG se compone de etapas independientes que pueden "
        "intercambiarse en caliente desde el gestor o fijarse por asistente. Cada "
        "etapa tiene su catálogo de opciones, y muchas requieren una clave de API:"
    )
    add_kv_table(
        doc,
        ["Etapa", "Opciones disponibles"],
        [
            ["Extracción doc. (etl_document)", "local, gemini_vision, gpt4o_mini_vision, deepseek_vision"],
            ["Transcripción audio (etl_audio)", "faster_whisper_local, whisper_api, deepgram"],
            ["Contextual retrieval", "none, deepseek, gpt4o_mini"],
            ["Embedding denso (dense)", "text-embedding-3-small, e5-large, jina-v3, MiniLM, openai-3-small"],
            ["Embedding disperso (sparse)", "bm25, bge_m3"],
            ["Reranking", "none, bge_local, cohere"],
            ["Generación", "gpt4o, gpt4o_mini, gemini, deepseek"],
        ],
    )
    doc.add_paragraph()

    # ── 5. Modelo de datos ─────────────────────────────────────────
    doc.add_heading("5. Modelo de Datos (PostgreSQL)", level=1)
    doc.add_paragraph(
        "El esquema relacional organiza la configuración del gestor, los asistentes "
        "y toda la analítica de uso. Sus tablas principales son:"
    )
    add_kv_table(
        doc,
        ["Tabla", "Contenido"],
        [
            ["usuarios", "Cuentas de acceso al gestor (admin y no admin, con hash de contraseña)."],
            ["rds", "Aulas RD (Responsabilidad Docente) con su embedding maestro fijado."],
            ["colecciones", "Vínculo entre courseid de Moodle, RD y colección de Qdrant; modo de migración."],
            ["asistentes", "Definición de cada chatbot público, su token, prompt y config RAG propia."],
            ["asistente_sesiones", "Una conversación = (asistente, curso, usuario Moodle)."],
            ["asistente_mensajes", "Historial completo de preguntas/respuestas, fuentes, tokens, tiempos y calificaciones."],
            ["asistente_carpetas", "Agrupación de conversaciones para la sidebar de historial."],
            ["curso_archivos", "Cola de indexación: un registro por archivo visto en el LMS."],
            ["configuracion", "Configuración mutable clave/valor y secretos (API keys)."],
        ],
    )
    doc.add_paragraph()
    doc.add_paragraph(
        "El esquema se versiona mediante archivos db_migration_*.sql y se siembra con "
        "las 14 RD iniciales del programa, cada una asociada a su courseid de Moodle."
    )

    # ── 6. Flujo de una consulta ───────────────────────────────────
    doc.add_heading("6. Cómo Responde el Asistente (Pipeline RAG)", level=1)
    doc.add_paragraph(
        "Cuando un estudiante envía una pregunta (answer_query en rag_pipeline.py), "
        "el sistema ejecuta esta secuencia:"
    )
    add_numbered(doc, [
        "Recuperación híbrida: se embebe la consulta con los modelos denso y "
        "disperso en paralelo y Qdrant fusiona ambos resultados (RRF).",
        "Reranking: un cross-encoder (o Cohere) reordena los candidatos por relevancia real.",
        "Boosts: se prioriza el material propio del curso sobre la bibliografía y "
        "se refuerza la coincidencia de semana/sesión mencionada en la pregunta.",
        "Evidencia estructurada: si la pregunta apunta a un dato de tabla (docente, "
        "créditos, ponderación), se inyecta el valor exacto desde el índice de campos.",
        "Query Router: reglas deterministas resuelven conteos de unidad/semana, "
        "navegación a actividades, fechas, horario, semana actual y progreso del alumno.",
        "Filtro de relevancia: se descartan candidatos con relevancia casi nula para "
        "evitar que el modelo responda con material irrelevante.",
        "Generación: el LLM produce la respuesta con un system prompt fijo que le "
        "exige basarse en las fuentes y citarlas, además del prompt maestro del asistente.",
        "Verificación de fuente: si el estudiante nombró un documento puntual "
        "(ej. 'el reglamento'), el sistema verifica por nombre de archivo real si ese documento existe.",
    ])

    # ── 7. Evaluación y pruebas ────────────────────────────────────
    doc.add_heading("7. Evaluación, Pruebas y Calidad", level=1)
    doc.add_paragraph(
        "El repositorio incluye un conjunto robusto de scripts y datasets de "
        "evaluación, lo que evidencia un enfoque de mejora continua basado en "
        "mediciones reales en lugar de intuiciones:"
    )
    add_bullets(doc, [
        "Datasets de 200 preguntas (v2, hallucinations y mentor) más datasets "
        "específicos por curso (1472, 1541) y adversariales.",
        "Comparativas antes/después de cambios (archivos *_PRE_3.7, *_BEFORE_fix).",
        "Evaluación de calidad de vectorización y de alucinaciones.",
        "Pruebas QA finales de formato y de contenido.",
        "Segundo asistente 'mentor' evaluado en paralelo.",
        "Resultados en JSON y CSV (precisión, fuentes, tiempos, etc.).",
    ])
    doc.add_paragraph(
        "Este andamiaje de evaluación es uno de los rasgos más destacables del "
        "proyecto: cada mejora del pipeline se valida contra un banco de preguntas "
        "reproducible antes de darse por buena."
    )

    # ── 8. Seguridad ───────────────────────────────────────────────
    doc.add_heading("8. Seguridad y Control de Acceso", level=1)
    add_bullets(doc, [
        "Login con sesión (SessionMiddleware) y middleware de control que protege "
        "todas las rutas salvo login, estáticos, chat público y documentación.",
        "Contraseñas hasheadas (pbkdf2_hmac + salt).",
        "Endpoints administrativos restringidos por require_admin.",
        "Las API keys se guardan de forma segura y nunca se exponen completas "
        "(solo valor enmascarado) al consultarlas.",
        "Aislamiento de datos del estudiante: el progreso (PROGRESS) siempre usa el "
        "userid de la propia sesión, nunca uno provisto por el cliente.",
        "Separación de entornos DESARROLLO/PRODUCCIÓN para PostgreSQL y Qdrant.",
    ])

    # ── 9. Funcionalidades destacadas ──────────────────────────────
    doc.add_heading("9. Funcionalidades Destacadas", level=1)
    add_bullets(doc, [
        "Ingesta multimodal: PDF, DOCX, PPTX, XLSX, audio/video (con transcripción), "
        "video de YouTube y scraping de páginas web.",
        "Conversión de tablas a Markdown y transcripción de fórmulas con OCR especializado.",
        "Migración periódica: re-sincronización automática con Moodle según intervalo configurable.",
        "Course Map y sección General: lee la estructura real del aula y campos administrativos.",
        "Structured Query Engine: conteos y rangos deterministas sobre unidad/semana.",
        "Índice de campos estructurados con detección genérica de campos y fallback semántico.",
        "Calificación y analítica por curso: mide la utilidad percibida de cada asistente.",
        "Despliegue contenedorizado y proxy inverso compartido con otras funcionalidades.",
    ])

    # ── 10. Conclusión ─────────────────────────────────────────────
    doc.add_heading("10. Conclusión", level=1)
    doc.add_paragraph(
        "Se trata de un proyecto maduro de IA educativa: un gestor RAG completo, "
        "configurable y auditable, que convierte el material disperso de los cursos "
        "de Moodle en asistentes académicos confiables, con citas de fuentes, "
        "respuestas deterministas para datos exactos y un fuerte aparato de "
        "evaluación. Su diseño modular (FastAPI + Strategy Pattern + Qdrant + "
        "PostgreSQL) prioriza la calidad de la respuesta (evitar alucinaciones) y "
        "la mantenibilidad, dejando una base sólida para escalar a más cursos, "
        "asistentes y estrategias de IA sin reescribir el sistema."
    )

    doc.save("Analisis_Proyecto_Gestor_RAG.docx")
    print("Documento generado: Analisis_Proyecto_Gestor_RAG.docx")


if __name__ == "__main__":
    build()
