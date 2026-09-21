"""
definitions.py
==============
Modelos Pydantic, VectorizationService y todas las funciones
compartidas por los tres mecanismos de vectorización: semiautomático,
automático y de actualización.
"""

# ──────────────────────────── IMPORTS ────────────────────────────
import re
import io
import hashlib
import unicodedata
import asyncio
import uuid
import requests
from pathlib import Path
from urllib.parse import urlparse
from datetime import datetime
from typing import List, Dict, Any, Optional, Tuple
from settings import settings, get_model_info, get_model_dimensions, uses_prefix, get_ocr_languages

from pydantic import BaseModel

# Procesamiento de documentos
import fitz          # PyMuPDF
import pdfplumber
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import openpyxl
import struct
import olefile

# Audio/video e imagenes
from faster_whisper import WhisperModel
import easyocr
import yt_dlp

# Vectorización
from sentence_transformers import SentenceTransformer
from qdrant_client import QdrantClient
from openai import AzureOpenAI, OpenAI
from qdrant_client.models import (
    Distance, VectorParams, PointStruct,
    Filter, FieldCondition, MatchValue
)

from langchain_text_splitters import RecursiveCharacterTextSplitter


# ═══════════════════════════════════════════════════════════════════
# MODELOS PYDANTIC
# ═══════════════════════════════════════════════════════════════════

class JobStatus(BaseModel):
    job_id: str
    status: str
    message: Optional[str] = None
    result: Optional[Dict[str, Any]] = None


class DocumentInfo(BaseModel):
    """Info de un archivo listado desde Moodle o desde colección."""
    filename: str
    filesize: Optional[int] = None
    fileurl: Optional[str] = None
    course_name: Optional[str] = None
    file_type: str = "other"
    is_duplicate: bool = False
    existing_chunks: int = 0
    existing_date: Optional[str] = None
    existing_source_type: Optional[str] = None
    moodle_upload_date: Optional[str] = None


class ValidationSummary(BaseModel):
    """Resumen de validación estructural de un documento."""
    filename: str
    file_type: str = "other"
    total_pages: int
    total_images: int
    total_tables: int
    duration_seconds: Optional[float] = None
    has_non_textual: bool
    warning_message: Optional[str] = None
    is_duplicate: bool = False
    existing_chunks: int = 0


class SemiUploadRequest(BaseModel):
    collection_name: str


class AutoPreviewRequest(BaseModel):
    collection_name: str
    curid: int


class AutoVectorizeRequest(BaseModel):
    collection_name: str
    curid: int
    selected_filenames: List[str]
    source_type: Optional[str] = None
    file_source_types: Optional[Dict[str, str]] = None


class UpdateListRequest(BaseModel):
    collection_name: str


class UpdateVectorizeRequest(BaseModel):
    collection_name: str
    filename_to_replace: str
    # Modo Moodle:
    curid: Optional[int] = None
    moodle_filename: Optional[str] = None
    course_name: Optional[str] = None


class ConfirmVectorizeRequest(BaseModel):
    job_id: str


class PipelineConfig(BaseModel):
    """Configuración del pipeline RAG (Strategy Pattern) para una
    vectorización contra una colección híbrida (dense+sparse). Se construye
    automáticamente en core.py a partir de la configuración activa
    (strategies.registry.get_all_active()) — el usuario no arma esto
    manualmente. Ausente (None) para colecciones legacy, que siguen usando
    VectorizationService sin cambios."""
    etl_document_strategy: Optional[str] = None
    etl_audio_strategy: Optional[str] = None
    contextual_strategy: Optional[str] = None
    dense_strategy: str
    sparse_strategy: str
    source_type: str = "curso_propio"


# ═══════════════════════════════════════════════════════════════════
# UTILIDADES DE TEXTO
# ═══════════════════════════════════════════════════════════════════

def _generate_hash(text: str) -> str:
    normalized = re.sub(r"\s+", " ", text.strip().lower())
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _sanitize_filename(name: str) -> str:
    """Sanitizar nombre para el sistema de archivos."""
    safe = "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789-_."
    result = "".join(c if c in safe else "_" for c in name)
    return result[:100]


def _normalize_filename(name: str) -> str:
    """
    Normalizar filename para búsquedas en Qdrant.
    """
    normalized = unicodedata.normalize("NFD", name.strip())
    normalized = re.sub(r"\s+", " ", normalized)
    return normalized


# ═══════════════════════════════════════════════════════════════════
# CLASIFICACIÓN DE TIPO DE ARCHIVO
# ═══════════════════════════════════════════════════════════════════

_EXTENSION_TYPE_MAP: Dict[str, str] = {
    ".pdf":  "pdf",
    ".docx": "docx",
    ".pptx": "pptx",
    ".ppt":  "ppt",
    ".xlsx": "xlsx",
    ".jpg":  "image", ".jpeg": "image", ".png": "image", ".webp": "image",
    ".mp4":  "video", ".mov": "video", ".webm": "video", ".mkv": "video",
    ".mp3":  "audio", ".wav": "audio", ".m4a": "audio", ".ogg": "audio",
    # Texto ya extraído de una fuente externa (página web scrapeada) — ver
    # scraping.py/get_course_resources(). Ningún archivo real subido pasa
    # por acá: .txt no está en settings.ALLOWED_EXTENSIONS.
    ".txt":  "web",
}


def classify_file_type(filename: str) -> str:
    """
    Clasifica un archivo por su extensión en:
    pdf | docx | pptx | xlsx | image | video | audio | web | other
    """
    ext = Path(filename).suffix.lower()
    return _EXTENSION_TYPE_MAP.get(ext, "other")


# Firmas (magic bytes) para validar descargas por tipo. Tipos no listados
# (ej. "other") no se validan por firma — solo se exige que el archivo no
# esté vacío, ya que puede no tener un magic byte conocido.
_MAGIC_BYTES: Dict[str, List[bytes]] = {
    "pdf":   [b"%PDF-"],
    "docx":  [b"PK\x03\x04"],  # OOXML = ZIP
    "pptx":  [b"PK\x03\x04"],
    "xlsx":  [b"PK\x03\x04"],
    "ppt":   [b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"],  # OLE2 compound file
    "image": [b"\xff\xd8\xff", b"\x89PNG\r\n\x1a\n", b"RIFF"],  # jpeg, png, webp(RIFF)
    "audio": [b"ID3", b"\xff\xfb", b"\xff\xf3", b"RIFF", b"OggS"],
    "video": [b"\x00\x00\x00", b"RIFF"],  # mp4/mov "ftyp" boxes vary; se valida por tamaño no vacío
}


def _validate_file_signature(file_type: str, header: bytes) -> bool:
    """
    Verifica que los primeros bytes del archivo coincidan con alguna firma
    conocida para su tipo. Si el tipo no tiene firmas registradas (o es
    "video", cuyo header varía según el contenedor), se acepta con solo
    que el archivo no esté vacío.
    """
    signatures = _MAGIC_BYTES.get(file_type)
    if not signatures or file_type == "video":
        return len(header) > 0
    return any(header.startswith(sig) for sig in signatures)


def _table_to_markdown(rows: List[List[Optional[str]]]) -> str:
    """Convierte filas de tabla (lista de listas, como las devuelve
    pdfplumber `.extract()` o una tabla de python-docx) a sintaxis Markdown
    GFM — la misma que ya sabe renderizar el chat (ver mdToHtml() en
    frontend/asistente_chat.html)."""
    clean_rows = [
        [(c or "").strip().replace("\n", " ") for c in row]
        for row in rows
    ]
    clean_rows = [r for r in clean_rows if any(c for c in r)]
    if not clean_rows:
        return ""
    header, *body = clean_rows
    width = len(header)
    lines = ["| " + " | ".join(header) + " |", "|" + "|".join(["---"] * width) + "|"]
    for row in body:
        row = (row + [""] * width)[:width]
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


# Un NOMBRE DE CAMPO estructurado real es corto: verificado con datos
# reales, el máximo en las dos colecciones con índice estructurado es 38
# caracteres. Una "clave" más larga que esto es, en la práctica, un
# fragmento de prosa mal detectado como tabla por pdfplumber (Fix A.5,
# Causa 2) — no se vuelca al índice estructurado.
_MAX_FIELD_NAME_LEN = 40


def _split_md_row(line: str) -> List[str]:
    """Celdas no vacías de una línea de tabla Markdown, en orden (comprime
    las celdas de relleno vacías que pdfplumber intercala por fusión de
    celdas)."""
    return [c.strip() for c in line.strip("|").split("|") if c.strip()]


def _markdown_table_header_and_body(md_text: str):
    """(lines, header_cells, body_lines) de una tabla Markdown GFM, o None
    si el texto no tiene forma de tabla (la 2ª línea no es el separador
    GFM esperado)."""
    lines = [ln.strip() for ln in md_text.strip().split("\n") if ln.strip()]
    if len(lines) < 2:
        return None
    if not re.match(r"^\|?[\s:-]+\|?[\s:|-]*$", lines[1]):
        return None
    return lines, _split_md_row(lines[0]), lines[2:]


def _looks_like_prose(text: str) -> bool:
    """Heurística genérica de "esto es prosa, no una fila de tabla real":
    texto largo, con varias oraciones completas, o con palabras cortadas a
    fin de línea (típico de un PDF justificado en columnas que pdfplumber
    confunde con una tabla). Usada para descartar filas de UNA sola celda
    (Fix A.5, Causa 2)."""
    t = (text or "").strip()
    if not t:
        return False
    if len(t) > 200:
        return True
    # Oración completa: minúscula + punto + espacio + mayúscula (exigir la
    # minúscula evita falsos positivos con etiquetas numeradas "0. Programa").
    if len(re.findall(r"[a-záéíóúñ]\.\s+[A-ZÁÉÍÓÚÑ]", t)) >= 2:
        return True
    if re.search(r"[a-záéíóúñ]- [a-záéíóúñ]", t):
        return True
    return False


def _looks_like_prose_key(key: str) -> bool:
    """True si `key` no puede ser un NOMBRE DE CAMPO real (es prosa mal
    detectada como clave): más largo que _MAX_FIELD_NAME_LEN, con palabras
    cortadas a fin de línea, o con oraciones completas."""
    k = (key or "").strip()
    if len(k) > _MAX_FIELD_NAME_LEN:
        return True
    if re.search(r"[a-záéíóúñ]- [a-záéíóúñ]", k):
        return True
    # Oración completa (minúscula + punto + mayúscula) — no confundir con
    # etiquetas numeradas tipo "0. Programa" o "9.1.1 ...".
    if re.search(r"[a-záéíóúñ]\.\s+[A-ZÁÉÍÓÚÑ]", k):
        return True
    return False


def _parse_markdown_table_rows(
    md_text: str, override_header: Optional[List[str]] = None,
) -> List[Dict[str, str]]:
    """Parsea una tabla Markdown GFM (la misma sintaxis que produce
    _table_to_markdown) de vuelta a filas {encabezado: valor} — genérico
    para cualquier tabla real, sin asumir nada sobre el significado de las
    columnas.

    Dos particularidades reales de pdfplumber que hay que tolerar (vistas
    en datos reales, no hipotéticas):
    1. Celdas de relleno vacías intercaladas por fusión de celdas del PDF
       original (ej. una fila real de 4 valores puede llegar como 12
       celdas, con 8 vacías en posiciones que varían de fila a fila) — por
       eso NUNCA se empareja por posición cruda; cada fila se comprime a
       sus celdas no vacías, en orden, antes de emparejar.
    2. Una tabla de 2 columnas suele ser una lista campo/valor SIN fila de
       encabezado real (la "primera fila" ya es un dato como cualquier
       otra, ej. "0. Programa de estudio | Odontología") — tratarla como
       encabezado descartaría ese primer dato y arruinaría el resto. Por
       eso el criterio de si hay encabezado es el número de columnas
       reales tras comprimir, no la posición de la fila.

    `override_header` (Fix A.5, Causa 1): encabezado real de la tabla
    inmediatamente anterior del mismo documento, usado para parsear un
    fragmento de tabla cuyo propio encabezado quedó roto por un salto de
    página. Cuando se pasa, la fila 0 del fragmento se descarta (es el
    encabezado roto) y las filas se emparejan contra ese encabezado."""
    parts = _markdown_table_header_and_body(md_text)
    if parts is None:
        return []
    lines, header_cells, body_lines = parts

    if override_header is not None:
        header_cells = list(override_header)
    elif len(header_cells) == 1:
        # Fix A.5-1 (seguridad): un header comprimido a 1 celda pero con
        # filas de cuerpo de más de 1 celda NO es una tabla de una columna
        # genuina — es un header ROTO (tabla partida por salto de página).
        # No adivinar una estructura falsa: sin filas.
        if any(len(_split_md_row(l)) > 1 for l in body_lines):
            return []

    if len(header_cells) == 1:
        # Tabla de una sola columna real — cada línea (incluida la
        # "header") ya es un dato autocontenido, sin key:value que armar.
        # Fix A.5-3: descarta párrafos de prosa mal detectados como tabla
        # (no los vuelca al índice estructurado).
        rows = []
        first = header_cells[0]
        if first and not _looks_like_prose(first):
            rows.append({"": first})
        for line in body_lines:
            cells = _split_md_row(line)
            if not cells:
                continue
            joined = " ".join(cells)
            if not _looks_like_prose(joined):
                rows.append({"": joined})
        return rows

    if len(header_cells) == 2 and override_header is None:
        # Lista campo/valor sin encabezado real — la fila 0 es un dato más.
        # Fix A.5-3: descarta filas cuya "clave" es prosa, no un campo real.
        rows = []
        for line in [lines[0]] + body_lines:
            cells = _split_md_row(line)
            if len(cells) == 2 and not _looks_like_prose_key(cells[0]):
                rows.append({cells[0]: cells[1]})
        return rows

    # 3+ columnas reales (o header provisto como override): la fila 0 (o el
    # override) son encabezados genuinos. Fix A.5-3: si algún encabezado es
    # prosa, la tabla entera es una mala detección — no se generan filas.
    if any(_looks_like_prose_key(h) for h in header_cells):
        return []
    rows = []
    for line in body_lines:
        cells = _split_md_row(line)
        if len(cells) != len(header_cells):
            continue  # fila con cantidad de columnas reales distinta — se omite, no se adivina
        row = {h: c for h, c in zip(header_cells, cells)}
        if row:
            rows.append(row)
    return rows


def _group_words_into_blocks(words: List[Dict[str, Any]]) -> List[Tuple[float, str]]:
    """Agrupa palabras (ya excluidas de tablas/encabezado/pie) en bloques de
    texto por proximidad vertical, conservando la posición ("top") de cada
    bloque — necesaria para poder intercalarlos con tablas y fórmulas en el
    orden real de lectura de la página (antes se unían todas las palabras
    de la página en un solo string, perdiendo esa posición)."""
    if not words:
        return []
    ordered = sorted(words, key=lambda w: (w["top"], w["x0"]))
    blocks: List[Tuple[float, List[str]]] = []
    block_top = ordered[0]["top"]
    block_words = [ordered[0]["text"]]
    prev_bottom = ordered[0]["bottom"]
    for w in ordered[1:]:
        gap = w["top"] - prev_bottom
        if gap > 8:  # salto vertical grande -> nuevo párrafo/bloque
            blocks.append((block_top, block_words))
            block_top = w["top"]
            block_words = []
        block_words.append(w["text"])
        prev_bottom = max(prev_bottom, w["bottom"])
    blocks.append((block_top, block_words))
    return [(top, " ".join(ws)) for top, ws in blocks]


def _merge_nearby_rects(
    rects: List[Tuple[float, float, float, float]], gap: float = 10
) -> List[Tuple[float, float, float, float]]:
    """Une rectángulos que están a menos de `gap` de distancia entre sí en
    uno solo (unión de bounding boxes) — usado para agrupar trazos
    vectoriales sueltos (glifos de una fórmula, líneas de un diagrama) en
    una sola región candidata a recortar y pasar por OCR."""
    boxes = [list(r) for r in rects]
    changed = True
    while changed:
        changed = False
        out: List[List[float]] = []
        used = [False] * len(boxes)
        for i in range(len(boxes)):
            if used[i]:
                continue
            cur = boxes[i]
            used[i] = True
            for j in range(i + 1, len(boxes)):
                if used[j]:
                    continue
                other = boxes[j]
                if not (cur[2] + gap < other[0] or other[2] + gap < cur[0] or
                        cur[3] + gap < other[1] or other[3] + gap < cur[1]):
                    cur = [min(cur[0], other[0]), min(cur[1], other[1]),
                           max(cur[2], other[2]), max(cur[3], other[3])]
                    used[j] = True
                    changed = True
            out.append(cur)
        boxes = out
    return [tuple(b) for b in boxes]


def _find_formula_regions(
    fitz_page,
    table_bboxes: List[Tuple[float, float, float, float]],
    header_bottom: float,
    footer_top: float,
    min_area: float = 2500,
) -> List[Tuple[float, float, float, float]]:
    """Identifica regiones de la página que probablemente contienen una
    fórmula o gráfico dibujado como trazos vectoriales (no como texto
    seleccionable — confirmado que esto ocurre en material real del curso:
    ni pdfplumber ni PyMuPDF recuperan esas fórmulas como texto bajo
    ninguna configuración). Se ignoran los trazos dentro de tablas ya
    detectadas, encabezado y pie de página, y las regiones muy chicas
    (líneas decorativas, viñetas).

    El OCR es un último recurso, no el camino principal — para que dispare
    lo menos posible (medido en Precálculo: 68% de las páginas, hasta 10
    llamadas en una sola), acá se fusionan trazos cercanos con más margen
    (gap=30 en vez de 10) para que fragmentos de UNA misma fórmula
    (subíndices, símbolos sueltos) formen una sola región en vez de varias
    llamadas de OCR separadas, y se subió el área mínima (2500) para
    descartar marcas chicas que no son fórmulas reales."""
    rects = []
    for d in fitz_page.get_drawings():
        r = d.get("rect")
        if r is None:
            continue
        if r.y0 < header_bottom or r.y1 > footer_top:
            continue
        in_table = any(
            not (r.x1 < tb[0] or r.x0 > tb[2] or r.y1 < tb[1] or r.y0 > tb[3])
            for tb in table_bboxes
        )
        if in_table:
            continue
        rects.append((r.x0, r.y0, r.x1, r.y1))
    if not rects:
        return []
    merged = _merge_nearby_rects(rects, gap=30)
    return [m for m in merged if (m[2] - m[0]) * (m[3] - m[1]) >= min_area]


_latex_ocr_model = None

# pix2tex es un modelo generativo sin mecanismo de rechazo propio: ante
# una región que NO es una fórmula real (una tabla sin líneas completas,
# un diagrama, un logo), en vez de fallar devuelve LaTeX con apariencia
# válida pero sin sentido — confirmado con datos reales de ambos cursos
# (ver evaluación de vectorización): tokens "scriptstyle"/"mathrm"
# repetidos decenas de veces, y un artefacto propio del modelo, el token
# literal "[object Object]", insertado como si fuera un argumento LaTeX
# real. Una fórmula real de material de curso (ej. interés compuesto:
# "P(1+\\frac{r}{n})^{nt}") es corta — estos falsos positivos midieron
# cientos a miles de caracteres. Genérico: no depende de qué documento es,
# solo de si el LaTeX generado es estructuralmente plausible.
_MAX_PLAUSIBLE_FORMULA_CHARS = 300
_MAX_TOKEN_REPETITIONS = 6


def _looks_like_plausible_formula(latex: str) -> bool:
    if not latex or len(latex) > _MAX_PLAUSIBLE_FORMULA_CHARS:
        return False
    if "object Object" in latex:
        return False
    token_counts: Dict[str, int] = {}
    for token in re.findall(r"[a-zA-Z]{4,}", latex):
        token_counts[token] = token_counts.get(token, 0) + 1
        if token_counts[token] > _MAX_TOKEN_REPETITIONS:
            return False
    return True


def _get_latex_ocr():
    """LaTeX-OCR (pix2tex): modelo especializado en imagen-de-fórmula ->
    LaTeX real, 100% local (descarga los pesos una sola vez, después corre
    sin red). Verificado contra EasyOCR en material real del curso: sobre
    la misma fórmula de interés compuesto, EasyOCR aplanaba la fracción
    ("P(1+ 5)\" 4.10") mientras pix2tex la reconstruye correctamente
    ("P(1+\\frac{r}{n})^{nt}")."""
    global _latex_ocr_model
    if _latex_ocr_model is None:
        import torch
        from munch import Munch
        from pix2tex.cli import LatexOCR
        # OJO: LatexOCR() sin argumentos usa 'no_cuda': True COMO DEFAULT
        # (ver pix2tex/cli.py LatexOCR.__init__) — corre en CPU aunque haya
        # GPU disponible, aunque no lo diga en ningún lado. Hay que pasar
        # explícitamente los mismos defaults con no_cuda en False (mismos
        # paths relativos que usa pix2tex internamente — los resuelve el
        # decorador @in_model_path() del propio __init__).
        use_gpu = torch.cuda.is_available()
        print(f"🔄 Cargando modelo LaTeX-OCR (pix2tex, gpu={use_gpu})...")
        arguments = Munch({
            "config": "settings/config.yaml",
            "checkpoint": "checkpoints/weights.pth",
            "no_cuda": not use_gpu,
            "no_resize": False,
        })
        _latex_ocr_model = LatexOCR(arguments)
    return _latex_ocr_model


async def _ocr_page_region(fitz_page, bbox: Tuple[float, float, float, float]) -> str:
    """Renderiza una región de la página a imagen y la transcribe en
    cascada: 1) pix2tex (LaTeX-OCR, especializado en fórmulas — preferido),
    2) EasyOCR como respaldo si pix2tex falla (p.ej. no se pudieron
    descargar/cargar sus pesos). Devuelve el bloque ya listo para insertar
    en el texto (fórmula LaTeX entre \\( \\) si fue pix2tex, o un marcador
    explícito de menor confianza si fue EasyOCR) — string vacío si ninguno
    de los dos reconoció nada."""
    loop = asyncio.get_event_loop()

    def _render() -> Path:
        clip = fitz.Rect(*bbox)
        pix = fitz_page.get_pixmap(clip=clip, dpi=200)
        tmp_path = settings.TEMP_DIR / f"_ocr_crop_{uuid.uuid4().hex}.png"
        pix.save(str(tmp_path))
        return tmp_path

    try:
        tmp_path = await loop.run_in_executor(None, _render)
    except Exception as e:
        # Una región con un bounding box degenerado (detectado a veces en
        # PDFs grandes con drawings malformados) puede hacer que MuPDF
        # rechace el pixmap (p.ej. "Invalid bandwriter header dimensions") —
        # sin este try, esa única región abortaba la extracción del
        # documento COMPLETO (ver el except de _extract_from_pdf). Se trata
        # igual que "ningún OCR reconoció nada": se omite esta región.
        print(f"⚠️ No se pudo renderizar una región para OCR, se omite: {e}")
        return ""

    try:
        try:
            from PIL import Image
            latex_model = await loop.run_in_executor(None, _get_latex_ocr)
            img = Image.open(tmp_path)
            latex = await loop.run_in_executor(None, lambda: latex_model(img))
            latex = (latex or "").strip()
            if latex and _looks_like_plausible_formula(latex):
                return f"\\( {latex} \\)"
            if latex:
                print(f"⚠️ pix2tex devolvió LaTeX no plausible (descartado), cae a EasyOCR: {latex[:120]}...")
        except Exception as e:
            print(f"⚠️ pix2tex falló, cae a EasyOCR: {e}")

        reader = await loop.run_in_executor(None, _get_ocr_reader)
        results = await loop.run_in_executor(
            None, lambda: reader.readtext(str(tmp_path), detail=0, paragraph=True)
        )
        text = " ".join(results).strip()
        return (
            f"[contenido gráfico detectado — posible fórmula, transcripción OCR]: {text}"
            if text else ""
        )
    finally:
        tmp_path.unlink(missing_ok=True)


async def _extract_from_pdf_blocks(file_path: str) -> List[Dict[str, Any]]:
    """Como _extract_from_pdf, pero conserva cada bloque por separado con
    su tipo ("table"/"text") y número de página, en vez de aplanarlo a un
    string — lo necesita chunk_structured_blocks() para no partir ni
    mezclar una tabla con el texto vecino (ver docstring de esa función).
    Excluye encabezados y pies de página."""
    doc = fitz.open(file_path)
    all_blocks: List[Dict[str, Any]] = []

    try:
        with pdfplumber.open(file_path) as pdf:
            for page_num, page in enumerate(pdf.pages, 1):
                page_height = page.height
                header_bottom = page_height * 0.10
                footer_top = page_height * 0.90
                fitz_page = doc[page_num - 1]

                tables = page.find_tables(table_settings={
                    "vertical_strategy": "lines",
                    "horizontal_strategy": "lines",
                    "intersection_tolerance": 4,
                })
                table_bboxes = [t.bbox for t in tables]

                # bloques con su posición vertical y tipo, para reconstruir
                # el orden real de lectura de la página al final.
                blocks: List[Tuple[float, str, str]] = []

                for t in tables:
                    md = _table_to_markdown(t.extract())
                    if md:
                        blocks.append((t.bbox[1], md, "table"))

                words = page.extract_words(x_tolerance=1, y_tolerance=1)
                text_words = []
                for word in words:
                    wb = (word["x0"], word["top"], word["x1"], word["bottom"])
                    if word["top"] < header_bottom or word["bottom"] > footer_top:
                        continue
                    in_table = any(
                        not (wb[2] < tb[0] or wb[0] > tb[2] or
                             wb[3] < tb[1] or wb[1] > tb[3])
                        for tb in table_bboxes
                    )
                    if not in_table:
                        text_words.append(word)
                for top, text in _group_words_into_blocks(text_words):
                    blocks.append((top, text, "text"))

                formula_regions = _find_formula_regions(
                    fitz_page, table_bboxes, header_bottom, footer_top
                )
                for region in formula_regions:
                    # _ocr_page_region ya devuelve el bloque formateado
                    # (LaTeX real si fue pix2tex, marcador de menor
                    # confianza si cayó a EasyOCR) — acá solo se cubre el
                    # caso de que ninguno de los dos reconociera nada.
                    content = await _ocr_page_region(fitz_page, region)
                    marker = content or "[hay contenido gráfico en esta sección que el sistema no pudo transcribir]"
                    blocks.append((region[1], marker, "text"))

                blocks.sort(key=lambda b: b[0])
                for _, text, btype in blocks:
                    if text.strip():
                        all_blocks.append({"type": btype, "text": text, "page": page_num})

    except Exception as e:
        raise ValueError(f"Error extrayendo PDF: {e}")
    finally:
        doc.close()

    return all_blocks


async def _extract_from_pdf(file_path: str) -> str:
    """Extrae texto de PDF: párrafos + tablas (convertidas a Markdown, ya
    no descartadas) + fórmulas/gráficos vectoriales (recorte + OCR local),
    intercalados en el orden vertical real de cada página. Excluye
    encabezados y pies de página.

    Compatibilidad: aplana _extract_from_pdf_blocks() a un string plano
    (mismo comportamiento que antes de existir el chunking estructural),
    para los caminos que todavía no lo usan (VectorizationService legacy,
    validate_document, etc.)."""
    blocks = await _extract_from_pdf_blocks(file_path)
    return "\n".join(b["text"] for b in blocks)


def _iter_docx_block_items(doc):
    """Recorre párrafos y tablas del cuerpo del documento en su orden real
    de aparición — python-docx no expone esto directo (`doc.paragraphs` y
    `doc.tables` son dos colecciones separadas, sin orden combinado), así
    que se lee el XML del body en orden y se envuelve cada hijo según su
    tipo (técnica estándar conocida como "iter_block_items")."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


async def _extract_from_docx_blocks(file_path: str) -> List[Dict[str, Any]]:
    """Como _extract_from_docx, pero conserva cada bloque por separado con
    su tipo ("table"/"text") — ver _extract_from_pdf_blocks(). DOCX no
    tiene concepto nativo de "página" (Word solo pagina al renderizar), así
    que `page` queda siempre None para este formato."""
    from docx.table import Table as DocxTable

    doc = DocxDocument(file_path)
    blocks: List[Dict[str, Any]] = []
    for item in _iter_docx_block_items(doc):
        if isinstance(item, DocxTable):
            rows = [[cell.text for cell in row.cells] for row in item.rows]
            md = _table_to_markdown(rows)
            if md:
                blocks.append({"type": "table", "text": md, "page": None})
        else:
            text = item.text.strip()
            if text:
                blocks.append({"type": "text", "text": text, "page": None})
    return blocks


async def _extract_from_docx(file_path: str) -> str:
    """Compatibilidad: aplana _extract_from_docx_blocks() a un string
    (ver _extract_from_pdf)."""
    blocks = await _extract_from_docx_blocks(file_path)
    return "\n\n".join(b["text"] for b in blocks)


async def _extract_from_pptx(file_path: str) -> str:
    """Extrae título, viñetas, tablas y notas del orador de cada slide."""
    prs = Presentation(file_path)
    slides_text = []

    for i, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(c.text.strip() for c in row.cells)
                    if row_text.strip(" |"):
                        parts.append(row_text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"Notas: {notes}")

        if parts:
            slides_text.append(f"[Diapositiva {i}]\n" + "\n".join(parts))

    return "\n\n".join(slides_text)


def _iter_ole_records(data: bytes, start: int, end: int):
    """
    Recorre el arbol de registros binarios del formato PowerPoint 97-2003
    (MS-PPT: cada registro tiene un header de 8 bytes: ver+instance (2B),
    recType (2B), recLen (4B)). recVer==0xF marca un contenedor (se recorre
    recursivamente); cualquier otro valor es un atomo (hoja).
    """
    offset = start
    while offset + 8 <= end:
        ver_instance, rec_type, rec_len = struct.unpack_from("<HHI", data, offset)
        rec_ver = ver_instance & 0x0F
        content_start = offset + 8
        content_end = content_start + rec_len
        if content_end > end:
            break
        yield rec_type, rec_ver, content_start, content_end
        offset = content_end


def _extract_ppt_legacy_text(data: bytes) -> str:
    """
    Extrae el texto de los atomos TextCharsAtom (0x0FA0, UTF-16LE) y
    TextBytesAtom (0x0FA8, un byte por caracter) del stream binario
    "PowerPoint Document" de un .ppt (97-2003).
    """
    texts = []

    def walk(start: int, end: int):
        for rec_type, rec_ver, c_start, c_end in _iter_ole_records(data, start, end):
            if rec_type == 0x0FA0:
                texts.append(data[c_start:c_end].decode("utf-16-le", errors="ignore"))
            elif rec_type == 0x0FA8:
                texts.append(data[c_start:c_end].decode("cp1252", errors="ignore"))
            elif rec_ver == 0x0F:
                walk(c_start, c_end)

    walk(0, len(data))
    cleaned = [t.replace("\x0b", "\n").strip() for t in texts]
    return "\n\n".join(t for t in cleaned if t)


async def _extract_from_ppt_legacy(file_path: str) -> str:
    ole = olefile.OleFileIO(file_path)
    try:
        if not ole.exists("PowerPoint Document"):
            raise ValueError("El archivo .ppt no contiene un stream 'PowerPoint Document' valido")
        data = ole.openstream("PowerPoint Document").read()
    finally:
        ole.close()
    return _extract_ppt_legacy_text(data)


async def _extract_from_xlsx(file_path: str) -> str:
    """Extrae cada hoja como filas de texto plano separadas por ' | '."""
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    sheets_text = []

    for sheet in wb.worksheets:
        rows_text = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(v).strip() for v in row if v is not None and str(v).strip() != ""]
            if values:
                rows_text.append(" | ".join(values))
        if rows_text:
            sheets_text.append(f"[Hoja: {sheet.title}]\n" + "\n".join(rows_text))

    wb.close()
    return "\n\n".join(sheets_text)


# ── OCR (imagenes) ──────────────────────────────────────────────────

_ocr_reader: Optional["easyocr.Reader"] = None


def _get_ocr_reader() -> "easyocr.Reader":
    global _ocr_reader
    if _ocr_reader is None:
        import torch
        use_gpu = torch.cuda.is_available()
        print(f"🔄 Cargando modelo OCR (idiomas: {settings.OCR_LANGUAGES}, gpu={use_gpu})...")
        _ocr_reader = easyocr.Reader(get_ocr_languages(), gpu=use_gpu)
    return _ocr_reader


async def _extract_from_image(file_path: str) -> str:
    loop = asyncio.get_event_loop()
    reader = await loop.run_in_executor(None, _get_ocr_reader)
    results = await loop.run_in_executor(
        None, lambda: reader.readtext(file_path, detail=0, paragraph=True)
    )
    return "\n".join(results)


# ── ASR (audio/video) ────────────────────────────────────────────────

_whisper_model: Optional[WhisperModel] = None


def _get_whisper_model() -> WhisperModel:
    global _whisper_model
    if _whisper_model is None:
        print(
            f"🔄 Cargando modelo Whisper ({settings.ASR_MODEL_SIZE}, "
            f"{settings.ASR_DEVICE}/{settings.ASR_COMPUTE_TYPE})..."
        )
        _whisper_model = WhisperModel(
            settings.ASR_MODEL_SIZE,
            device=settings.ASR_DEVICE,
            compute_type=settings.ASR_COMPUTE_TYPE,
        )
    return _whisper_model


def get_media_duration(file_path: str) -> float:
    """Duración en segundos leyendo solo metadata del contenedor (sin decodificar)."""
    try:
        import av
        with av.open(file_path) as container:
            if container.duration:
                return round(container.duration / 1_000_000, 1)
    except Exception as e:
        print(f"⚠️ No se pudo leer duración de {file_path}: {e}")
    return 0.0


def extract_audio_track(file_path: str) -> Path:
    """
    Extrae la pista de audio de un archivo de video a un .wav mono 16kHz
    temporal, usando PyAV (decodifica y remuxa sin depender del binario
    ffmpeg). Usada por las estrategias ETL de audio que requieren subir un
    archivo de audio puro (APIs externas), a diferencia de faster-whisper
    local que decodifica el video directamente.
    """
    import av

    input_container = av.open(file_path)
    try:
        try:
            input_audio_stream = next(
                s for s in input_container.streams if s.type == "audio"
            )
        except StopIteration:
            raise ValueError(f"El archivo '{file_path}' no tiene pista de audio")

        output_path = settings.TEMP_DIR / f"{Path(file_path).stem}_audio.wav"
        output_container = av.open(str(output_path), mode="w")
        try:
            output_stream = output_container.add_stream("pcm_s16le", rate=16000)
            output_stream.layout = "mono"
            resampler = av.AudioResampler(format="s16", layout="mono", rate=16000)

            for frame in input_container.decode(input_audio_stream):
                for resampled in resampler.resample(frame):
                    for packet in output_stream.encode(resampled):
                        output_container.mux(packet)
            for packet in output_stream.encode(None):
                output_container.mux(packet)
        finally:
            output_container.close()
    finally:
        input_container.close()

    return output_path


async def transcribe_audio_or_video(file_path: str) -> List[Dict[str, Any]]:
    """
    Transcribe audio/video con faster-whisper.
    Retorna una lista de segmentos: [{"start": float, "end": float, "text": str}, ...]
    """
    loop = asyncio.get_event_loop()
    model = await loop.run_in_executor(None, _get_whisper_model)

    def _run():
        # vad_filter=True descarta silencio antes de transcribir — cuando
        # TODO el archivo es silencio (ej. una grabación de pantalla sin
        # narración hablada), es un comportamiento conocido de faster-whisper
        # que su fusión interna de timestamps de voz calcule max() sobre una
        # lista vacía y reviente ANTES de devolver algo. Sin este try/except
        # ese crash llegaba crudo hasta el usuario ("max() arg is an empty
        # sequence") en vez del mensaje ya existente y claro
        # ("Transcripción sin contenido de texto suficiente") que el caller
        # ya produce correctamente para una lista de segmentos vacía.
        try:
            segments_iter, _info = model.transcribe(file_path, vad_filter=True)
            return [
                {"start": s.start, "end": s.end, "text": s.text.strip()}
                for s in segments_iter if s.text.strip()
            ]
        except Exception as e:
            print(f"⚠️ Transcripción falló (posible video/audio sin voz detectable): {e}")
            return []

    return await loop.run_in_executor(None, _run)


async def _extract_text(file_path: str) -> str:
    path = Path(file_path)
    file_type = classify_file_type(path.name)
    if file_type == "pdf":
        return await _extract_from_pdf(file_path)
    elif file_type == "docx":
        return await _extract_from_docx(file_path)
    elif file_type == "pptx":
        return await _extract_from_pptx(file_path)
    elif file_type == "ppt":
        return await _extract_from_ppt_legacy(file_path)
    elif file_type == "xlsx":
        return await _extract_from_xlsx(file_path)
    elif file_type == "image":
        return await _extract_from_image(file_path)
    elif file_type == "web":
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Formato no soportado: {path.suffix}")


async def _normalize_text(text: str) -> str:
    # Whitelist ampliada con operadores/símbolos matemáticos comunes
    # (=+^{}[]<>√±≈∞×÷·, superíndices/subíndices Unicode) — antes se
    # borraban aunque la fórmula sí hubiera sobrevivido a la extracción
    # como texto real, desfigurándola igual. El pipe "|" también está en
    # la whitelist — es la sintaxis de columnas de las tablas Markdown
    # (_table_to_markdown/chunk_structured_blocks); sin él, la tabla
    # llegaba al chunker ya sin separadores de columna, rompiendo tanto el
    # chunk de tabla completa como el parseo de filas individuales
    # (_parse_markdown_table_rows) — encontrado con datos reales.
    text = re.sub(
        r"[^\w\s.,;:()/\-%¿?!áéíóúñÁÉÍÓÚÑ=+^{}\[\]<>√±≈∞×÷·²³¹⁰-ₜ|]",
        "",
        text,
    )
    text = re.sub(r"\x00", "", text)
    text = re.sub(r"[\x01-\x08\x0b-\x0c\x0e-\x1f\x7f-\x9f]", "", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"(\w+)-\s*\n\s*(\w+)", r"\1\2", text)
    return text.strip()


async def extract_normalize_and_hash(file_path: str) -> Dict[str, Any]:
    raw = await _extract_text(file_path)
    normalized = await _normalize_text(raw)
    if len(normalized) < 100:
        raise ValueError("Documento sin contenido de texto suficiente")
    return {
        "normalized_text": normalized,
        "document_hash": _generate_hash(normalized),
    }


async def chunk_text(text: str) -> List[Dict[str, Any]]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.CHUNK_SIZE,
        chunk_overlap=settings.CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", ", ", " ", ""],
    )
    chunks = []
    for i, chunk in enumerate(splitter.split_text(text)):
        urls = re.findall(r"https?://\S+", chunk)
        chunks.append({
            "chunk": i,
            "text": chunk.strip(),
            "urls": urls if urls else None,
        })
    return chunks


# Formatos con extracción por bloques tipados (tabla/texto) disponible —
# los únicos que hoy pasan por chunk_structured_blocks() en vez del
# chunk_text() clásico por conteo de caracteres. pptx/xlsx/ppt/imagen
# siguen el camino de siempre sin cambios.
_STRUCTURED_BLOCK_EXTRACTORS = {"pdf", "docx"}


async def _extract_blocks_for_structured_chunking(file_path: str) -> Optional[List[Dict[str, Any]]]:
    """None si el tipo de archivo no tiene extracción por bloques todavía
    (el caller debe caer al camino clásico extract_normalize_and_hash() +
    chunk_text())."""
    file_type = classify_file_type(Path(file_path).name)
    if file_type == "pdf":
        return await _extract_from_pdf_blocks(file_path)
    if file_type == "docx":
        return await _extract_from_docx_blocks(file_path)
    return None


async def extract_normalize_and_hash_blocks(file_path: str) -> Optional[Dict[str, Any]]:
    """Variante de extract_normalize_and_hash() que conserva los bloques
    tipados (tabla/texto/página) en vez de aplanarlos — la normalización
    (limpieza de caracteres de control, colapso de espacios) se aplica
    bloque por bloque para no perder la distinción tabla/texto. Devuelve
    None si el formato no tiene extracción por bloques (ver
    _STRUCTURED_BLOCK_EXTRACTORS) — el caller debe usar el camino clásico
    en ese caso."""
    blocks = await _extract_blocks_for_structured_chunking(file_path)
    if blocks is None:
        return None

    normalized_blocks = []
    for b in blocks:
        norm_text = await _normalize_text(b["text"])
        if norm_text:
            normalized_blocks.append({**b, "text": norm_text})

    full_text = "\n".join(b["text"] for b in normalized_blocks)
    if len(full_text) < 100:
        raise ValueError("Documento sin contenido de texto suficiente")

    return {
        "blocks": normalized_blocks,
        "document_hash": _generate_hash(full_text),
    }


def table_signature(md_text: str) -> Optional[Tuple[str, ...]]:
    """"Huella" de una tabla, usada para detectar si su ESTRUCTURA se
    repite en otros documentos del mismo curso (ej. cada sesión semanal
    trae su propia mini tabla "Datos Generales" sin completar, con las
    mismas 14 etiquetas de campo que la del sílabo real, pero vacía) —
    ver caller en el script de vectorización por lote, que decide con
    esto qué tablas NO conviene partir en filas (ver
    chunk_structured_blocks(..., repeated_signatures=...)).

    Para una tabla de 3+ columnas reales, la firma son sus encabezados de
    columna. Para una lista campo/valor de 2 columnas (sin encabezado
    real — ver _parse_markdown_table_rows), la firma es el conjunto de
    NOMBRES DE CAMPO (columna izquierda) de todas sus filas — dos
    documentos con la misma lista de campos (aunque los valores sean
    distintos) tienen la misma firma. None si la tabla no tiene una
    estructura reconocible (una sola columna)."""
    rows = _parse_markdown_table_rows(md_text)
    if not rows:
        return None
    if len(rows) == 1 and "" in rows[0]:
        return None  # tabla de una sola columna, sin estructura campo/valor
    if all(len(r) == 1 for r in rows):
        # 2 columnas (lista campo/valor): la firma son los nombres de campo.
        keys = tuple(sorted({k for r in rows for k in r if k}))
        return keys or None
    # 3+ columnas reales: la firma son los encabezados de columna (mismas
    # claves en cada fila).
    keys = tuple(sorted(rows[0].keys()))
    return keys or None


def chunk_structured_blocks(
    blocks: List[Dict[str, Any]],
    repeated_signatures: Optional[set] = None,
) -> List[Dict[str, Any]]:
    """Chunking consciente de tablas: agrupa bloques de texto/prosa
    consecutivos hasta acercarse a CHUNK_SIZE (mismo criterio de
    separadores que chunk_text()/RecursiveCharacterTextSplitter), pero a
    diferencia de chunk_text() una TABLA nunca se mezcla con texto vecino
    ni se parte a la mitad de una fila — cada tabla es su propio chunk (o,
    si es enorme, se parte solo dentro de sí misma, por filas/líneas de su
    propio Markdown).

    Nace de un caso real: la tabla "DATOS GENERALES" de un sílabo (14
    campos administrativos, incluido el nombre del docente) quedaba
    mezclada con texto narrativo vecino por chunk_text() clásico — el
    chunk resultante tenía tan poca especificidad temática (14 campos sin
    relación entre sí) que perdía en el ranking de retrieval contra
    chunks de prosa que solo repetían la palabra buscada. Separar la
    tabla en su propio chunk no cambia el contenido, pero le da un
    embedding/score BM25 enfocado solo en esos 14 campos, sin dilución.

    Cada chunk resultante suma dos campos nuevos frente a chunk_text():
    `chunk_type` ("table"/"prose"/"table_row") y `page` (número de página
    real para PDF — antes no existía a nivel de chunk; sigue None para
    DOCX, que no tiene concepto nativo de página).

    Fase 3.5 — además del chunk de tabla completa, cada FILA de la tabla
    se suma como un chunk propio (`chunk_type="table_row"`, vía
    _parse_markdown_table_rows). Es la capa que faltaba para tablas con
    varias filas heterogéneas: aislar la tabla entera no alcanza cuando la
    pregunta apunta a UNA fila puntual entre varias temáticamente
    distintas (visto en un caso real: "Examen sumativo 20%" perdido dentro
    de una tabla de 6 componentes de evaluación). No reemplaza la tabla
    completa — ambas conviven, cada una útil para preguntas distintas
    ("¿cuánto vale el examen?" vs "dame toda la tabla de evaluación").

    `repeated_signatures`: firmas (ver table_signature()) de tablas cuya
    ESTRUCTURA se repite en otros documentos del mismo curso — ej. cada
    sesión semanal trae su propia mini tabla "Datos Generales" sin
    completar, con las mismas etiquetas de campo que la del sílabo real
    pero vacía ("Docente: Apellidos y nombres / correo institucional").
    Si la firma de una tabla está en este set, NO se parte en filas (se
    deja solo como chunk de tabla completa) — partir 13 plantillas vacías
    casi idénticas en filas individuales genera decenas de fragmentos
    cortos y genéricos que, por pura cantidad, ahogan a la única fila real
    (la del documento que sí tiene los datos completos). Confirmado con
    datos reales: sin este filtro, "Docente: Apellidos y nombres / correo
    institucional" (plantilla vacía) ocupaba 4 de los primeros 10 puestos
    para la pregunta "¿quién es el docente?", por encima de la fila real
    con el nombre verdadero."""
    prose_splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.CHUNK_SIZE,
        chunk_overlap=settings.CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", ", ", " ", ""],
    )
    # Sin overlap y separando por línea/fila — una tabla enorme se parte
    # por filas de su propio Markdown, nunca a mitad de celda, y nunca se
    # le pega overlap de la tabla vecina.
    table_splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.CHUNK_SIZE, chunk_overlap=0,
        separators=["\n", " ", ""],
    )

    chunks: List[Dict[str, Any]] = []
    prose_buffer: List[Dict[str, Any]] = []
    # Encabezado más cercano detectado por heurística genérica (bloque de
    # texto corto, sin punto final — no se ancla a ningún texto específico
    # de ningún documento) — se propaga a los chunks siguientes hasta que
    # aparece un encabezado nuevo, igual que en cualquier documento real.
    last_heading: Optional[str] = None
    # Fix A.5-2: encabezado REAL de la última tabla vista (mismo documento),
    # para reusarlo cuando la tabla siguiente sea una continuación con el
    # header roto por un salto de página.
    prev_table_header: Optional[List[str]] = None

    def _looks_like_heading(text: str) -> bool:
        t = text.strip()
        return 0 < len(t) <= 80 and not t.endswith((".", ":", ";", ","))

    def _flush_prose():
        if not prose_buffer:
            return
        combined = "\n\n".join(b["text"] for b in prose_buffer)
        pages = sorted({b["page"] for b in prose_buffer if b.get("page") is not None})
        for piece in prose_splitter.split_text(combined):
            piece = piece.strip()
            if not piece:
                continue
            urls = re.findall(r"https?://\S+", piece)
            chunks.append({
                "chunk": len(chunks),
                "text": piece,
                "urls": urls or None,
                "chunk_type": "prose",
                "page": pages[0] if pages else None,
                "heading": last_heading,
            })
        prose_buffer.clear()

    for block in blocks:
        if block["type"] == "table":
            _flush_prose()
            for piece in table_splitter.split_text(block["text"]):
                piece = piece.strip()
                if not piece:
                    continue
                urls = re.findall(r"https?://\S+", piece)
                chunks.append({
                    "chunk": len(chunks),
                    "text": piece,
                    "urls": urls or None,
                    "chunk_type": "table",
                    "page": block.get("page"),
                    "heading": last_heading,
                })
            # Fase 3.5 — además de la tabla completa (arriba), cada FILA se
            # suma como su propio chunk aislado. Necesario para tablas con
            # varias filas heterogéneas (ej. ponderación de evaluación:
            # Participación, RSU, Examen sumativo...): el chunk de la tabla
            # completa diluye la relevancia de una fila puntual entre las
            # demás — visto en un caso real donde "Examen sumativo 20%"
            # quedó enterrado (rank 133/189) dentro de una tabla de 6 filas
            # de temas distintos. No reemplaza la tabla completa (sigue
            # sirviendo para "dame toda la tabla") — es una capa adicional,
            # igual que ya se documentó en el diseño de esta fase.
            #
            # EXCEPCIÓN: si la firma de esta tabla se repite en otros
            # documentos del curso (plantilla semanal, ver docstring), no
            # se parte en filas — solo el chunk de tabla completa de
            # arriba. Partir 13 plantillas casi idénticas en filas
            # individuales generaba decenas de fragmentos cortos que
            # ahogaban a la única fila real por pura cantidad.
            # Fix A.5-2: detectar si este bloque es la CONTINUACIÓN de la
            # tabla anterior con el header roto (tabla partida por salto de
            # página). Si el header comprimido tiene menos celdas que el
            # máximo de celdas reales de sus filas de cuerpo, es un header
            # roto; se reusa el header real de la tabla inmediatamente
            # anterior si su cantidad de columnas coincide con esas filas.
            override_header: Optional[List[str]] = None
            parts = _markdown_table_header_and_body(block["text"])
            if parts is not None:
                _, header_cells, body_lines = parts
                max_body = max((len(_split_md_row(l)) for l in body_lines), default=0)
                if max_body > len(header_cells):
                    if prev_table_header is not None and len(prev_table_header) == max_body:
                        override_header = prev_table_header
                else:
                    prev_table_header = header_cells or None
            else:
                prev_table_header = None

            sig = table_signature(block["text"]) if repeated_signatures else None
            if sig is not None and sig in repeated_signatures:
                continue
            for row in _parse_markdown_table_rows(block["text"], override_header=override_header):
                row_text = " | ".join(f"{k}: {v}" if k else v for k, v in row.items())
                if not row_text.strip():
                    continue
                # Una fila normalmente es corta (unos pocos campos), pero
                # una celda con un párrafo largo dentro (visto con datos
                # reales: filas de 5000+ caracteres, 5-6x CHUNK_SIZE) no
                # tenía ningún splitter — a diferencia de "table"/"prose"
                # arriba. Igual que esas dos, se parte con table_splitter
                # si excede CHUNK_SIZE; una fila normal (la inmensa
                # mayoría) nunca lo toca porque split_text() devuelve el
                # texto intacto cuando ya entra en un solo pedazo.
                for piece in table_splitter.split_text(row_text):
                    piece = piece.strip()
                    if not piece:
                        continue
                    chunks.append({
                        "chunk": len(chunks),
                        "text": piece,
                        "urls": None,
                        "chunk_type": "table_row",
                        "page": block.get("page"),
                        "heading": last_heading,
                    })
        else:
            # Un bloque de prosa rompe la continuidad: la próxima tabla ya no
            # es continuación de la anterior.
            prev_table_header = None
            if _looks_like_heading(block["text"]):
                last_heading = block["text"].strip()
            prose_buffer.append(block)
    _flush_prose()
    return chunks


def chunk_audio_segments(segments: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """
    Agrupa segmentos consecutivos de una transcripción (con start/end en
    segundos) hasta acercarse a CHUNK_SIZE caracteres, conservando el
    start_time/end_time de cada chunk resultante — así el asistente puede
    citar "video X, minuto mm:ss" en vez de perder los tiempos al
    re-particionar el texto plano.
    """
    chunks = []
    current_text = []
    current_start = None
    current_end = None

    def _flush():
        if not current_text:
            return
        text = " ".join(current_text).strip()
        if text:
            chunks.append({
                "chunk": len(chunks),
                "text": text,
                "urls": None,
                "start_time": current_start,
                "end_time": current_end,
            })

    for seg in segments:
        candidate = " ".join(current_text + [seg["text"]])
        if current_text and len(candidate) > settings.CHUNK_SIZE:
            _flush()
            current_text = [seg["text"]]
            current_start = seg["start"]
            current_end = seg["end"]
        else:
            if current_start is None:
                current_start = seg["start"]
            current_text.append(seg["text"])
            current_end = seg["end"]

    _flush()
    return chunks


# ═══════════════════════════════════════════════════════════════════
# VALIDACIÓN ESTRUCTURAL
# ═══════════════════════════════════════════════════════════════════

async def validate_document(file_path: str) -> Dict[str, Any]:
    """
    Valida estructura del documento.
    """
    path = Path(file_path)
    file_type = classify_file_type(path.name)
    stats = {"total_pages": 0, "total_images": 0, "total_tables": 0, "duration_seconds": None}

    if path.suffix.lower() == ".pdf":
        try:
            with pdfplumber.open(file_path) as pdf:
                stats["total_pages"] = len(pdf.pages)
                for page in pdf.pages:
                    tables = page.find_tables(table_settings={
                        "vertical_strategy": "lines",
                        "horizontal_strategy": "lines",
                        "intersection_tolerance": 5,
                        "snap_tolerance": 5,
                    })
                    stats["total_tables"] += len(tables)
                    relevant = [
                        img for img in page.images
                        if float(img.get("width", 0)) > 50
                        and float(img.get("height", 0)) > 50
                    ]
                    stats["total_images"] += len(relevant)

            doc = fitz.open(file_path)
            for page in doc:
                for img in page.get_images(full=True):
                    if img[2] > 100 and img[3] > 100:
                        stats["total_images"] += 1
            doc.close()
        except Exception as e:
            raise ValueError(f"Error validando PDF: {e}")

    elif path.suffix.lower() == ".docx":
        try:
            doc = DocxDocument(file_path)
            stats["total_tables"] = len(doc.tables)
            stats["total_images"] = sum(
                1 for rel in doc.part.rels.values() if "image" in rel.target_ref
            )
            total_words = sum(len(p.text.split()) for p in doc.paragraphs)
            stats["total_pages"] = max(1, total_words // 500)
        except Exception as e:
            raise ValueError(f"Error validando DOCX: {e}")

    elif file_type == "pptx":
        try:
            prs = Presentation(file_path)
            slides = list(prs.slides)
            stats["total_pages"] = len(slides)
            stats["total_images"] = sum(
                1 for slide in slides for shape in slide.shapes
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
            )
            stats["total_tables"] = sum(
                1 for slide in slides for shape in slide.shapes if shape.has_table
            )
        except Exception as e:
            raise ValueError(f"Error validando PPTX: {e}")

    elif file_type == "ppt":
        try:
            text = await _extract_from_ppt_legacy(file_path)
            total_words = len(text.split())
            stats["total_pages"] = max(1, total_words // 150)
        except Exception as e:
            raise ValueError(f"Error validando PPT: {e}")

    elif file_type == "xlsx":
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True)
            stats["total_pages"] = len(wb.worksheets)
            wb.close()
        except Exception as e:
            raise ValueError(f"Error validando XLSX: {e}")

    elif file_type == "image":
        stats["total_pages"] = 1

    elif file_type in ("audio", "video"):
        stats["duration_seconds"] = get_media_duration(file_path)

    elif file_type == "web":
        # Texto ya extraído (scraping.py) — validación liviana, mismo
        # criterio de "páginas" que ya usa DOCX (palabras / 500).
        text = Path(file_path).read_text(encoding="utf-8", errors="ignore")
        stats["total_pages"] = max(1, len(text.split()) // 500)

    else:
        raise ValueError(f"Formato no soportado para vectorización: '{path.suffix}'")

    warning = None
    if stats["total_images"] > 0 or stats["total_tables"] > 0:
        parts = []
        if stats["total_images"]:
            parts.append(f"{stats['total_images']} imagen(es)")
        if stats["total_tables"]:
            parts.append(f"{stats['total_tables']} tabla(s)")
        warning = (
            f"Se detectaron {' y '.join(parts)} en el documento. "
            "Solo se vectorizará el contenido textual."
        )
    elif file_type == "image":
        warning = "Imagen procesada con OCR — la calidad del texto depende de la nitidez de la imagen."
    elif file_type in ("audio", "video"):
        warning = "Contenido transcrito automáticamente con reconocimiento de voz — puede contener errores."
    elif file_type == "web":
        warning = "Página web externa — el sistema puede no distinguir contenido de menús/publicidad residual."

    return {**stats, "file_type": file_type, "warning_message": warning}


# ═══════════════════════════════════════════════════════════════════
# VECTORIZATION SERVICE
# ═══════════════════════════════════════════════════════════════════

class VectorizationService:
    """
    Servicio de vectorización ligado a un modelo de embeddings concreto.
    """

    def __init__(self, model_name: str):
        """
        model_name debe existir en AVAILABLE_EMBEDDING_MODELS (settings.py).
        Lanza ValueError si el modelo no está en el catálogo.
        """
        # Valida y obtiene la configuración del modelo
        info = get_model_info(model_name)

        self._model_name: str   = model_name
        self._vector_size: int  = info["dimensions"]
        self._use_prefix: bool  = info.get("use_prefix", False)
        self._use_task: bool    = info.get("use_task", False)
        self._task_passage: str = info.get("task_passage", "retrieval.passage")
        self._task_query: str   = info.get("task_query",   "retrieval.query")
        self._provider: str     = info.get("provider", "local")

        print(f"🔄 Cargando modelo : {self._model_name}")
        print(f"   Dimensión       : {self._vector_size}")
        print(f"   Proveedor       : {self._provider}")

        self._azure_client   = None
        self._azure_deployment = None
        self._openai_client  = None
        self._openai_model   = None

        if self._provider == "azure":
            self._model        = None
            self._azure_client = AzureOpenAI(
                api_key=settings.AZURE_EMBEDDING_KEY,
                azure_endpoint=settings.AZURE_EMBEDDING_ENDPOINT,
                api_version=settings.AZURE_EMBEDDING_API_VERSION,
            )
            self._azure_deployment = settings.AZURE_EMBEDDING_DEPLOYMENT
            print(f"   Endpoint Azure  : {settings.AZURE_EMBEDDING_ENDPOINT}")
            print(f"   Deployment      : {self._azure_deployment}")
        elif self._provider == "openai":
            # OpenAI API directa (no Azure) — mismo modelo pero sin pasar por
            # un deployment de Azure. model_name trae el prefijo "openai/"
            # (para distinguirlo del catálogo del mismo nombre bajo Azure);
            # la API de OpenAI espera el nombre sin ese prefijo.
            self._model         = None
            self._openai_client = OpenAI(api_key=settings.OPENAI_API_KEY)
            self._openai_model  = self._model_name.split("/", 1)[-1]
            print(f"   Modelo OpenAI   : {self._openai_model}")
        else:
            print(f"   Usa prefijos    : {self._use_prefix}")
            print(f"   Usa task LoRA   : {self._use_task}")
            self._model = SentenceTransformer(
                self._model_name,
                trust_remote_code=self._use_task,
            )

        print(f"🔄 Conectando a Qdrant: {settings.QDRANT_HOST}:{settings.QDRANT_PORT}")
        self._client = QdrantClient(
            url=f"http://{settings.QDRANT_HOST}:{settings.QDRANT_PORT}",
            timeout=60,
        )
        print(f"✅ VectorizationService listo [{self._model_name}]")

    # ── Propiedades ───────────────────────────────────────────────

    @property
    def model_name(self) -> str:
        return self._model_name

    @property
    def vector_size(self) -> int:
        return self._vector_size

    # ── Validación de compatibilidad con colección ─────────────────

    def validate_model_for_collection(self, collection_name: str) -> None:
        """
        Verifica que la dimensión del modelo coincida con la colección Qdrant.
        Lanza ValueError descriptivo si hay incompatibilidad.

        Para colecciones híbridas (vectores nombrados dense+sparse) no
        aplica: la dimensión efectiva la determina la estrategia de
        embedding denso activa (ver strategies/registry.py y
        core._build_pipeline_config), no el modelo fijo de este servicio
        — así que se omite la validación y se deja pasar.
        """
        if not self.collection_exists(collection_name):
            # Colección inexistente: la validación la hará el endpoint
            return

        try:
            info = self._client.get_collection(collection_name)
            vectors_config = info.config.params.vectors
        except Exception as e:
            raise ValueError(f"No se pudo consultar la colección '{collection_name}': {e}")

        if isinstance(vectors_config, dict):
            return

        col_dim = vectors_config.size
        if col_dim != self._vector_size:
            raise ValueError(
                f"⚠️  Incompatibilidad de dimensiones: la colección '{collection_name}' "
                f"fue creada con vectores de {col_dim} dims, pero el modelo "
                f"'{self._model_name}' genera {self._vector_size} dims. "
                f"Elige un modelo con {col_dim} dimensiones para esta colección."
            )

    # ── Colecciones ──────────────────────────────────────────────

    def list_collections(self) -> List[str]:
        resp = self._client.get_collections()
        return [c.name for c in resp.collections]

    def collection_exists(self, name: str) -> bool:
        return name in self.list_collections()

    # ── Documentos en colección ──────────────────────────────────

    def list_documents_in_collection(self, collection_name: str) -> List[Dict[str, Any]]:
        """
        Lista los documentos únicos vectorizados en una colección.
        """
        if not self.collection_exists(collection_name):
            return []

        seen: Dict[str, Dict] = {}
        offset = None

        while True:
            results, next_offset = self._client.scroll(
                collection_name=collection_name,
                limit=200,
                offset=offset,
                with_payload=["filename", "total_chunks", "date", "course_name", "section_name"],
                with_vectors=False,
            )
            for point in results:
                fn = point.payload.get("filename")
                if fn and fn not in seen:
                    seen[fn] = {
                        "filename": fn,
                        "total_chunks": point.payload.get("total_chunks", 0),
                        "date": point.payload.get("date", ""),
                        "course_name": point.payload.get("course_name", ""),
                        "section_name": point.payload.get("section_name", ""),
                    }
            if next_offset is None:
                break
            offset = next_offset

        return list(seen.values())

    # ── Check duplicado por filename ──────────────────────────────

    def check_duplicate(self, filename: str, collection_name: str) -> Dict[str, Any]:
        """Verifica si un filename ya está vectorizado en la colección."""
        if not self.collection_exists(collection_name):
            return {"exists": False, "total_chunks": 0, "document_hash": None}

        normalized_filename = _normalize_filename(filename)

        try:
            results, _ = self._client.scroll(
                collection_name=collection_name,
                limit=1,
                scroll_filter=Filter(must=[
                    FieldCondition(
                        key="filename_normalized",
                        match=MatchValue(value=normalized_filename)
                    )
                ]),
                with_payload=["document_hash", "total_chunks", "date", "source_type"],
                with_vectors=False,
            )

            if not results:
                results, _ = self._client.scroll(
                    collection_name=collection_name,
                    limit=1,
                    scroll_filter=Filter(must=[
                        FieldCondition(key="filename", match=MatchValue(value=filename))
                    ]),
                    with_payload=["document_hash", "total_chunks", "date", "source_type"],
                    with_vectors=False,
                )

            if not results:
                return {"exists": False, "total_chunks": 0, "document_hash": None, "date": None, "source_type": None}

            doc_hash = results[0].payload.get("document_hash")
            date = results[0].payload.get("date")
            source_type = results[0].payload.get("source_type")

            all_results, _ = self._client.scroll(
                collection_name=collection_name,
                limit=10000,
                scroll_filter=Filter(must=[
                    FieldCondition(key="document_hash", match=MatchValue(value=doc_hash))
                ]),
                with_payload=False,
                with_vectors=False,
            )
            return {
                "exists": True,
                "total_chunks": len(all_results),
                "document_hash": doc_hash,
                "date": date,
                "source_type": source_type,
            }
        except Exception as e:
            print(f"⚠️ Error check_duplicate: {e}")
            return {"exists": False, "total_chunks": 0, "document_hash": None, "date": None, "source_type": None}

    # ── Eliminar documento por filename ──────────────────────────

    def delete_document(self, filename: str, collection_name: str) -> Dict[str, Any]:
        check = self.check_duplicate(filename, collection_name)
        if not check["exists"]:
            return {"success": False, "message": f"'{filename}' no existe en la colección"}

        normalized_filename = _normalize_filename(filename)
        doc_hash = check["document_hash"]

        results, _ = self._client.scroll(
            collection_name=collection_name,
            limit=1,
            scroll_filter=Filter(must=[
                FieldCondition(
                    key="filename_normalized",
                    match=MatchValue(value=normalized_filename)
                )
            ]),
            with_payload=["document_hash"],
            with_vectors=False,
        )

        if not results:
            return {"success": False, "message": f"'{filename}' no encontrado en la colección"}

        self._client.delete(
            collection_name=collection_name,
            points_selector=Filter(must=[
                FieldCondition(key="document_hash", match=MatchValue(value=doc_hash))
            ]),
        )
        print(f"🗑️ Eliminado: {filename} ({check['total_chunks']} chunks)")
        return {
            "success": True,
            "deleted_chunks": check["total_chunks"],
            "filename": filename,
        }

    # ── Embed ─────────────────────────────────────────────────────

    async def embed(self, texts: List[str], is_query: bool = False) -> List[List[float]]:
        """
        Genera embeddings para la lista de textos.
        """
        loop = asyncio.get_event_loop()

        if self._provider == "azure":
            # Azure OpenAI Embeddings — batches de 100 (límite recomendado)
            async def _azure_embed() -> List[List[float]]:
                results = []
                batch_size = 100
                for i in range(0, len(texts), batch_size):
                    batch = texts[i:i + batch_size]
                    response = await loop.run_in_executor(
                        None,
                        lambda b=batch: self._azure_client.embeddings.create(
                            model=self._azure_deployment,
                            input=b,
                        )
                    )
                    results.extend([item.embedding for item in response.data])
                return results
            return await _azure_embed()

        if self._provider == "openai":
            # OpenAI Embeddings API directa — mismo esquema de batching que Azure.
            async def _openai_embed() -> List[List[float]]:
                results = []
                batch_size = 100
                for i in range(0, len(texts), batch_size):
                    batch = texts[i:i + batch_size]
                    response = await loop.run_in_executor(
                        None,
                        lambda b=batch: self._openai_client.embeddings.create(
                            model=self._openai_model,
                            input=b,
                        )
                    )
                    results.extend([item.embedding for item in response.data])
                return results
            return await _openai_embed()

        if self._use_task:
            # Jina Embeddings v3
            task = self._task_query if is_query else self._task_passage
            return await loop.run_in_executor(
                None,
                lambda: self._model.encode(
                    texts,
                    task=task,
                    normalize_embeddings=True,
                ).tolist(),
            )

        if self._use_prefix:
            # Familia E5
            prefix = "query: " if is_query else "passage: "
            inputs = [prefix + t for t in texts]
        else:
            inputs = texts

        return await loop.run_in_executor(
            None,
            lambda: self._model.encode(inputs, normalize_embeddings=True).tolist(),
        )

    # ── Store ─────────────────────────────────────────────────────

    async def store_vectors(
        self,
        *,
        collection_name: str,
        chunks: List[Dict[str, Any]],
        embeddings: List[List[float]],
        metadata: Dict[str, Any],
    ) -> Dict[str, Any]:
        """
        Almacena vectores en Qdrant.
        metadata debe incluir: document_hash, filename, format,
        total_pages, total_chunks, course_name.

        Valida compatibilidad de dimensiones antes de insertar.
        """
        if not self.collection_exists(collection_name):
            raise ValueError(f"La colección '{collection_name}' no existe en Qdrant.")

        self.validate_model_for_collection(collection_name)

        doc_hash = metadata["document_hash"]
        now = datetime.utcnow().isoformat()
        filename_normalized = _normalize_filename(metadata["filename"])
        points = []

        for chunk, vector in zip(chunks, embeddings):
            pid = hashlib.md5(
                f"{doc_hash}_{chunk['chunk']}".encode()
            ).hexdigest()
            points.append(PointStruct(
                id=pid,
                vector=vector,
                payload={
                    "document_hash":      doc_hash,
                    "filename":           metadata["filename"],
                    "filename_normalized": filename_normalized,
                    "format":             metadata["format"],
                    "file_type":          metadata.get("file_type", metadata["format"]),
                    "total_pages":        metadata["total_pages"],
                    "total_chunks":       metadata["total_chunks"],
                    "course_name":        metadata.get("course_name", "modo manual"),
                    "date":               now,
                    "chunk":              chunk["chunk"],
                    "urls":               chunk["urls"],
                    "text":               chunk["text"],
                    "start_time":         chunk.get("start_time"),
                    "end_time":           chunk.get("end_time"),
                    "embedding_model":    self._model_name,
                },
            ))

        batch_size = 100
        for i in range(0, len(points), batch_size):
            self._client.upsert(
                collection_name=collection_name,
                points=points[i:i + batch_size],
                wait=True,
            )

        print(f"💾 {len(points)} vectores → {collection_name} [{self._model_name}]")
        return {
            "success": True,
            "vectors_stored": len(points),
            "collection": collection_name,
            "document_hash": doc_hash,
            "embedding_model": self._model_name,
        }


# ═══════════════════════════════════════════════════════════════════
# REGISTRY DE SERVICIOS (una instancia por modelo)
# ═══════════════════════════════════════════════════════════════════

_service_registry: Dict[str, VectorizationService] = {}


def get_service(model_name: Optional[str] = None) -> VectorizationService:
    """
    Retorna (o crea) la instancia de VectorizationService para el modelo dado.
    """
    key = model_name or settings.EMBEDDING_MODEL

    # Valida que el modelo exista en el catálogo antes de instanciar
    get_model_info(key)  # lanza ValueError si no está

    if key not in _service_registry:
        print(f"🆕 Instanciando servicio para modelo: {key}")
        _service_registry[key] = VectorizationService(model_name=key)

    return _service_registry[key]


# ═══════════════════════════════════════════════════════════════════
# MOODLE — funciones compartidas
# ═══════════════════════════════════════════════════════════════════

def _add_moodle_token(url: str) -> str:
    """
    Agrega el token de Moodle a una URL de pluginfile.php si no lo trae ya.
    Necesario para que el enlace "Ver" del listado (que el frontend abre
    directo, sin pasar por el backend) pueda abrir el archivo sin sesión.
    """
    if not url or "/pluginfile.php" not in url or "token=" in url:
        return url
    sep = "&" if "?" in url else "?"
    return f"{url}{sep}token={settings.MOODLE_TOKEN}"


_YOUTUBE_RE = re.compile(
    r"(?:youtube\.com/(?:watch\?v=|embed/|shorts/)|youtu\.be/)[\w-]{6,}",
    re.IGNORECASE,
)

# Enlaces "url" de Moodle que apuntan a herramientas funcionales (un
# formulario, una videollamada) en vez de contenido real — no tiene
# sentido scrapearlos como página web (no hay artículo que extraer, en el
# mejor caso texto de interfaz sin valor, en el peor una página de login).
# No es (ni pretende ser) una lista exhaustiva: el filtro real y general
# es que HybridVectorizationService.process_and_store exige mínimo 100
# caracteres de texto extraído — esto solo evita gastar una consulta HTTP
# en los casos más comunes y obvios.
_NON_CONTENT_URL_RE = re.compile(
    r"forms\.gle|docs\.google\.com/forms|zoom\.us|meet\.google\.com",
    re.IGNORECASE,
)


def _institutional_domain() -> str:
    """Dominio propio de la institución (derivado de MOODLE_URL, quitando
    el subdominio de Moodle) — usado para distinguir una página web
    "externa" real de un enlace interno a otra herramienta de la propia
    institución. Visto en cursos reales: enlaces "url" de Moodle que
    apuntan al ERP/biblioteca, al repositorio de tesis, o incluso a este
    mismo gestor (educaia.uladech.edu.pe/f1, /f2, /f3) — ninguno de esos
    es "contenido" vectorizable, y varios de esos GET tienen efectos
    secundarios reales (p.ej. crear una sesión de chat) que no hay que
    disparar por accidente al scrapear."""
    host = (urlparse(settings.MOODLE_URL).hostname or "").lower()
    parts = host.split(".")
    return ".".join(parts[1:]) if len(parts) > 2 else host


def _is_external_web_link(link: str) -> bool:
    """True si `link` es una página web real fuera del ecosistema de la
    institución — ver _institutional_domain()."""
    if not link.startswith(("http://", "https://")):
        return False
    if _NON_CONTENT_URL_RE.search(link):
        return False
    host = (urlparse(link).hostname or "").lower()
    institutional = _institutional_domain()
    if institutional and (host == institutional or host.endswith("." + institutional)):
        return False
    # Herramientas propias corriendo en local/red interna (visto en un
    # curso real: un link a "localhost:7002" apuntando a una herramienta
    # de evaluación) — nunca son contenido externo real, sin importar el
    # dominio configurado en MOODLE_URL.
    if host == "localhost" or host.startswith(("127.", "10.", "192.168.")):
        return False
    return True


def get_moodle_user_info(userid: int) -> Dict[str, Optional[str]]:
    """Nombre completo y username del usuario de Moodle (para el saludo del
    asistente público y para dejar `moodle_username` bien poblado en
    asistente_sesiones — antes solo se pedía `fullname`, aunque
    core_user_get_users_by_field ya trae `username` en la misma respuesta
    sin costo adicional). Usa core_user_get_users_by_field — verificado en
    vivo que el MOODLE_TOKEN actual tiene permiso para esta función.
    Retorna {"fullname": None, "username": None} si el usuario no existe o
    la llamada falla (el saludo genérico es el fallback, no debe romper la
    ruta pública)."""
    empty = {"fullname": None, "username": None}
    try:
        url = f"{settings.MOODLE_URL}/webservice/rest/server.php"
        params = {
            "wstoken": settings.MOODLE_TOKEN,
            "wsfunction": "core_user_get_users_by_field",
            "moodlewsrestformat": "json",
            "field": "id",
            "values[0]": userid,
        }
        r = requests.get(url, params=params, verify=False, timeout=10)
        r.raise_for_status()
        data = r.json()
        if isinstance(data, dict) and "exception" in data:
            print(f"⚠️ Moodle no pudo resolver el usuario {userid}: {data.get('message')}")
            return empty
        if data and isinstance(data, list):
            return {"fullname": data[0].get("fullname"), "username": data[0].get("username")}
    except Exception as e:
        print(f"⚠️ Error consultando nombre de usuario Moodle {userid}: {e}")
    return empty


def get_moodle_users_info(userids: List[int]) -> Dict[int, Dict[str, Optional[str]]]:
    """Igual que get_moodle_user_info pero para varios userids en una sola
    llamada (core_user_get_users_by_field acepta varios `values[]`) — usado
    para el backfill masivo de moodle_username en asistente_sesiones, para
    no hacer una consulta HTTP por sesión."""
    if not userids:
        return {}
    try:
        url = f"{settings.MOODLE_URL}/webservice/rest/server.php"
        params = {
            "wstoken": settings.MOODLE_TOKEN,
            "wsfunction": "core_user_get_users_by_field",
            "moodlewsrestformat": "json",
            "field": "id",
        }
        for i, uid in enumerate(userids):
            params[f"values[{i}]"] = uid
        r = requests.get(url, params=params, verify=False, timeout=20)
        r.raise_for_status()
        data = r.json()
        if isinstance(data, dict) and "exception" in data:
            print(f"⚠️ Moodle no pudo resolver el lote de usuarios: {data.get('message')}")
            return {}
        if data and isinstance(data, list):
            return {
                u["id"]: {"fullname": u.get("fullname"), "username": u.get("username")}
                for u in data
            }
    except Exception as e:
        print(f"⚠️ Error consultando lote de usuarios Moodle: {e}")
    return {}


def get_moodle_course_name(courseid: int) -> Optional[str]:
    """Nombre completo (`fullname`) del curso de Moodle, para confirmar en
    el paso "Curso" del wizard de asistentes. Usa core_course_get_courses
    (options[ids][]) — verificado en vivo que el MOODLE_TOKEN actual tiene
    permiso para esta función; core_course_get_courses_by_field está
    bloqueada para este token. Retorna None si el curso no existe o la
    llamada falla."""
    try:
        url = f"{settings.MOODLE_URL}/webservice/rest/server.php"
        params = {
            "wstoken": settings.MOODLE_TOKEN,
            "wsfunction": "core_course_get_courses",
            "moodlewsrestformat": "json",
            "options[ids][0]": courseid,
        }
        r = requests.get(url, params=params, verify=False, timeout=10)
        r.raise_for_status()
        data = r.json()
        if isinstance(data, dict) and "exception" in data:
            print(f"⚠️ Moodle no pudo resolver el curso {courseid}: {data.get('message')}")
            return None
        if data and isinstance(data, list):
            return data[0].get("fullname")
    except Exception as e:
        print(f"⚠️ Error consultando nombre de curso Moodle {courseid}: {e}")
    return None


def get_course_resources(curid: int) -> List[Dict[str, Any]]:
    """
    Consulta la API de Moodle y retorna los recursos (módulos) del curso.

    Usa core_course_get_contents en vez de mod_resource_get_resources_by_courses
    porque este último solo devuelve módulos de tipo "Archivo" (mod_resource) y
    omite los archivos dentro de módulos "Carpeta" (mod_folder). core_course_get_contents
    trae todos los tipos de módulo con sus archivos, así que se filtran aquí
    "resource" y "folder" y se normaliza la salida al formato
    {"id", "name", "contentfiles": [...]} que ya consume el resto del sistema.

    También procesa módulos "url" (enlaces externos): si el enlace apunta a
    YouTube se agrega como un archivo de video sintético (source="youtube")
    que el resto del pipeline transcribe (subtítulos si el video los tiene,
    o audio+ASR si no — ver scraping.py y core.py). Cualquier otro enlace
    http(s) real se agrega como una página web sintética (source="web") que
    se scrapea (texto principal del artículo, ver scraping.py). Enlaces sin
    URL real (formularios, Zoom sin link directo, etc.) se ignoran porque
    no son contenido descargable/vectorizable.
    """
    url = f"{settings.MOODLE_URL}/webservice/rest/server.php"
    params = {
        "wstoken": settings.MOODLE_TOKEN,
        "wsfunction": "core_course_get_contents",
        "moodlewsrestformat": "json",
        "courseid": curid,
    }
    r = requests.get(url, params=params, verify=False)
    r.raise_for_status()
    data = r.json()

    if isinstance(data, dict) and "exception" in data:
        raise RuntimeError(f"Moodle API error: {data.get('message', data)}")

    resources = []
    for section in data:
        section_name = section.get("name") or ""
        for module in section.get("modules", []):
            modname = module.get("modname")
            if modname not in ("resource", "folder", "url"):
                continue

            if modname == "url":
                contentfiles = []
                for item in module.get("contents", []):
                    link = item.get("fileurl") or ""
                    if item.get("type") != "url":
                        continue
                    # contenthash/timemodified: casi nunca vienen para un
                    # recurso "url" (no es un archivo real subido a Moodle),
                    # pero se pasan igual cuando existan — es lo único que
                    # curso_archivos (ledger.py) tiene para detectar que el
                    # profesor editó el enlace. Limitación real y conocida:
                    # nunca detecta que el CONTENIDO de la página/video
                    # externo cambió por su cuenta, solo que el bloque de
                    # Moodle se tocó.
                    if _YOUTUBE_RE.search(link):
                        module_name = module.get("name") or "video_youtube"
                        contentfiles.append({
                            "filename": f"{module_name}.mp4",
                            "filesize": None,
                            "fileurl":  link,
                            "mimetype": "video/youtube",
                            "source":   "youtube",
                            "contenthash": item.get("contenthash"),
                            "timemodified": item.get("timemodified"),
                        })
                    elif _is_external_web_link(link):
                        # Página web externa (no YouTube, fuera del propio
                        # dominio institucional) — se scrapea su texto
                        # principal en vez de descargarse como archivo.
                        # Enlaces no-http(s), a herramientas funcionales, o
                        # a la propia institución (Moodle, este gestor, el
                        # ERP, la biblioteca) se ignoran: no son contenido
                        # vectorizable — ver _is_external_web_link().
                        module_name = module.get("name") or "pagina_web"
                        contentfiles.append({
                            "filename": f"{module_name}.txt",
                            "filesize": None,
                            "fileurl":  link,
                            "mimetype": "text/html",
                            "source":   "web",
                            "contenthash": item.get("contenthash"),
                            "timemodified": item.get("timemodified"),
                        })
            else:
                contentfiles = [
                    {
                        "filename": f["filename"],
                        "filesize": f.get("filesize"),
                        "fileurl":  _add_moodle_token(f.get("fileurl") or ""),
                        "mimetype": f.get("mimetype"),
                        # contenthash: SHA1 nativo de Moodle sobre los bytes del
                        # archivo — cambia solo si el archivo en sí cambió (a
                        # diferencia de timemodified, que Moodle puede tocar por
                        # cambios de visibilidad/permisos sin tocar el contenido).
                        # Usado por curso_archivos (ver rds.py) como filtro barato
                        # de "vale la pena re-extraer para confirmar" antes de
                        # gastar cómputo en re-extraer y re-embeder.
                        "contenthash": f.get("contenthash"),
                        "timemodified": f.get("timemodified"),
                    }
                    for f in module.get("contents", [])
                    if f.get("type") == "file"
                ]

            if not contentfiles:
                continue

            resources.append({
                "id":            module.get("id"),
                "name":          module.get("name"),
                "section_name":  section_name,
                "contentfiles":  contentfiles,
            })

    return resources


def _unique_temp_name(filename: str) -> str:
    """Nombre de archivo temporal, único por descarga (prefijo aleatorio de
    8 hex). settings.TEMP_DIR se comparte entre CUALQUIER descarga
    concurrente del proceso — dos migraciones corriendo al mismo tiempo
    (periódica + manual, o dos colecciones distintas cuyo curso real
    comparte nombres de archivo, ej. una RD y su colección "hija" clonada)
    antes escribían las dos al MISMO path (derivado solo del filename
    sanitizado) y la limpieza de una borraba el archivo que la otra
    todavía estaba usando — encontrado con datos reales en producción
    (colisión real entre Curso_1472 y Curso_173, mismos nombres de
    archivo, ambos con actividad de vectorización al mismo tiempo).
    Ningún caller depende de que el nombre del archivo temporal sea
    exactamente el filename original — ese nombre real siempre se pasa
    aparte (`original_filename`/`filename`) a quien vectoriza."""
    return f"{uuid.uuid4().hex[:8]}_{_sanitize_filename(filename)}"


def download_pdf_to_temp(fileurl: str, filename: str) -> Path:
    """
    Descarga un archivo (de cualquier tipo soportado) directamente desde el
    webservice de Moodle usando el token (wstoken habilitado para descarga
    de archivos) — ya no requiere login por sesión con usuario/contraseña.
    Valida la firma binaria del archivo descargado según su tipo.
    """
    download_url = fileurl
    if "token=" not in download_url:
        sep = "&" if "?" in download_url else "?"
        download_url = f"{download_url}{sep}token={settings.MOODLE_TOKEN}"

    dest = settings.TEMP_DIR / _unique_temp_name(filename)
    file_type = classify_file_type(filename)

    with requests.get(download_url, stream=True, allow_redirects=True, verify=False) as r:
        r.raise_for_status()
        with open(dest, "wb") as f:
            for chunk in r.iter_content(256 * 1024):
                if chunk:
                    f.write(chunk)

    with open(dest, "rb") as f:
        header = f.read(16)
    if not header or not _validate_file_signature(file_type, header):
        dest.unlink(missing_ok=True)
        raise ValueError(f"Archivo corrupto o formato inesperado: {filename}")

    print(f"⬇️ Descargado en temp: {dest}")
    return dest


def write_temp_text(filename: str, text: str) -> Path:
    """Guarda texto ya extraído (página web escrapeada, o el texto plano de
    una transcripción de YouTube por subtítulos) como un .txt real en
    temp_docs/ — así sigue el mismo camino que cualquier archivo descargado
    (validate_document/_extract_text lo leen igual, cleanup_temp_file lo
    borra igual al terminar)."""
    dest = settings.TEMP_DIR / _unique_temp_name(filename)
    dest.write_text(text, encoding="utf-8")
    return dest


def download_youtube_audio(youtube_url: str, filename: str) -> Path:
    """
    Descarga solo el audio (mejor calidad disponible) de un video de YouTube
    a la carpeta temporal, usando yt-dlp. No requiere ffmpeg porque no se
    mezcla audio+video (el contenedor descargado se transcribe tal cual con
    faster-whisper, que decodifica cualquier formato vía PyAV).
    """
    safe_stem = Path(_unique_temp_name(filename)).stem
    outtmpl = str(settings.TEMP_DIR / f"{safe_stem}.%(ext)s")

    ydl_opts = {
        "format": "bestaudio/best",
        "outtmpl": outtmpl,
        "quiet": True,
        "no_warnings": True,
        "noplaylist": True,
    }

    with yt_dlp.YoutubeDL(ydl_opts) as ydl:
        info = ydl.extract_info(youtube_url, download=True)
        dest = Path(ydl.prepare_filename(info))

    if not dest.exists() or dest.stat().st_size == 0:
        raise ValueError(f"No se pudo descargar el audio de YouTube: {youtube_url}")

    print(f"⬇️ Audio de YouTube descargado: {dest}")
    return dest


def cleanup_temp_file(path: Path) -> None:
    """Elimina un archivo de la carpeta temporal."""
    try:
        path.unlink(missing_ok=True)
        print(f"🗑️ Temp eliminado: {path.name}")
    except Exception as e:
        print(f"⚠️ No se pudo eliminar temp {path}: {e}")


# ═══════════════════════════════════════════════════════════════════
# PIPELINE HÍBRIDO (Strategy Pattern) — colecciones con vectores nombrados
# ═══════════════════════════════════════════════════════════════════

class HybridVectorizationService:
    """Vectorización dirigida por configuración (Strategy Pattern) para
    colecciones híbridas (dense+sparse nombrados). No reemplaza a
    VectorizationService, que sigue sirviendo intacta a las colecciones
    legacy (vector único sin nombre) — ver vectorize_file_with_pages()."""

    def __init__(self, config: "PipelineConfig"):
        self._config = config

    async def process_and_store(
        self,
        *,
        file_path: Path,
        collection_name: str,
        course_name: str,
        total_pages: int,
        original_filename: str = "",
        section_name: str = "",
        pretranscribed_segments: Optional[List[Dict[str, Any]]] = None,
    ) -> Dict[str, Any]:
        from strategies import registry as strategy_registry
        from qdrant_admin import get_qdrant_admin

        stored_filename = original_filename if original_filename else file_path.name
        file_type = classify_file_type(stored_filename)

        etl_document = strategy_registry.get_etl_document_strategy(
            self._config.etl_document_strategy
        )
        etl_audio = strategy_registry.get_etl_audio_strategy(self._config.etl_audio_strategy)
        contextual = strategy_registry.get_contextual_strategy(self._config.contextual_strategy)
        dense = strategy_registry.get_dense_strategy(self._config.dense_strategy)
        sparse = strategy_registry.get_sparse_strategy(self._config.sparse_strategy)

        use_contextual = self._config.contextual_strategy not in (None, "none")

        if file_type in ("audio", "video"):
            # `pretranscribed_segments`: transcripción de YouTube obtenida
            # por subtítulos/CC ya existentes (ver scraping.py/core.py) — se
            # usa tal cual, sin volver a transcribir con la estrategia
            # configurada, conservando start_time/end_time por chunk.
            segments = (
                pretranscribed_segments if pretranscribed_segments is not None
                else await etl_audio.transcribe(str(file_path))
            )
            full_text = " ".join(s["text"] for s in segments)
            if len(full_text.strip()) < 20:
                raise ValueError("Transcripción sin contenido de texto suficiente")
            document_hash = _generate_hash(full_text)
            chunks = chunk_audio_segments(segments)
        elif file_type == "web":
            # Página web externa: el texto ya viene extraído (scraping.py) —
            # nunca pasa por etl_document_strategy (algunas estrategias
            # asumen PDF real/render de páginas, incompatible con texto
            # plano ya escrapeado).
            raw_text = Path(file_path).read_text(encoding="utf-8", errors="ignore")
            full_text = await _normalize_text(raw_text)
            if len(full_text) < 100:
                raise ValueError(
                    "La página no contiene suficiente texto extraíble (mínimo 100 caracteres)"
                )
            document_hash = _generate_hash(full_text)
            chunks = await chunk_text(full_text)
        else:
            # Chunking consciente de tablas (ver chunk_structured_blocks): una
            # tabla nunca se mezcla con texto vecino ni se parte a mitad de
            # fila, y cada fila se suma como su propio chunk — es lo que hace
            # posible el Índice de Campos Estructurados (campos_estructurados.py).
            # Disponible hoy solo para pdf/docx (_STRUCTURED_BLOCK_EXTRACTORS);
            # cualquier otro formato de documento (pptx/xlsx/imagen) cae al
            # camino clásico sin cambios.
            #
            # `repeated_signatures` se deja siempre vacío ACÁ: distinguir una
            # tabla-plantilla real (su estructura se repite en 2+ documentos)
            # de una tabla única requiere ver TODO el curso a la vez, no un
            # archivo aislado — y este método procesa un archivo por llamada.
            # Con el valor por defecto, cada tabla se parte en filas como si
            # fuera única (seguro: nunca produce una respuesta incorrecta,
            # como mucho unas filas de plantilla vacía de más hasta la
            # siguiente sincronización). La clasificación correcta se corrige
            # DESPUÉS, sin re-embeber nada, con
            # campos_estructurados.sync_after_vectorization() — ver el caller
            # en core.py, que la llama una vez terminado todo el curso/lote
            # (con visión de la colección completa, ahí sí se puede decidir
            # bien qué tablas son plantillas repetidas).
            blocks_result = await extract_normalize_and_hash_blocks(str(file_path))
            if blocks_result is not None:
                document_hash = blocks_result["document_hash"]
                full_text = "\n".join(b["text"] for b in blocks_result["blocks"])
                chunks = chunk_structured_blocks(blocks_result["blocks"], repeated_signatures=None)
            else:
                raw_text = await etl_document.extract(str(file_path))
                full_text = await _normalize_text(raw_text)
                if len(full_text) < 100:
                    raise ValueError(
                        "El documento no contiene suficiente texto extraíble "
                        "(mínimo 100 caracteres)"
                    )
                document_hash = _generate_hash(full_text)
                chunks = await chunk_text(full_text)

        # Contextual Retrieval: cada chunk de PROSA se enriquece con una
        # cabecera de 50-100 palabras (si la estrategia activa no es "none")
        # y el texto que se vectoriza es cabecera+chunk, aunque el payload
        # conserva el texto crudo por separado. Un chunk de tabla/fila de
        # tabla (chunk_type table/table_row) NO pasa por el enriquecimiento
        # LLM — un resumen generado de "Docente: Fulano de Tal" agrega poco
        # y cuesta una llamada más; en su lugar lleva una cabecera
        # determinística (curso/documento/sección/tipo), igual que ya se
        # probó y validó para Curso_1472/Curso_1541 esta sesión.
        embed_texts: List[str] = []
        for chunk in chunks:
            chunk_type = chunk.get("chunk_type", "prose")
            if chunk_type in ("table", "table_row"):
                header_parts = [f"Curso: {course_name}"] if course_name else []
                header_parts.append(f"Documento: {stored_filename}")
                if chunk.get("heading"):
                    header_parts.append(f"Sección: {chunk['heading']}")
                header_parts.append(
                    f"Tipo: {'tabla de datos' if chunk_type == 'table' else 'fila de tabla'}"
                )
                header = " | ".join(header_parts)
            else:
                header = await contextual.enrich(full_text, chunk["text"]) if use_contextual else ""
            chunk["contextual_header"] = header
            embed_texts.append(f"{header}\n\n{chunk['text']}" if header else chunk["text"])

        dense_vectors = await dense.embed(embed_texts)
        sparse_vectors = await sparse.embed(embed_texts)

        filename_normalized = _normalize_filename(stored_filename)
        now = datetime.utcnow().isoformat()
        points = []
        for chunk, dense_vec, sparse_vec in zip(chunks, dense_vectors, sparse_vectors):
            pid = hashlib.md5(f"{document_hash}_{chunk['chunk']}".encode()).hexdigest()
            points.append({
                "id": pid,
                "dense_vector": dense_vec,
                "sparse_indices": sparse_vec["indices"],
                "sparse_values": sparse_vec["values"],
                "payload": {
                    "document_hash": document_hash,
                    "filename": stored_filename,
                    "filename_normalized": filename_normalized,
                    "format": file_path.suffix.lstrip("."),
                    "file_type": file_type,
                    "total_pages": total_pages,
                    "total_chunks": len(chunks),
                    "course_name": course_name,
                    "section_name": section_name,
                    "date": now,
                    "chunk": chunk["chunk"],
                    "urls": chunk.get("urls"),
                    "text": chunk["text"],
                    "contextual_header": chunk.get("contextual_header") or None,
                    "start_time": chunk.get("start_time"),
                    "end_time": chunk.get("end_time"),
                    "slide_number": chunk.get("slide_number"),
                    "source_url": chunk.get("source_url"),
                    "source_type": self._config.source_type,
                    "embedding_model": self._config.dense_strategy,
                    "sparse_model": self._config.sparse_strategy,
                    # chunk_type/page/is_repeated_template: solo tienen valor
                    # real para pdf/docx vía chunk_structured_blocks (arriba);
                    # para el resto de formatos chunk_type queda en "prose"
                    # (valor por defecto, sin romper compatibilidad con el
                    # campos_estructurados.py existente, que solo actúa sobre
                    # chunk_type table/table_row). is_repeated_template se dejar
                    # siempre False acá — se corrige después, ver
                    # campos_estructurados.sync_after_vectorization().
                    "chunk_type": chunk.get("chunk_type", "prose"),
                    "page": chunk.get("page"),
                    "is_repeated_template": False,
                },
            })

        admin = get_qdrant_admin()
        result = admin.upsert_hybrid_points(collection_name=collection_name, points=points)

        return {
            "success": True,
            "vectors_stored": result["points_upserted"],
            "collection": collection_name,
            "document_hash": document_hash,
            "embedding_model": self._config.dense_strategy,
            "sparse_model": self._config.sparse_strategy,
            "pipeline": "hybrid",
        }


# ═══════════════════════════════════════════════════════════════════
# PIPELINE DE VECTORIZACIÓN (función central compartida)
# ═══════════════════════════════════════════════════════════════════

async def vectorize_file(
    file_path: Path,
    collection_name: str,
    course_name: str,
    service: VectorizationService,
) -> Dict[str, Any]:
    """
    Pipeline completo para un archivo ya guardado en disco:
    extract → normalize → hash → chunk → embed → store.
    """
    extracted = await extract_normalize_and_hash(str(file_path))
    normalized_text = extracted["normalized_text"]
    document_hash   = extracted["document_hash"]

    chunks     = await chunk_text(normalized_text)
    embeddings = await service.embed([c["text"] for c in chunks])

    metadata = {
        "document_hash": document_hash,
        "filename":      file_path.name,
        "format":        file_path.suffix.lstrip("."),
        "total_pages":   0,
        "total_chunks":  len(chunks),
        "course_name":   course_name,
    }

    return await service.store_vectors(
        collection_name=collection_name,
        chunks=chunks,
        embeddings=embeddings,
        metadata=metadata,
    )


async def vectorize_file_with_pages(
    file_path: Path,
    collection_name: str,
    course_name: str,
    total_pages: int,
    service: VectorizationService,
    original_filename: str = "",
    pipeline_config: Optional["PipelineConfig"] = None,
    section_name: str = "",
    pretranscribed_segments: Optional[List[Dict[str, Any]]] = None,
) -> Dict[str, Any]:
    """
    Pipeline completo: extract → normalize → hash → chunk → embed → store.
    Para audio/video usa transcripción + chunking por segmentos (conserva
    timestamps); para el resto usa el pipeline de texto plano habitual.

    Si `pipeline_config` no es None (colección híbrida con la configuración
    RAG activa — ver core.py), delega en HybridVectorizationService y usa
    ETL/contextual/embedding denso+disperso según lo configurado. Si es
    None (colección legacy), el resto de esta función queda intacto: mismo
    comportamiento que antes de introducir el pipeline configurable.

    `section_name`: nombre de la sección/semana de Moodle que contiene el
    módulo (distinto de `course_name`, que es el nombre del módulo/recurso
    en sí) — permite responder preguntas de navegación como "¿en qué semana
    está este archivo?" sin depender del contenido del PDF.
    """
    if pipeline_config is not None:
        hybrid_service = HybridVectorizationService(pipeline_config)
        return await hybrid_service.process_and_store(
            file_path=file_path,
            collection_name=collection_name,
            course_name=course_name,
            total_pages=total_pages,
            original_filename=original_filename,
            section_name=section_name,
            pretranscribed_segments=pretranscribed_segments,
        )

    stored_filename = original_filename if original_filename else file_path.name
    file_type = classify_file_type(stored_filename)

    if file_type in ("audio", "video"):
        segments = (
            pretranscribed_segments if pretranscribed_segments is not None
            else await transcribe_audio_or_video(str(file_path))
        )
        full_text = " ".join(s["text"] for s in segments)
        if len(full_text.strip()) < 20:
            raise ValueError("Transcripción sin contenido de texto suficiente")
        document_hash = _generate_hash(full_text)
        chunks = chunk_audio_segments(segments)
    else:
        extracted = await extract_normalize_and_hash(str(file_path))
        document_hash = extracted["document_hash"]
        chunks = await chunk_text(extracted["normalized_text"])

    embeddings = await service.embed([c["text"] for c in chunks])

    metadata = {
        "document_hash": document_hash,
        "filename":      stored_filename,
        "format":        file_path.suffix.lstrip("."),
        "file_type":     file_type,
        "total_pages":   total_pages,
        "total_chunks":  len(chunks),
        "course_name":   course_name,
        "section_name":  section_name,
    }

    return await service.store_vectors(
        collection_name=collection_name,
        chunks=chunks,
        embeddings=embeddings,
        metadata=metadata,
    )
